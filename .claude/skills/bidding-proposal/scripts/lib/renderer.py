from __future__ import annotations

import base64
import copy
import os
import re
import tempfile
from pathlib import Path
from typing import Dict
from typing import List
from typing import Optional
from typing import Tuple

from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt

from lib.content_condenser import (
    DEFAULT_BULLET,
    condense,
    count_items,
    describe,
    emitted_items,
)
from lib.schemas import ContentItem
from lib.schemas import ContentType
from lib.schemas import SlideContent
from lib.templates import SVNProposalTemplate
from lib.templates.base import ContentKind, ShapeRole, SlideRoleConfig
from lib.profile_schema import ProjectProfile, section_has_content
from lib.shape_discovery import ShapeDiscovery



class PPTXRenderer:
    """Renderer for filling content into existing template slides"""

    # OOXML unit: 914400 EMU = 1 inch.  All other EMU constants derive from this.
    _EMU_PER_INCH: int = 914400

    # Rows are compressed to fit the safe area; this floor keeps text legible.
    _MIN_ROW_HEIGHT_EMU: int = int(0.25 * _EMU_PER_INCH)

    # `_estimate_max_chars` floors at 8 for a degenerately small box. At or
    # barely above that floor, truncating still cuts mid-word — a category
    # pill like "Intake & Registration" becomes the meaningless "Intake &…".
    # Below this threshold, skip truncation entirely and let auto-fit shrink
    # the full text instead (small-but-complete beats legible-but-meaningless).
    _MIN_MEANINGFUL_TRUNCATION_CHARS: int = 12

    # A remainder shorter than this does not earn a slide of its own. Below it,
    # the full text stays in the box and normAutofit shrinks to suit: one
    # trailing sentence is not worth a page, and dropping it is not on offer.
    _MIN_OVERFLOW_APPENDIX_CHARS: int = 150

    # Minimum height for a data row: ONE line of 8pt Noto Sans JP plus the cell's
    # padding, with a hair of slack for what a renderer adds of its own.
    #
    # It was 0.35" — two lines — which charged every row for a second line
    # whether it used one or not. `_estimate_row_content_height_emu` already
    # measures the lines a row really wraps to, so the floor only has to cover
    # the shortest row there can be; at 0.35" a seven-row table whose rows need
    # 1.7 lines on average was paged as six rows and a slide holding one.
    _MIN_DATA_ROW_HEIGHT_EMU: int = int(0.28 * _EMU_PER_INCH)

    # Font metrics for estimating table-cell wrap (8pt Noto Sans JP). Shared by
    # the pre-fill split-size estimate and the post-fill row-height application
    # so both agree on what a row actually needs.
    _TABLE_FONT_PT: int = 8
    _TABLE_CJK_W_INCH: float = _TABLE_FONT_PT / 72          # ~0.111" per CJK char
    _TABLE_LATIN_W_INCH: float = _TABLE_FONT_PT * 0.6 / 72  # ~0.067" per Latin char
    _TABLE_LINE_H_EMU: int = int(_TABLE_FONT_PT * 1.6 / 72 * _EMU_PER_INCH)
    _TABLE_CELL_PADDING_EMU: int = int(0.07 * _EMU_PER_INCH)
    # OOXML's default cell insets are 0.1" left and right, and the text has only
    # what is left between them. Measuring wrap against the FULL column width
    # made every estimate three latin characters too generous: a `Count` header
    # was given a column exactly as wide as the word and came out as `Cou nt`.
    _TABLE_CELL_SIDE_PADDING_EMU: int = int(0.2 * _EMU_PER_INCH)
    _TABLE_CELL_SIDE_PADDING_IN: float = 0.2

    def __init__(self):
        self.template_config = None
        self._output_dir: Optional[Path] = None
        # Safe content area bottom — computed dynamically from slide height in render_from_profile.
        # Fallback: 5" (suits a 5.625" widescreen slide with footer at ~5.25").
        self._SVN_SAFE_BOTTOM_EMU: int = int(5.0 * self._EMU_PER_INCH)
        # Actual slide width — set from prs.slide_width in render_from_profile.
        self._slide_width_emu: int = 0
        # Text-role items too long for their fixed template box: truncated in
        # place, the remainder queued here as
        # {(section, role_name): [(template slide number, remainder), ...]}
        # and rendered as an appendix extra-slide after the main fill pass —
        # reset per render_from_profile call.
        #
        # The slide number rides along because a remainder is only real if the
        # slide that could not hold it SHIPS. Slide 15 is filled like every
        # other configured slide and then pruned (no section claims it), and its
        # truncations used to survive it as an appendix slide of fragments
        # cut from text the deck prints in full elsewhere.
        self._overflow_extras: Dict[Tuple[str, str], List[Tuple[int, str]]] = {}
        # The template slide currently being filled — see `_overflow_extras`.
        self._filling_slide_num: int = 0
        # Set per render from the profile; False keeps every box on the
        # pour-and-cut behavior that predates the condenser.
        self._shape_content: bool = False
        # {slide number: (shape tree, layout, source part)} for the template
        # designs an appendix chapter may borrow — captured before pruning.
        self._donor_library: Dict[int, tuple] = {}
        # {family name: [(shape, text), ...]} for roles that must be SET at one
        # size rather than left to the renderer's per-box auto-fit. Filled as
        # the roles of one slide are written, spent by `_flush_shared_fonts`
        # once they all have been — a family spans roles (slide 16's card
        # carries its headline and its body in two different roles), so the
        # size cannot be decided while either is still unwritten.
        self._font_families: Dict[str, dict] = {}

    def render(
        self,
        slides: List[SlideContent],
        template: str,
        output_name: str,
        output_dir: Optional[str] = None,
    ) -> Path:
        """
        Fill content into existing template slides

        Args:
            slides: List of slide content to fill
            template: Template PPTX file path (e.g., 'SVN Proposal Menu.pptx')
            output_name: Output file name (without .pptx extension)
            output_dir: Deprecated/ignored. Output directory is resolved from env `SLIDE_GENERATOR__OUTPUTS_PATH`.

        Returns:
            Path to generated PPTX file
        """
        print(f'Filling content into template slides: {output_name}')

        # Resolve output_dir once so image path resolution can use it
        self._output_dir = self._resolve_output_dir(output_dir)

        # Auto-detect template config from template name
        self._detect_template_config(template)

        # Load template presentation
        prs = self._load_template(template)

        # Fill content into existing slides
        self._fill_existing_slides(prs, slides)

        # Save presentation
        output_file = self._save_presentation(prs, output_name, output_dir)

        return output_file

    @staticmethod
    def _resolve_output_dir(output_dir: Optional[str]) -> Path:
        """Resolve output directory path.

        Priority: explicit output_dir > env SLIDE_GENERATOR__OUTPUTS_PATH > CWD/outputs
        """
        if output_dir:
            return Path(output_dir)
        env_outputs = os.getenv('SLIDE_GENERATOR__OUTPUTS_PATH')
        if env_outputs:
            return Path(env_outputs).expanduser()
        return Path.cwd() / 'outputs'

    def _detect_template_config(self, template: str):
        """Auto-detect and load template config from template name"""
        if 'svn' in template.lower():
            self.template_config = SVNProposalTemplate()
            print(f'Auto-detected SVN template config: {len(self.template_config.slide_configs)} slides configured')
        else:
            self.template_config = None
            print('No specific template config detected, using default rendering')

    def _load_template(self, template: str) -> Presentation:
        """Load template presentation — supports absolute paths and relative names."""
        template_path = Path(template)
        if template_path.is_absolute() and template_path.exists():
            prs = Presentation(str(template_path))
            print(f'Loaded template: {template_path.name} ({len(prs.slides)} slides)')
            return prs
        rel_path = Path(__file__).parent / 'templates' / template
        if rel_path.exists():
            prs = Presentation(str(rel_path))
            print(f'Loaded template: {template} from {rel_path} ({len(prs.slides)} slides)')
            return prs
        print(f'Template not found: {template} at {template_path} or {rel_path}')
        return Presentation()

    def _save_presentation(self, prs: Presentation, output_name: str, output_dir: Optional[str] = None) -> Path:
        """Save presentation to output directory.

        Priority: explicit output_dir > env SLIDE_GENERATOR__OUTPUTS_PATH > CWD/outputs
        """
        output_path = self._output_dir or self._resolve_output_dir(output_dir)

        try:
            output_path.mkdir(parents=True, exist_ok=True)
        except PermissionError as e:
            # Make error clearer for API layer
            raise PermissionError(f'Permission denied creating outputs directory: {output_path}') from e
        output_file = output_path / f'{output_name}.pptx'

        prs.save(str(output_file))
        print(f'Saved PPTX to: {output_file}')

        return output_file

    def _fill_existing_slides(self, prs: Presentation, slides_content: List[SlideContent]):
        """Fill content into existing slides"""
        print(f'Filling content into {len(slides_content)} existing slides')

        # Sort by slide number so that the offset tracking works correctly when
        # overflow slides are inserted in the middle of the presentation.
        sorted_contents = sorted(slides_content, key=lambda s: s.slide_number)

        # Tracks how many extra slides have been inserted before the current one.
        slide_offset = 0

        for slide_content in sorted_contents:
            slide_num = slide_content.slide_number

            # Check if slide is protected
            if self.template_config and not self.template_config.should_fill_slide(slide_num):
                print(f'Slide {slide_num} is protected - skipping')
                continue

            # Adjusted index accounts for any extra slides inserted before this slide.
            adjusted_index = slide_num - 1 + slide_offset

            # Check if slide exists
            if adjusted_index < len(prs.slides):
                slide = prs.slides[adjusted_index]
                print(f'Filling content into slide {slide_num} (template index: {adjusted_index})')

                # Fill content; if a table overflows it will insert extra slides and
                # return the count so we can update the offset.
                extra = self._fill_with_overflow_handling(prs, slide, slide_content, adjusted_index)
                slide_offset += extra
            else:
                print(
                    f'Slide {slide_num} not found in template '
                    f'(total slides: {len(prs.slides)}, adjusted index: {adjusted_index})',
                )

    # ------------------------------------------------------------------
    # Table-overflow helpers
    # ------------------------------------------------------------------

    def _fill_with_overflow_handling(
        self,
        prs: Presentation,
        slide,
        slide_content: SlideContent,
        slide_index: int,
    ) -> int:
        """Fill a slide, splitting into multiple slides when a table is too tall
        to fit within the safe content area above the footer band.

        Returns the number of *extra* slides inserted (0 when no overflow).
        """
        shape_contents = getattr(slide_content, 'shape_contents', {}) or {}

        # Find the first table-type shape content for this slide
        table_key, table_data = self._find_table_content(slide_content, shape_contents)

        if not table_key:
            # No table present – fill normally
            self._fill_slide_content(slide, slide_content)
            return 0

        # Legacy path: shape position is unknown here, so estimate from safe area.
        # Assume content starts ~1.5" below the top (title band); minimum row
        # height gives the theoretical maximum rows that could ever fit.
        data_area_est = self._SVN_SAFE_BOTTOM_EMU - int(1.5 * self._EMU_PER_INCH)
        legacy_max_rows = max(1, data_area_est // self._MIN_ROW_HEIGHT_EMU)
        chunks = self._split_table_data(table_data, legacy_max_rows)

        if len(chunks) <= 1:
            # Fits on one slide – fill normally
            self._fill_slide_content(slide, slide_content)
            return 0

        print(
            f'  - Slide {slide_content.slide_number}: table split '
            f'({len(chunks)} page(s), up to {legacy_max_rows} rows/page)',
        )

        # Fill the original slide with the first chunk
        first_content = self._build_chunked_slide_content(
            slide_content, table_key, chunks[0], is_first=True,
        )
        self._fill_slide_content(slide, first_content)

        # Create one continuation slide per remaining chunk
        extra_slides = 0
        current_index = slide_index
        for i, chunk in enumerate(chunks[1:], 1):
            new_slide = self._duplicate_slide(prs, current_index)
            extra_slides += 1
            current_index += 1

            cont_content = self._build_chunked_slide_content(
                slide_content, table_key, chunk, is_first=False,
            )
            self._fill_slide_content(new_slide, cont_content)
            print(
                f'  - Added continuation slide {i}/{len(chunks) - 1} '
                f'at template position {current_index + 1}',
            )

        return extra_slides

    def _find_table_content(
        self,
        slide_content: SlideContent,
        shape_contents: dict,
    ):
        """Return (content_key, markdown_table_string) for the first table-type
        shape found in this slide's config.  Returns (None, None) if none found.
        """
        config = (
            self.template_config.get_slide_config(slide_content.slide_number)
            if self.template_config
            else None
        )
        if not config or not config.shape_targets:
            return None, None

        for target in config.shape_targets:
            content = shape_contents.get(target.content_key, '')
            if content and self._is_table_data(content):
                return target.content_key, content

        return None, None

    def _split_table_data(self, table_data: str, max_rows: int,
                          has_header: bool = True) -> List[str]:
        """Split a markdown table string into chunks of at most *max_rows* data rows.

        Each chunk keeps the original header row(s) and the separator row so that
        each chunk is a valid, self-contained markdown table.

        `has_header=False` is the row-only markdown `fill_cols` roles emit (the
        template supplies the header). Without this flag the header/data split
        below keys off finding a `|---|` separator line, and a table that has
        none puts EVERY line in `header_lines` — leaving `data_lines` empty, so
        the length check passes and the table is never split however far it
        overflows. That silently disabled paging on all three 前提条件 slides.

        Returns a list of markdown table strings (length 1 when no split needed).
        """
        lines = [line for line in table_data.strip().split('\n') if line.strip()]
        if not lines:
            return [table_data]

        if not has_header:
            if len(lines) <= max_rows:
                return [table_data]
            per_page = self._rows_per_page(len(lines), max_rows)
            return ['\n'.join(lines[i:i + per_page])
                    for i in range(0, len(lines), per_page)]

        # Separate header / separator lines from data lines
        header_lines: List[str] = []
        data_lines: List[str] = []
        separator_found = False

        for line in lines:
            stripped = line.strip()
            is_separator = (
                not separator_found
                and '|' in stripped
                and all(c in '|-: \t' for c in stripped)
            )
            if is_separator:
                header_lines.append(line)
                separator_found = True
            elif not separator_found:
                header_lines.append(line)
            else:
                data_lines.append(line)

        if len(data_lines) <= max_rows:
            return [table_data]

        header_block = '\n'.join(header_lines)
        per_page = self._rows_per_page(len(data_lines), max_rows)
        chunks: List[str] = []
        for i in range(0, len(data_lines), per_page):
            row_block = '\n'.join(data_lines[i : i + per_page])
            chunks.append(f'{header_block}\n{row_block}')

        return chunks

    @staticmethod
    def _rows_per_page(rows: int, max_rows: int) -> int:
        """Rows to put on each page — the same number on all of them.

        Filling every page to `max_rows` and letting the last take the remainder
        leaves a page holding one row of a table whose other pages hold six. The
        page COUNT is fixed by capacity either way, so spreading the rows evenly
        across that many pages costs nothing and no page ends up nearly empty.
        """
        if rows <= max_rows or max_rows < 1:
            return max(1, max_rows)
        pages = -(-rows // max_rows)
        return -(-rows // pages)

    def _build_chunked_slide_content(
        self,
        original: SlideContent,
        table_key: str,
        table_chunk: str,
        is_first: bool,
    ) -> SlideContent:
        """Return a new SlideContent with *table_key* replaced by *table_chunk*.

        For continuation slides (``is_first=False``) all non-table text shapes
        are cleared so description text is not repeated across slides.
        """
        new_shape_contents = dict(getattr(original, 'shape_contents', {}) or {})
        new_shape_contents[table_key] = table_chunk

        if not is_first:
            config = (
                self.template_config.get_slide_config(original.slide_number)
                if self.template_config
                else None
            )
            if config and config.shape_targets:
                for target in config.shape_targets:
                    key = target.content_key
                    if key != table_key and key in new_shape_contents:
                        if not self._is_table_data(new_shape_contents[key]):
                            new_shape_contents[key] = ''  # Clear description text

        return SlideContent(
            slide_number=original.slide_number,
            layout=original.layout,
            title=original.title,
            subtitle=original.subtitle,
            content=original.content,
            shape_contents=new_shape_contents,
        )

    @staticmethod
    def _ensure_unique_slide_partname(prs: Presentation, slide) -> None:
        """Work around a python-pptx limitation: ``Slides.add_slide()``
        assigns the new slide's partname from ``len(sldIdLst) + 1`` — a
        naive count of the CURRENT slide list — not from the set of
        partnames actually in use. Once slides have been deleted out of
        order (as ``_prune_and_reorder`` does via manual XML surgery,
        working around python-pptx having no public slide-delete API),
        that count drifts below the real max partname index and collides
        with an original template slide that's still alive. python-pptx's
        zip writer doesn't error on this — it silently emits two different
        slide XML parts under the same path, and whichever is written last
        wins for BOTH relationships that point at that name (the earlier
        one's content becomes unreachable). Call right after every
        ``add_slide()`` to reassign a name checked against every part
        actually in the package.
        """
        from pptx.opc.package import PackURI

        package = prs.part.package
        tmpl = '/ppt/slides/slide%d.xml'
        prefix = tmpl[: (tmpl % 42).find('42')]
        others = {p.partname for p in package.iter_parts()
                  if p.partname.startswith(prefix) and p is not slide.part}
        if slide.part.partname in others:
            n = len(others) + 1
            while tmpl % n in others:
                n += 1
            slide.part.partname = PackURI(tmpl % n)

    def _duplicate_slide(self, prs: Presentation, slide_index: int):
        """Duplicate the slide at *slide_index* and insert the copy immediately
        after it in the presentation.

        Returns the new slide object (at position ``slide_index + 1``).
        """
        source_slide = prs.slides[slide_index]

        # 1. Add a new blank slide with the same layout – this is appended at the end.
        new_slide = prs.slides.add_slide(source_slide.slide_layout)
        self._ensure_unique_slide_partname(prs, new_slide)

        # 2. Replace the new slide's shape tree with a deep copy of the source's.
        source_spTree = source_slide.shapes._spTree
        new_spTree = new_slide.shapes._spTree

        for child in list(new_spTree):
            new_spTree.remove(child)
        for child in source_spTree:
            new_child = copy.deepcopy(child)
            self._remap_image_relationships(source_slide.part, new_slide.part, new_child)
            new_spTree.append(new_child)

        # 3. Move the new slide from the end to slide_index + 1 by reordering
        #    the <p:sldIdLst> in the presentation XML.
        pml_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        try:
            presentation_elem = prs.part._element
            sldIdLst = presentation_elem.find(f'{{{pml_ns}}}sldIdLst')
            if sldIdLst is not None:
                sld_id_elems = list(sldIdLst)
                if sld_id_elems:
                    new_sldId = sld_id_elems[-1]          # just appended → at end
                    source_sldId = sld_id_elems[slide_index]  # source position
                    sldIdLst.remove(new_sldId)
                    source_sldId.addnext(new_sldId)
                    print(
                        f'  - Duplicated slide at index {slide_index}, '
                        f'inserted at index {slide_index + 1}',
                    )
        except Exception as e:
            print(f'  - Error repositioning duplicated slide: {e}')

        return prs.slides[slide_index + 1]

    @staticmethod
    def _remove_shape(shape) -> None:
        """Remove a shape entirely (not just its text) from its slide's shape tree."""
        if shape is None:
            return
        elem = shape._element
        parent = elem.getparent()
        if parent is not None:
            parent.remove(elem)

    @staticmethod
    def _remove_decorations_around(slide, ref_shapes: list, padding_emu: int = 228600) -> int:
        """Remove sibling shapes (card background, icon badge, icon picture)
        that visually belong to an unused title/body (or text-item) slot but
        aren't tracked by the role config, so blanking the tracked shapes'
        text alone would leave an empty-looking card/box behind.

        Purely geometric — the role config only knows about title_shape/
        body_shape, not per-slot decoration, and this template's decoration
        shapes aren't at any predictable index/name offset from their title
        (e.g. an icon picture can sit anywhere else in the shape list).
        Any top-level shape whose bounding box is contained within the union
        of `ref_shapes`'s boxes, expanded by `padding_emu` (default 0.25" —
        covers a card background that bleeds slightly past its title+body,
        without reaching a neighboring card across this template's ~0.3"+
        gaps), is removed. `ref_shapes` themselves are excluded.
        """
        boxes = [
            (s.left, s.top, s.left + s.width, s.top + s.height)
            for s in ref_shapes if s is not None and s.left is not None
        ]
        if not boxes:
            return 0
        # Compare by the underlying XML element, not the shape wrapper's own
        # id() — python-pptx returns a fresh wrapper instance each time
        # `.shapes` is iterated/indexed, so two wrappers for the SAME shape
        # never share an id() even though `sh1._element is sh2._element`.
        # Using id(wrapper) here would make every ref_shape match itself as
        # "contained within its own bounding box" and get deleted too.
        exclude_ids = {id(s._element) for s in ref_shapes if s is not None}
        l = min(b[0] for b in boxes) - padding_emu
        t = min(b[1] for b in boxes) - padding_emu
        r = max(b[2] for b in boxes) + padding_emu
        b = max(b[3] for b in boxes) + padding_emu
        removed = 0
        for sh in list(slide.shapes):
            if id(sh._element) in exclude_ids or sh.left is None:
                continue
            sl, st = sh.left, sh.top
            sr, sb = sl + sh.width, st + sh.height
            if sl >= l and st >= t and sr <= r and sb <= b:
                PPTXRenderer._remove_shape(sh)
                removed += 1
        return removed

    @staticmethod
    def _remap_image_relationships(source_part, new_part, element) -> None:
        """Fix up `r:embed` references after deep-copying XML between parts.

        A deep-copied <p:pic>/<a:blip> keeps its source slide's relationship
        ID (e.g. "rId1") verbatim, but that ID is local to the source slide's
        own .rels — the destination slide (freshly created via `add_slide`,
        or the same slide the element is merely being duplicated within)
        doesn't necessarily have a matching relationship at that ID. Left
        unfixed, the reference silently dangles and the picture renders as
        nothing (no error, just a blank image) in PowerPoint/LibreOffice.
        """
        if source_part is new_part:
            return
        from pptx.oxml.ns import qn
        r_embed = qn('r:embed')
        for el in element.iter():
            old_rid = el.get(r_embed)
            if not old_rid:
                continue
            try:
                image_part = source_part.related_part(old_rid)
                reltype = source_part.rels[old_rid].reltype
                new_rid = new_part.relate_to(image_part, reltype)
                el.set(r_embed, new_rid)
            except Exception:
                pass

    @classmethod
    def _clone_shape(cls, slide, source_shape, top: Optional[int] = None):
        """Deep-copy *source_shape* onto *slide* (same slide or another), optionally
        repositioning the copy's top. Returns the new Shape."""
        new_elem = copy.deepcopy(source_shape._element)
        source_part = source_shape.part
        cls._remap_image_relationships(source_part, slide.part, new_elem)
        slide.shapes._spTree.append(new_elem)
        new_shape = slide.shapes[-1]
        if top is not None:
            new_shape.top = top
        return new_shape

    # ------------------------------------------------------------------

    def _fill_slide_content(self, slide, slide_content: SlideContent):
        """Fill content into an existing slide"""
        slide_num = slide_content.slide_number
        config = self.template_config.get_slide_config(slide_num) if self.template_config else None

        # Strategy 1: Shape targeting (if shape_contents exists)
        if self._has_shape_contents(slide_content) and config and config.shape_targets:
            print(f'  - Using shape targeting for slide {slide_num}')
            self._fill_with_shape_targets(slide, slide_content, config)
            return

        # Strategy 2: Default filling
        self._fill_slide_default(slide, slide_content)

    def _has_shape_contents(self, slide_content: SlideContent) -> bool:
        """Check if slide_content has shape_contents attribute"""
        return bool(
            hasattr(slide_content, '__dict__') and
            'shape_contents' in slide_content.__dict__ and
            slide_content.shape_contents,
        )

    def _fill_with_shape_targets(self, slide, slide_content, config):
        """Fill content using shape targets from config"""
        shape_contents = getattr(slide_content, 'shape_contents', {})

        if not shape_contents:
            print('  - No shape_contents found in slide_content')
            return

        print(f'  - Found {len(shape_contents)} shape contents to fill')

        # Update title if exists
        self._update_slide_title(slide, slide_content.title)

        # Fill each shape target
        for target in config.shape_targets:
            self._fill_single_shape_target(slide, target, shape_contents, config)

    def _update_slide_title(self, slide, title: Optional[str]):
        """Update slide title if exists"""
        if slide.shapes.title and title:
            slide.shapes.title.text = title
            print(f'  - Updated title: {title}')

    def _fill_single_shape_target(self, slide, target, shape_contents: Dict[str, str], config=None):
        """Fill a single shape target with content"""
        content_key = target.content_key

        # Check if content exists for this key
        if content_key not in shape_contents:
            print(f"  - Content key '{content_key}' not found in markdown")
            return

        content = shape_contents[content_key]

        # Find shape by index or name
        shape = self._find_shape(slide, target)

        if not shape:
            print(f'  - Shape not found for target: {content_key}')
            return

        # Detect content type and fill accordingly
        if self._is_image_path(content):
            # Content is an image path
            self._replace_shape_with_image(slide, shape, content, target)
        elif target.fill_cols or self._is_table_data(content):
            # fill_cols always implies a table target (even single-row);
            # otherwise fall back to markdown table detection.
            self._fill_table(shape, content, target)
        elif shape.has_text_frame:
            # Content is text
            shape.text_frame.text = content
            self._apply_text_font(shape.text_frame)
            
            # Apply 10pt line spacing for slide 4
            if config and config.slide_number == 4:
                self._apply_line_spacing(shape.text_frame, Pt(10))
                print(f"  - Applied 10pt line spacing to slide 4 shape '{content_key}'")
            
            print(f"  Filled '{content_key}' into {shape.name} ({len(content)} chars)")
        else:
            print(f'  - Shape {shape.name} has no text frame and content is not an image or table')

    def _is_image_path(self, content: str) -> bool:
        """Check if content is an image file path or base64 data URI"""
        content = content.strip()
        # Check base64 data URI: data:image/...;base64,...
        if content.startswith('data:image/'):
            return True
        # Check if it's a file path with image extension
        image_extensions = ['.png', '.jpg', '.jpeg', '.gif', '.bmp', '.svg']
        return any(content.lower().endswith(ext) for ext in image_extensions)

    def _decode_base64_image(self, data_uri: str) -> Optional[str]:
        """Decode a base64 data URI to a temporary file and return the file path.

        Supports format: data:image/<ext>;base64,<data>

        Returns:
            Absolute path to the temporary file, or None on error.
        """
        try:
            # Parse header: data:image/png;base64,<data>
            header, encoded = data_uri.split(',', 1)
            # Extract mime type, e.g. image/png
            mime_part = header.split(':')[1].split(';')[0]  # e.g. "image/png"
            ext = '.' + mime_part.split('/')[1]  # e.g. ".png"
            # Handle special cases
            if ext == '.jpeg':
                ext = '.jpg'
            elif ext == '.svg+xml':
                ext = '.svg'

            image_bytes = base64.b64decode(encoded)
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
            tmp.write(image_bytes)
            tmp.flush()
            tmp.close()
            print(f'  - Decoded base64 image to temp file: {tmp.name} ({len(image_bytes)} bytes)')
            return tmp.name
        except Exception as e:
            print(f'  - Failed to decode base64 image: {e}')
            return None

    def _replace_shape_with_image(self, slide, old_shape, image_path: str, target):
        """Replace a shape with an image (supports file path or base64 data URI)"""
        tmp_file: Optional[str] = None
        try:
            # Get shape position and size
            left = old_shape.left
            top = old_shape.top
            width = old_shape.width
            height = old_shape.height

            image_path = image_path.strip()

            # Handle base64 data URI
            if image_path.startswith('data:image/'):
                tmp_file = self._decode_base64_image(image_path)
                if not tmp_file:
                    print(f"  - Cannot decode base64 image for shape '{target.content_key}'")
                    return
                image_path = tmp_file
            else:
                # Resolve file path
                if not Path(image_path).is_absolute():
                    rel_path = Path(image_path)
                    if rel_path.exists():
                        image_path = str(rel_path)
                    elif self._output_dir:
                        resolved = self._output_dir / image_path
                        if resolved.exists():
                            image_path = str(resolved)
                    else:
                        outputs_path = Path('outputs') / image_path
                        if outputs_path.exists():
                            image_path = str(outputs_path)

                # Check if image exists
                if not Path(image_path).exists():
                    print(f'  - Image not found: {image_path}')
                    return

            # Add new image at same position (overlays the placeholder shape)
            try:
                slide.shapes.add_picture(
                    str(image_path),
                    left, top,
                    width=width,
                    height=height,
                )
                print(f"  Replaced shape '{target.content_key}' with image: {Path(image_path).name}")
            except Exception as e:
                print(f'  - Error adding image: {e}')

        except Exception as e:
            print(f'  - Error replacing shape with image: {e}')
        finally:
            # Clean up temp file if created from base64
            if tmp_file and Path(tmp_file).exists():
                try:
                    Path(tmp_file).unlink()
                    print(f'  - Cleaned up temp file: {tmp_file}')
                except Exception:
                    pass

    def _is_table_data(self, content: str) -> bool:
        """Check if content is table data (markdown table format).

        Accepts two formats:
        1. Standard markdown table with header + separator:
               | col1 | col2 |
               |------|------|
               | val1 | val2 |
        2. No-header row-only table (used for fill_cols targeting):
               | val1 |
               | val2 |
               | val3 |
        """
        content = content.strip()
        lines = [line.strip() for line in content.split('\n') if line.strip()]

        if len(lines) < 2:
            return False

        # All lines must contain at least one pipe
        if not all('|' in line for line in lines):
            return False

        # Format 1: standard table with separator line
        for line in lines[1:3]:
            if all(c in '|-: \t' for c in line):
                return True

        # Format 2: no-separator — every line starts and ends with | (row-only table)
        if all(line.startswith('|') and line.endswith('|') for line in lines):
            return True

        return False

    def _fill_table(self, shape, table_data: str, target):
        """Fill an existing table shape with data from markdown table format"""
        try:
            # Check if shape has table
            if not shape.has_table:
                print(f"  - Shape '{target.content_key}' is not a table")
                return

            table = shape.table
            rows_data = self._parse_md_table_rows(table_data)

            if not rows_data:
                print('  - No valid table data found')
                return

            # Get current table dimensions
            current_rows = len(table.rows)
            current_cols = len(table.columns)
            needed_rows = len(rows_data)
            needed_cols = len(rows_data[0]) if rows_data else 0

            # Adjust table size if needed
            if needed_rows > current_rows:
                # Add rows if we need more
                rows_to_add = needed_rows - current_rows
                print(f'  - Adding {rows_to_add} rows to table (current: {current_rows}, needed: {needed_rows})')
                self._add_table_rows(table, rows_to_add)
                # After adding rows, update current_rows to reflect the new count.
                # Note: python-pptx's table.rows collection is live, so this should
                # reflect the new size immediately, but we recalculate for safety.
                current_rows = len(table.rows)
            elif needed_rows < current_rows:
                # Fewer rows needed - remove the extra rows from the XML so they
                # don't appear as empty rows in the rendered slide.
                extra_rows = current_rows - needed_rows
                print(
                    '  - Table has %s rows but only %s rows of data '
                    '(removing %s unused rows)',
                    current_rows,
                    needed_rows,
                    extra_rows,
                )
                self._remove_table_rows(table, needed_rows)
                current_rows = len(table.rows)

            # Reshape the grid to the data first, so the fill below neither
            # leaves empty columns eating width the real text needs nor drops
            # data the markdown supplied. A no-op when they already agree, and
            # for `fill_cols` roles by design.
            self._fit_table_cols_to_data(table, rows_data, target.fill_cols)
            current_cols = len(table.columns)

            if needed_cols > current_cols:
                # Only reachable when the grid could not be reshaped (merged
                # cells) or under fill_cols — otherwise data is lost silently.
                print(f'  - WARNING: data has {needed_cols} columns but the table '
                      f'still has {current_cols}; dropping '
                      f'{needed_cols - current_cols} column(s) of data')
                needed_cols = current_cols  # Limit to available columns
            elif needed_cols < current_cols and not target.fill_cols:
                print(f'  - Data has {needed_cols} columns but table has {current_cols} columns '
                      f'(clearing {current_cols - needed_cols} unmatched column(s) instead of '
                      f'leaving stale template placeholder text)')

            # Fill table cells
            filled_count = 0
            for row_idx, row_data in enumerate(rows_data):
                # After adding rows via XML, we can fill up to needed_rows
                # Use try-except to handle any edge cases where table structure might not match
                try:
                    # If fill_cols is specified, map row_data positionally onto those columns;
                    # columns not listed are left untouched (preserving template content).
                    if target.fill_cols:
                        col_pairs = list(zip(target.fill_cols, row_data))
                    else:
                        col_pairs = list(enumerate(row_data))

                    filled_col_idxs = set()
                    for col_idx, cell_value in col_pairs:
                        if col_idx >= current_cols:
                            break  # Skip extra columns in data

                        cell = table.cell(row_idx, col_idx)
                        # Preserving the cell's own run, not `cell.text = ...`:
                        # that setter styles the text from scratch, and this
                        # template styles cells per run — the assumptions grid's
                        # label column is white bold on gold, and written the
                        # blunt way it came out as plain body text on the same
                        # gold, reading like a cell nobody had finished.
                        #
                        # Only where the cell speaks with ONE voice, though. A
                        # cell whose sample text picked out a phrase in red bold
                        # (slide 25's does) would lend that emphasis to whatever
                        # replaced it, and a whole assumption came out shouting.
                        if self._cell_style_is_uniform(cell.text_frame):
                            self._set_text_preserving_format(cell.text_frame, cell_value)
                        else:
                            cell.text = cell_value
                        self._apply_table_cell_font(cell)
                        filled_count += 1
                        filled_col_idxs.add(col_idx)

                    # Without fill_cols, the data is expected to cover every
                    # column; any column the data doesn't reach gets cleared
                    # so stale template placeholder text can't leak through.
                    # fill_cols mode is exempt — it deliberately targets a
                    # subset of columns and leaves the rest (e.g. a fixed
                    # label column) untouched by design.
                    if not target.fill_cols:
                        for col_idx in range(current_cols):
                            if col_idx not in filled_col_idxs:
                                table.cell(row_idx, col_idx).text = ''
                except (IndexError, AttributeError) as e:
                    print(f'  - Cannot fill row {row_idx}: {e}')
                    break  # Stop filling if we hit an error

            print(f"  Filled '{target.content_key}' table: {len(rows_data)} rows x {needed_cols} cols ({filled_count} cells)")
            if target.fill_cols:
                self._normalize_table_col_widths(table)
            self._apply_content_row_heights(table)
            self._fit_table_to_bounds(shape)

        except Exception as e:
            print(f'  - Error filling table: {e}')

    def _add_table_rows(self, table, num_rows: int):
        """Add rows to an existing table by manipulating XML

        python-pptx uses lxml internally, so we can manipulate the XML directly
        to add new rows by cloning the structure of the last row.
        """
        try:
            from lxml import etree

            # Get the table XML element (this is an lxml element)
            tbl = table._tbl

            # Get the number of existing rows
            current_row_count = len(table.rows)

            # Get the last row to use as template for new rows
            if current_row_count == 0:
                print('  - Cannot add rows: table has no existing rows to use as template')
                return

            # Use positive index instead of negative index (table.rows doesn't support negative indexing)
            last_row = table.rows[current_row_count - 1]
            last_row_xml = last_row._tr

            # Define namespace for XPath queries
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

            # Add new rows by cloning the last row
            for i in range(num_rows):
                # Create a deep copy of the last row by serializing and parsing
                new_row_xml = etree.fromstring(etree.tostring(last_row_xml))

                # Clear text content in all cells of the new row
                for tc in new_row_xml.xpath('.//a:tc', namespaces=ns):
                    txBody = tc.find('.//a:txBody', namespaces=ns)
                    if txBody is not None:
                        # Get all paragraphs
                        paragraphs = txBody.findall('.//a:p', namespaces=ns)
                        if paragraphs:
                            # Clear text in first paragraph
                            first_p = paragraphs[0]
                            for t_elem in first_p.xpath('.//a:t', namespaces=ns):
                                t_elem.text = ''
                            # Remove other paragraphs (keep structure but remove extra content)
                            for p in paragraphs[1:]:
                                txBody.remove(p)

                # Append the new row to the table
                tbl.append(new_row_xml)

            print(f'  - Successfully added {num_rows} rows to table')

        except ImportError:
            print('  - lxml not available. python-pptx requires lxml, please install it: pip install lxml')
        except Exception as e:
            print(f'  - Error adding rows to table: {e}')
            import traceback
            print(f'  - Traceback: {traceback.format_exc()}')

    def _remove_table_rows(self, table, keep_rows: int):
        """Remove all rows after index *keep_rows* from the table by manipulating XML.

        Args:
            table: python-pptx Table object
            keep_rows: Number of rows to keep (rows at index 0 … keep_rows-1 are retained).
        """
        try:
            from lxml import etree  # noqa: F401 (just to confirm availability)

            tbl = table._tbl
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            tr_elements = tbl.findall('a:tr', namespaces=ns)

            rows_to_remove = tr_elements[keep_rows:]
            for tr in rows_to_remove:
                tbl.remove(tr)

            print(
                f'  - Removed {len(rows_to_remove)} extra row(s); '
                f'table now has {keep_rows} row(s)',
            )
        except Exception as e:
            print(f'  - Error removing table rows: {e}')

    # ------------------------------------------------------------------
    # Column reshaping — make the template grid follow the source markdown
    # ------------------------------------------------------------------

    def _fit_table_cols_to_data(self, table, rows_data: List[List[str]],
                                fill_cols=None) -> None:
        """Reshape the template's table grid to the markdown table's column count.

        The template ships each table with a fixed width and a fixed number of
        columns; the markdown a proposal is written from rarely has exactly that
        many. Left alone both mismatches damage the deck:

          - Fewer data columns than the grid — a 2-column table filled into the
            6-column 機能一覧 grid kept four empty columns that still took two
            thirds of the width, so the real text wrapped inside a ~1.1" ribbon,
            rows grew tall, and five rows split across three slides.
          - More data columns than the grid — the surplus was dropped silently.

        So the grid follows the data: surplus columns are deleted, missing ones
        appended, and the width is then shared out by what each surviving column
        actually holds. The table's position and total width never change.

        No-op when the data already matches the grid, and when `fill_cols` maps
        the data onto a deliberate subset of columns (the assumptions slides,
        where column 0 is a fixed template label the data never supplies).
        """
        if fill_cols or not rows_data:
            return
        needed = max(len(row) for row in rows_data)
        current = len(table.columns)
        if needed < 1:
            return
        if needed == current:
            return
        if self._has_merged_cells(table):
            print(f'  - Table has merged cells — keeping its {current} template '
                  f"column(s) rather than reshaping to the data's {needed}")
            return
        # Read the table's full width BEFORE the reshape: deleting a column
        # takes its width out of the grid and cloning one adds width on top, so
        # measuring afterwards would shrink or inflate the table instead of
        # redistributing what the template authored.
        total_width = sum(col.width or 0 for col in table.columns)
        if needed < current:
            self._remove_table_cols(table, needed)
        else:
            self._add_table_cols(table, needed - current)
        self._spread_col_widths(table, rows_data, total_width)

    def _unmerge_table_cells(self, table) -> bool:
        """Flatten every merged cell in the table back into ordinary cells.

        Some template tables group rows under one spanning label — slide 39's
        想定成果物 matrix merges its 対応タスク column down eight rows. Writing a
        flat list of rows into such a grid silently loses text: a cell hidden by
        a merge renders nothing, so most of the column came out blank.

        The merge is decoration on sample content, not structure the profile
        knows about, so it is dropped rather than worked around. Returns True
        when anything was unmerged.
        """
        from pptx.oxml.ns import qn

        rows = table._tbl.findall(qn('a:tr'))
        removed = 0
        for tr in rows:
            for cell in tr.findall(qn('a:tc')):
                for attr in ('gridSpan', 'hMerge', 'rowSpan', 'vMerge'):
                    if cell.get(attr) is not None:
                        del cell.attrib[attr]
                        removed += 1
        if not removed:
            return False
        print(f'  - Unmerged {removed} merged-cell attribute(s) so every row '
              f'can be written')
        self._level_column_cell_style(rows)
        return True

    @staticmethod
    def _level_column_cell_style(rows: list) -> None:
        """Give every data cell in a column the fill AND borders of the first one.

        A grouped template table decorates its spanning label cells to mark the
        groups: a tint to separate them, and no inner horizontal rule between the
        rows a label spans. Once the merges are gone, both land on individual
        rows — the column reads as banded at random and loses the lines between
        the rows that used to share a label.

        A flat table wants a flat column, so the first data row sets both for the
        rest. The header row is left alone: its own styling is the table's
        design, not group decoration.
        """
        from pptx.oxml.ns import qn

        if len(rows) < 3:
            return
        # Fill and the four cell borders, in the order CT_TableCellProperties
        # requires them — a:tcPr is a sequence, so appending out of order gives
        # a file PowerPoint rejects.
        styled_tags = [qn('a:lnL'), qn('a:lnR'), qn('a:lnT'), qn('a:lnB'),
                       qn('a:noFill'), qn('a:solidFill'), qn('a:gradFill'),
                       qn('a:blipFill'), qn('a:pattFill'), qn('a:grpFill')]
        order = {tag: i for i, tag in enumerate(styled_tags)}

        def styled_children(cell):
            props = cell.find(qn('a:tcPr'))
            if props is None:
                return []
            return [child for child in props if child.tag in order]

        first_data = rows[1].findall(qn('a:tc'))
        for col_idx, source in enumerate(first_data):
            template_style = styled_children(source)
            if not template_style:
                continue
            for tr in rows[2:]:
                cells = tr.findall(qn('a:tc'))
                if col_idx >= len(cells):
                    continue
                props = cells[col_idx].find(qn('a:tcPr'))
                if props is None:
                    continue
                for existing in styled_children(cells[col_idx]):
                    props.remove(existing)
                # Inserted at the front in reverse order, so they end up in
                # schema order ahead of whatever else `a:tcPr` holds (`a:headers`,
                # `a:extLst`), which must stay last.
                for child in sorted(template_style, key=lambda c: -order[c.tag]):
                    props.insert(0, copy.deepcopy(child))

    @staticmethod
    def _has_merged_cells(table) -> bool:
        """True when any cell spans columns.

        Such a grid cannot be reshaped by adding or dropping whole columns
        without corrupting the merge, so reshaping bails out instead.
        """
        from pptx.oxml.ns import qn

        for tr in table._tbl.findall(qn('a:tr')):
            for tc in tr.findall(qn('a:tc')):
                if tc.get('gridSpan') or tc.get('hMerge'):
                    return True
        return False

    def _remove_table_cols(self, table, keep_cols: int) -> None:
        """Delete every column after index *keep_cols* — grid entry and cells.

        Mirrors `_remove_table_rows`. A PPTX table is a `tblGrid` of `gridCol`
        widths plus one `tc` per column in every `tr`, so both have to go or the
        file is malformed. The freed width is handed back by `_spread_col_widths`.
        """
        try:
            from pptx.oxml.ns import qn

            tbl = table._tbl
            grid = tbl.find(qn('a:tblGrid'))
            removed = 0
            for gridcol in grid.findall(qn('a:gridCol'))[keep_cols:]:
                grid.remove(gridcol)
                removed += 1
            for tr in tbl.findall(qn('a:tr')):
                for cell in tr.findall(qn('a:tc'))[keep_cols:]:
                    tr.remove(cell)
            print(f'  - Removed {removed} unused column(s); '
                  f'table now has {keep_cols} column(s)')
        except Exception as e:
            print(f'  - Error removing table columns: {e}')

    def _add_table_cols(self, table, num_cols: int) -> None:
        """Append *num_cols* columns by cloning the last one.

        Cloning carries the template's cell fill, borders and margins onto the
        new columns — the same reason `_add_table_rows` clones the last row.
        New elements are inserted next to the last existing one rather than
        appended to the parent, so they stay ahead of any trailing `extLst`.
        """
        try:
            from pptx.oxml.ns import qn

            tbl = table._tbl
            grid = tbl.find(qn('a:tblGrid'))
            for _ in range(num_cols):
                last_gridcol = grid.findall(qn('a:gridCol'))[-1]
                last_gridcol.addnext(copy.deepcopy(last_gridcol))
                for tr in tbl.findall(qn('a:tr')):
                    last_cell = tr.findall(qn('a:tc'))[-1]
                    last_cell.addnext(copy.deepcopy(last_cell))
            print(f'  - Added {num_cols} column(s); '
                  f'table now has {len(table.columns)} column(s)')
        except Exception as e:
            print(f'  - Error adding table columns: {e}')

    def _spread_col_widths(self, table, rows_data: List[List[str]],
                           total: int) -> None:
        """Share *total* EMU of table width out by how much text each column holds.

        Called for every data-driven grid, whether or not the column count
        changed: the template's own proportions describe the sample text it
        shipped with, not this proposal's. *total* is the table's width as the
        template authored it, so the grid ends up exactly as wide as the one it
        replaced.

        Two figures per column, and the difference between them is the whole
        method:

          * its FLOOR is the widest single word it holds. Narrower than that and
            the renderer breaks words mid-syllable, which no amount of extra
            height repairs;
          * its DEMAND is the widest whole cell, i.e. the width at which nothing
            in that column wraps at all.

        Floors are paid first, and what is left over is shared in proportion to
        what each column still wants. A flat one-inch floor — the rule this
        replaces — spent 2" of a 8.06" grid on `#` and `Count`, two columns
        holding at most two digits, and took it from the two columns holding
        sentences: `Authentication & Onboarding` then wrapped to two lines in a
        1.25" ribbon and was truncated to `Authentication & Onboa…`.
        """
        cols = list(table.columns)
        if not total or not cols:
            return
        floors, demands = [], []
        for j in range(len(cols)):
            texts = [row[j] for row in rows_data if j < len(row) and row[j]]
            words = [word for text in texts for word in text.split()]
            floor = max((self._text_width_in(word, self._TABLE_FONT_PT)
                         for word in words), default=0.0)
            demand = max((self._text_width_in(text, self._TABLE_FONT_PT)
                          for text in texts), default=0.0)
            floors.append(min(max(floor, self._MIN_COL_WIDTH_IN), self._MAX_COL_FLOOR_IN))
            demands.append(max(demand, floors[-1]) + self._TABLE_CELL_SIDE_PADDING_IN)
        padding = self._TABLE_CELL_SIDE_PADDING_EMU
        floor_emu = [int(width * self._EMU_PER_INCH) + padding for width in floors]
        # Floors alone over-subscribe the grid only when the words themselves are
        # too wide for it; then every column is scaled down together and the
        # renderer breaks what it must.
        if sum(floor_emu) >= total:
            widths = [int(width * total / sum(floor_emu)) for width in floor_emu]
        else:
            spare = total - sum(floor_emu)
            wants = [max(0.0, demand - floor - self._TABLE_CELL_SIDE_PADDING_IN)
                     for demand, floor in zip(demands, floors)]
            share = sum(wants)
            widths = [
                base + (int(want * spare / share) if share else spare // len(cols))
                for base, want in zip(floor_emu, wants)
            ]
        widths[-1] += total - sum(widths)  # absorb rounding drift, keep total exact
        for col, width in zip(cols, widths):
            col.width = width
        print('  - Column widths by content: '
              f'{[round(w / self._EMU_PER_INCH, 2) for w in widths]}"')

    # A column narrower than this is unreadable whatever it holds; wider than
    # the floor cap, and one long word would starve every other column.
    _MIN_COL_WIDTH_IN: float = 0.3
    _MAX_COL_FLOOR_IN: float = 1.6

    # ------------------------------------------------------------------
    # Cell trimming — keep a table inside the slide the template designed
    # ------------------------------------------------------------------

    def _trim_table_cells_to_fit(self, shape, rows_data: List[List[str]],
                                 role: ShapeRole, has_header: bool,
                                 rows_per_slide: Optional[int] = None) -> bool:
        """Trim over-long cells so the whole table fits above the footer.

        The bespoke template slides size each table for a *summary*: the
        前提条件 cell holds roughly 200 characters. A chapter copied verbatim
        out of a proposal runs 500-1200, and the renderer's own measurement
        says so — `_fit_table_to_bounds` reports a content floor of 8.89" for a
        4.14" gap, then leaves the table overflowing the footer because
        compressing further would clip text mid-sentence.

        Splitting instead would be honest but useless here: one assumption per
        slide turns three 前提条件 slides into eight. So the table keeps the
        shape the template designed — as many rows as the source has — and each
        cell keeps as much text as its own share of the height allows, cut at a
        sentence boundary.

        The trimmed remainder is NOT queued for an appendix slide, unlike an
        over-long text box. These slides are summaries by design: what falls off
        the end is detail the deck deliberately leaves in the source document,
        not content in need of rescuing. Sending it to an appendix produced a
        slide holding six 700-1200 character assumptions as bullets — itself
        overflowing, and duplicating what the three summary slides had just
        said. How much was dropped is reported instead.

        Returns True when any cell was trimmed (the caller then rebuilds the
        markdown from `rows_data`, which this mutates in place).
        """
        if not shape.has_table or not rows_data:
            return False
        table = shape.table
        data_rows = rows_data[1:] if has_header else rows_data
        available = max(0, self._SVN_SAFE_BOTTOM_EMU - shape.top)
        if not data_rows or not available:
            return False

        # Reserve the first template row as a header only when the data HAS a
        # header. A `fill_cols` role writes row-only markdown into a grid whose
        # first row is data — the assumptions slides are exactly that — and
        # charging it to a header that does not exist cost the row its share of
        # the height: the one-row slide came out with a single line's budget and
        # dropped 445 characters of a 500-character assumption. `
        # _max_rows_for_shape` reserves it on the same condition, so the two
        # still agree on the space available.
        header_height = (table.rows[0].height or 0) if has_header and len(table.rows) else 0
        # Every row gets an equal share of what the header leaves behind. Equal
        # rather than proportional on purpose: a table whose rows are visibly
        # the same height reads as a summary, which is what these slides are.
        col_widths = [max(col.width / self._EMU_PER_INCH, 0.5) for col in table.columns]
        grid_of = self._grid_positions(role.fill_cols, len(col_widths))

        # Subtract the padding `_estimate_row_content_height_emu` adds back when
        # the split pass re-measures these same rows. Without this the two
        # disagree by one line per row and the table is paged in two anyway —
        # defeating the trim, which exists precisely to keep it on one slide.
        # A grid that may be PAGED charges each row against a whole slide, not
        # against a share of one: the split pass will give it as many slides as
        # its rows need, so cutting text to make room here throws away content to
        # save a page nobody asked to save. Only a grid pinned to one slide —
        # `fill_cols`, the 前提条件 summaries — divides the height between its
        # rows.
        if not role.fill_cols:
            rows_per_slide = 1
        on_slide = max(1, min(rows_per_slide or len(data_rows), len(data_rows)))
        budget = int((available - header_height
                      - on_slide * self._TABLE_CELL_PADDING_EMU)
                     / self._TABLE_LINE_H_EMU)
        needs = [self._row_line_need(row, grid_of, col_widths) for row in data_rows]
        # A table that will be paged gets its budget shared among the rows of ONE
        # page, then that allowance is applied to every row: each page then holds
        # rows of the size its own height allows.
        page_budget = max(on_slide, budget)
        allowance = (self._share_lines(needs, page_budget) if on_slide == len(data_rows)
                     else self._share_lines_per_page(needs, page_budget, on_slide))

        trimmed = 0
        dropped = 0
        for row, lines_allowed in zip(data_rows, allowance):
            for pos, text in enumerate(row):
                col_idx = grid_of[pos] if pos < len(grid_of) else pos
                if col_idx >= len(col_widths) or not text:
                    continue
                max_chars = max(self._MIN_MEANINGFUL_TRUNCATION_CHARS,
                                self._cell_chars_per_line(col_widths[col_idx],
                                                          sample=text)
                                * lines_allowed)
                if len(text) <= max_chars:
                    continue
                row[pos] = self._truncate_at_sentence(text, max_chars)
                dropped += len(text) - len(row[pos])
                trimmed += 1
        if trimmed:
            print(f'  - Trimmed {trimmed} over-long cell(s) ({dropped} chars dropped) so the '
                  f'table fits its slide; rows given {allowance} line(s) of the '
                  f'{budget} available')
        return bool(trimmed)

    def _cell_chars_per_line(self, col_width_inch: float, sample: str = '') -> int:
        """Characters of *sample*'s own script that fit one line of this column.

        Measured in the script actually being set, not in the widest one there
        is. Charging latin text the CJK width — 0.111" a character against
        0.067" — cut a column's capacity by 40%: `Authentication & Onboarding`
        was judged too long for a 1.25" column that holds it in two lines, and
        was truncated to `Authentication & Onboa…` instead.
        """
        per_char = (self._text_width_in(sample, self._TABLE_FONT_PT) / len(sample)
                    if sample else self._TABLE_CJK_W_INCH)
        usable = max(0.1, col_width_inch - self._TABLE_CELL_SIDE_PADDING_IN)
        return max(1, int(usable / max(per_char, 0.001)))

    def _row_line_need(self, row: List[str], grid_of: List[int],
                       col_widths: List[float]) -> int:
        """Lines this row wants — its tallest cell decides, as rendering does."""
        need = 1
        for pos, text in enumerate(row):
            col_idx = grid_of[pos] if pos < len(grid_of) else pos
            if col_idx >= len(col_widths) or not text:
                continue
            per_line = self._cell_chars_per_line(col_widths[col_idx], sample=text)
            need = max(need, -(-len(text) // per_line))
        return need

    def _share_lines_per_page(self, needs: List[int], page_budget: int,
                              rows_per_page: int) -> List[int]:
        """Share a PAGE's lines among rows, page by page, for a table that splits.

        Rows are taken in the order they will be laid out, `rows_per_page` at a
        time, and each group divides one page's budget between its own members —
        so a long row on a sparse page keeps its lines instead of being cut to
        the average of a table it never shares a slide with.
        """
        allowance: List[int] = []
        for start in range(0, len(needs), rows_per_page):
            group = needs[start:start + rows_per_page]
            allowance.extend(self._share_lines(group, max(len(group), page_budget)))
        return allowance

    @staticmethod
    def _share_lines(needs: List[int], budget: int) -> List[int]:
        """Divide *budget* lines among rows wanting *needs*, and waste none.

        An equal share per row was the old rule, and on a real chapter it read
        as a mistake: five assumptions of wildly different lengths were each cut
        to four lines, so two rows ended mid-sentence in `…` while the three
        below them sat half empty. A row that needs two lines cannot use four,
        and the two it does not need belong to the row that ran out.

        So: fill every row to an equal share, hand back what the short rows
        could not use, and repeat until either every row has what it needs or
        the budget is spent. Nobody goes below one line — a row rendered as
        nothing at all is worse than a row cut short.
        """
        allowance = [0] * len(needs)
        remaining, open_rows = budget, [i for i in range(len(needs))]
        while open_rows and remaining >= len(open_rows):
            share = remaining // len(open_rows)
            for i in list(open_rows):
                take = min(share, needs[i] - allowance[i])
                allowance[i] += take
                remaining -= take
                if allowance[i] >= needs[i]:
                    open_rows.remove(i)
            if not share:
                break
        # Crumbs — fewer lines left than rows still wanting them: the row
        # furthest from what it needs takes each one.
        for i in sorted(open_rows, key=lambda i: needs[i] - allowance[i], reverse=True):
            if remaining <= 0:
                break
            allowance[i] += 1
            remaining -= 1
        return [max(1, lines) for lines in allowance]

    @staticmethod
    def _grid_positions(fill_cols, n_cols: int) -> List[int]:
        """Map each position in a data row to its column index in the grid.

        A `fill_cols` role's row holds only the columns it writes, so position
        in the row is not position in the grid: the assumptions row's single
        cell is column 1 (the wide body), not column 0 (the narrow label).
        Anything measuring a cell against the wrong column gets its width wrong
        by a factor of five — too narrow, so it over-estimates wrapping and
        either over-trims the text or decides one row per slide.
        """
        return list(fill_cols) if fill_cols else list(range(n_cols))

    @classmethod
    def _row_to_grid(cls, cells: List[str], fill_cols, n_cols: int) -> List[str]:
        """Spread a data row across the full grid, blank in columns it skips."""
        if not fill_cols:
            return cells
        grid = [''] * n_cols
        for col_idx, value in zip(fill_cols, cells):
            if 0 <= col_idx < n_cols:
                grid[col_idx] = value
        return grid

    @staticmethod
    def _md_table_from_rows(rows_data: List[List[str]], has_header: bool) -> str:
        """Rebuild a markdown pipe table from parsed rows.

        Mirrors `_rows_to_md_table`'s two shapes: header + separator + data, or
        the row-only form `fill_cols` roles use.
        """
        lines = ['| ' + ' | '.join(cell for cell in row) + ' |' for row in rows_data]
        if has_header and lines:
            lines.insert(1, '|' + '|'.join(['---'] * len(rows_data[0])) + '|')
        return '\n'.join(lines)

    @staticmethod
    def _is_cjk_char(ch: str) -> bool:
        return '　' <= ch <= '鿿' or '＀' <= ch <= '￯'

    def _estimate_row_content_height_emu(
        self, cells: List[str], col_widths_inch: List[float],
    ) -> int:
        """Estimate the EMU height *cells* need to render without further wrap,
        given each column's width. Shared by the pre-fill split-size estimate
        (_max_rows_for_shape) and the post-fill row-height pass
        (_apply_content_row_heights) so both agree on what a row actually needs.
        """
        max_lines = 1
        for j, text in enumerate(cells):
            if j >= len(col_widths_inch) or not text:
                continue
            col_w = max(0.1, col_widths_inch[j] - self._TABLE_CELL_SIDE_PADDING_IN)
            text_width = sum(
                self._TABLE_CJK_W_INCH if self._is_cjk_char(ch) else self._TABLE_LATIN_W_INCH
                for ch in text
            )
            lines = max(1, int(text_width / col_w) + 1)
            max_lines = max(max_lines, lines)
        return int(max_lines * self._TABLE_LINE_H_EMU + self._TABLE_CELL_PADDING_EMU)

    @staticmethod
    def _parse_md_table_rows(table_data: str) -> List[List[str]]:
        """Parse a markdown pipe-table string into row cell-lists (header first).

        Drops separator lines (e.g. ``|---|---|``). Shared by `_fill_table` (which
        writes the cells) and `_max_rows_for_shape` (which estimates how many rows
        of this exact data will fit) so both operate on the same parsed rows.

        An EMPTY cell is a cell. Only the pipes at the two ends of the line are
        punctuation; every `||` between them is a column the row says nothing in,
        and dropping those shifted every later value one column to the left. A
        source row of `| | Total | | 36 |` came out as `Total` under `#` and `36`
        under `Group` — a total of thirty-six groups, which the document never
        said.
        """
        lines = [line.strip() for line in table_data.strip().split('\n') if line.strip()]
        rows_data = []
        for line in lines:
            if all(c in '|-: \t' for c in line):
                continue  # separator line
            cells = line.split('|')
            if cells and not cells[0].strip():
                cells = cells[1:]
            if cells and not cells[-1].strip():
                cells = cells[:-1]
            cells = [cell.strip() for cell in cells]
            if any(cells):
                rows_data.append(cells)
        return rows_data

    def _apply_content_row_heights(self, table) -> None:
        """Bump data-row heights so each row is tall enough for its actual text content.

        The template may have very compact rows (e.g. 0.262") designed for short
        placeholder text.  When real content is longer and word-wraps, some renderers
        (LibreOffice, Google Slides preview) auto-expand rows — causing the table to
        grow past the footer even though the python-pptx shape is technically in-bounds.

        This method estimates the required row height from text length + column width,
        then sets each data row to max(template_height, content_height, MIN_DATA_ROW).
        After this call, _fit_table_to_bounds sees the true height and can compress
        or trigger a split correctly.
        """
        try:
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            emu = self._EMU_PER_INCH

            cols = list(table.columns)
            col_widths_inch = [max(c.width / emu, 0.5) for c in cols] if cols else []

            tbl = table._tbl
            tr_elements = tbl.findall('a:tr', namespaces=ns)
            changed = 0

            for idx, tr in enumerate(tr_elements):
                if idx == 0:
                    continue  # skip header row

                current_h = int(tr.get('h', '0'))
                min_h = max(current_h, self._MIN_DATA_ROW_HEIGHT_EMU)

                cells = []
                for j, tc in enumerate(tr.findall('a:tc', namespaces=ns)):
                    if j >= len(col_widths_inch):
                        break
                    texts = tc.findall('.//a:t', namespaces=ns)
                    cells.append(''.join(t.text or '' for t in texts).strip())

                content_h = self._estimate_row_content_height_emu(cells, col_widths_inch)
                new_h = max(min_h, content_h)

                if new_h != current_h:
                    tr.set('h', str(new_h))
                    changed += 1

            if changed:
                print(f'  - Content-adjusted {changed} row height(s) '
                      f'(min {self._MIN_DATA_ROW_HEIGHT_EMU/emu:.3f}")')
        except Exception as e:
            print(f'  - _apply_content_row_heights error: {e}')

    def _max_rows_for_shape(
        self, shape, rows_data: Optional[List[List[str]]] = None,
        fill_cols=None,
    ) -> int:
        """Compute max data rows that fit above the footer for *shape*.

        `fill_cols` is the role's column mapping, needed so each cell is
        measured against the column it will actually land in — see
        `_grid_positions`.

        When *rows_data* (the actual cell text about to be filled, header
        excluded) is given, each row's real wrap-estimated height is used — the
        same heuristic `_apply_content_row_heights` applies after filling — so
        the split point matches what will actually be rendered instead of
        assuming every row is as short as the template's placeholder text.
        Without *rows_data*, falls back to the template row heights (each
        bumped to at least `_MIN_DATA_ROW_HEIGHT_EMU`) plus the last template
        row's height as the per-new-row cost.
        """
        emu = self._EMU_PER_INCH
        default_row_h = emu // 3  # ~0.33" when a row has no explicit height

        try:
            avail_height = max(0, self._SVN_SAFE_BOTTOM_EMU - shape.top)

            if not shape.has_table:
                return max(1, avail_height // self._MIN_ROW_HEIGHT_EMU)

            table = shape.table
            all_rows = list(table.rows)
            if not all_rows:
                return max(1, avail_height // self._MIN_ROW_HEIGHT_EMU)

            # No header to reserve under `fill_cols`: that mode's markdown is
            # row-only and the grid's first row carries data. See
            # `_trim_table_cells_to_fit`, which reserves it on the same terms.
            header_h = 0 if fill_cols else (all_rows[0].height or default_row_h)
            avail_for_data = max(0, avail_height - header_h)

            if rows_data:
                col_widths_inch = [max(c.width / emu, 0.5) for c in table.columns]
                result = 0
                used = 0
                for cells in rows_data:
                    row_h = max(
                        self._MIN_DATA_ROW_HEIGHT_EMU,
                        self._estimate_row_content_height_emu(
                            self._row_to_grid(cells, fill_cols, len(col_widths_inch)),
                            col_widths_inch),
                    )
                    if result and used + row_h > avail_for_data:
                        break
                    used += row_h
                    result += 1

                print(f'  - max_rows_for_shape (content-aware): avail={avail_height/emu:.2f}" '
                      f'header={header_h/emu:.2f}" data={used/emu:.2f}" → {result} total data rows')
                return max(1, result)

            data_rows = all_rows[1:] if len(all_rows) > 1 else all_rows

            # _apply_content_row_heights will bump every data row to at least
            # _MIN_DATA_ROW_HEIGHT_EMU after filling.  Use that same minimum here so
            # the pre-fill split threshold is consistent with the post-fill actual heights.
            effective_min = self._MIN_DATA_ROW_HEIGHT_EMU

            # Total height assuming template rows are bumped to at least effective_min.
            template_data_h = sum(
                max(r.height or default_row_h, effective_min) for r in data_rows
            )
            template_data_count = len(data_rows)

            # New rows cloned from last template row — also subject to effective_min.
            last_row_h = (data_rows[-1].height or default_row_h) if data_rows else default_row_h
            new_row_h = max(last_row_h, effective_min)

            remaining = max(0, avail_for_data - template_data_h)
            max_new = int(remaining / new_row_h) if new_row_h else 0
            result = template_data_count + max_new

            print(f'  - max_rows_for_shape: avail={avail_height/emu:.2f}" '
                  f'header={header_h/emu:.2f}" '
                  f'tmpl_data={template_data_h/emu:.2f}" ({template_data_count} rows) '
                  f'new_row={new_row_h/emu:.2f}" → {result} total data rows')
            return max(1, result)
        except Exception as e:
            fallback = max(1, self._SVN_SAFE_BOTTOM_EMU // self._MIN_ROW_HEIGHT_EMU)
            print(f'  - _max_rows_for_shape error ({e}); fallback={fallback}')
            return fallback

    def _normalize_table_col_widths(self, table, min_col_width_emu: int = 914400):
        """Ensure no column is narrower than *min_col_width_emu* (default 1 inch).

        When a column is below the minimum, it is expanded to the minimum and the
        surplus is subtracted proportionally from the wider columns, so the total
        table width stays the same.
        """
        try:
            cols = list(table.columns)
            if not cols:
                return
            widths = [c.width for c in cols]
            total_width = sum(widths)

            needs_expand = [w < min_col_width_emu for w in widths]
            if not any(needs_expand):
                return  # all columns already wide enough

            new_widths = list(widths)
            for i, below in enumerate(needs_expand):
                if below:
                    new_widths[i] = min_col_width_emu

            # Amount already locked to narrow columns
            locked_total = sum(w for i, w in enumerate(new_widths) if needs_expand[i])
            remaining = total_width - locked_total
            free_total = sum(w for i, w in enumerate(widths) if not needs_expand[i])

            if free_total > 0 and remaining > 0:
                scale = remaining / free_total
                for i, below in enumerate(needs_expand):
                    if not below:
                        new_widths[i] = max(min_col_width_emu, int(widths[i] * scale))

            # Fix rounding drift so total is exactly preserved
            diff = total_width - sum(new_widths)
            if diff != 0:
                # Apply diff to the widest free column
                widest = max((i for i, b in enumerate(needs_expand) if not b),
                             key=lambda i: new_widths[i], default=0)
                new_widths[widest] += diff

            for col, w in zip(cols, new_widths):
                col.width = w

            print(f'  - Normalized col widths: {[round(w / 914400, 2) for w in new_widths]}" each')
        except Exception as e:
            print(f'  - Error normalizing table col widths: {e}')

    def _fit_table_to_bounds(self, shape):
        """Constrain the table inside the slide's safe content area.

        After rows are added/cloned the table can grow taller than the area
        above the footer.  This method:
          1. Measures total row height after fill.
          2. If it fits → updates shape.height to the actual content height
             (removes dead space / avoids shape box drifting below table).
          3. If it overflows → scales row heights down toward each row's real
             content-required floor, but never below it. A PPTX row's `h` is
             only a *minimum* — PowerPoint/LibreOffice auto-grow a row to fit
             its real text regardless of what we declare, so compressing past
             the content-required height doesn't actually shrink anything; it
             just makes this tool believe the table fits while the real
             renderer still overflows it into the footer.
        """
        try:
            if not shape.has_table:
                return
            table = shape.table
            rows = list(table.rows)
            if not rows:
                return

            total_h = sum(r.height or 0 for r in rows)
            avail_h = max(0, self._SVN_SAFE_BOTTOM_EMU - shape.top)

            if total_h <= avail_h:
                shape.height = total_h   # sync shape box to actual content
                return

            # Need to compress — but only the slack above each row's real
            # content-required height; never below what the text itself needs.
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            col_widths_inch = [max(c.width / self._EMU_PER_INCH, 0.5) for c in table.columns]
            tr_elements = table._tbl.findall('a:tr', namespaces=ns)
            floors = []
            for idx, tr in enumerate(tr_elements):
                if idx == 0:
                    floors.append(self._MIN_ROW_HEIGHT_EMU)
                    continue
                cells = [
                    ''.join(t.text or '' for t in tc.findall('.//a:t', namespaces=ns)).strip()
                    for tc in tr.findall('a:tc', namespaces=ns)
                ]
                floors.append(max(
                    self._MIN_ROW_HEIGHT_EMU,
                    self._estimate_row_content_height_emu(cells, col_widths_inch),
                ))

            floor_total = sum(floors)
            if floor_total >= avail_h:
                print(f'  - _fit_table_to_bounds: content floor '
                      f'{floor_total / 914400:.2f}" already ≥ avail '
                      f'{avail_h / 914400:.2f}" — cannot compress further without '
                      f'clipping text; table may overflow the footer')
                return

            slack_total = total_h - floor_total
            scale = (avail_h - floor_total) / slack_total if slack_total else 0
            new_heights = [
                floor + int(((r.height or 0) - floor) * scale)
                for r, floor in zip(rows, floors)
            ]

            # Fix rounding drift: adjust the tallest row
            diff = avail_h - sum(new_heights)
            if diff != 0:
                idx = max(range(len(new_heights)), key=lambda i: new_heights[i])
                new_heights[idx] += diff

            for row, h in zip(rows, new_heights):
                row.height = h
            shape.height = sum(r.height for r in rows)
            print(f'  - Fit table to safe area: '
                  f'{total_h / 914400:.2f}" → {shape.height / 914400:.2f}" '
                  f'(safe bottom {self._SVN_SAFE_BOTTOM_EMU / 914400:.2f}")')
        except Exception as e:
            print(f'  - _fit_table_to_bounds error: {e}')

    def _find_shape(self, slide, target):
        """Find shape by index or name.

        When only ``shape_name`` is provided (no ``shape_index``), the search
        is performed recursively through GROUP shapes so that shapes nested
        inside groups can be targeted by name.
        """
        try:
            # Try by index first (faster, top-level only)
            if target.shape_index is not None and target.shape_index < len(slide.shapes):
                shape = slide.shapes[target.shape_index]
                print(f'  - Found shape by index {target.shape_index}: {shape.name}')
                return shape

            # Fallback to recursive name search (traverses group children)
            if target.shape_name:
                found = self._find_shape_recursive(slide.shapes, target.shape_name)
                if found:
                    print(f'  - Found shape by name (recursive): {target.shape_name}')
                    return found
        except Exception as e:
            print(f'  - Error finding shape: {e}')

        return None

    def _find_shape_recursive(self, shapes, name: str):
        """Recursively search for a shape by name, traversing GROUP children.

        Args:
            shapes: A shape collection (``slide.shapes`` or ``group_shape.shapes``).
            name: The exact shape name to search for.

        Returns:
            The matching shape, or ``None`` if not found.
        """
        try:
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            for shape in shapes:
                if shape.name == name:
                    return shape
                # Recurse into group shapes
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        found = self._find_shape_recursive(shape.shapes, name)
                        if found:
                            return found
                except AttributeError:
                    pass
        except Exception as e:
            print(f'  - _find_shape_recursive error: {e}')
        return None

    def _fill_slide_default(self, slide, slide_content: SlideContent):
        """Default filling strategy (original logic)"""

        # Update title
        self._update_slide_title(slide, slide_content.title)

        # Get available text containers
        placeholders = self._get_text_placeholders(slide)
        text_shapes = self._get_text_shapes(slide)

        # Fill content items
        placeholder_idx = 0
        for content_item in slide_content.content:
            if content_item.type == ContentType.TEXT:
                placeholder_idx = self._fill_text_content(
                    slide, content_item, placeholders, text_shapes, placeholder_idx,
                )

            elif content_item.type == ContentType.LIST:
                placeholder_idx = self._fill_list_content(
                    slide, content_item, placeholders, placeholder_idx,
                )

            elif content_item.type == ContentType.TABLE:
                self._add_table(slide, content_item, Inches(1.0), Inches(2.5), Inches(8.0))

            elif content_item.type == ContentType.IMAGE:
                self._add_image(slide, content_item, Inches(1.0), Inches(2.5), Inches(8.0))

    def _get_text_placeholders(self, slide):
        """Get all text placeholders in slide (excluding title)"""
        return [
            s for s in slide.shapes
            if s.is_placeholder and s.has_text_frame and s != slide.shapes.title
        ]

    def _get_text_shapes(self, slide):
        """Get all text shapes in slide (excluding placeholders and title)"""
        return [
            s for s in slide.shapes
            if s.has_text_frame and not s.is_placeholder and s != slide.shapes.title
        ]

    def _fill_text_content(self, slide, content_item, placeholders, text_shapes, placeholder_idx):
        """Fill text content into placeholders or text shapes"""
        if placeholder_idx < len(placeholders):
            placeholders[placeholder_idx].text_frame.text = content_item.data
            self._apply_text_font(placeholders[placeholder_idx].text_frame)
            print(f'  - Filled placeholder {placeholder_idx}: {content_item.data[:50]}...')
            return placeholder_idx + 1

        elif text_shapes:
            text_shapes[0].text_frame.text = content_item.data
            self._apply_text_font(text_shapes[0].text_frame)
            text_shapes.pop(0)
            return placeholder_idx

        else:
            # Add new textbox if no placeholders left
            self._add_content_at_position(slide, content_item, Inches(1.0), Inches(2.0), Inches(8.0))
            return placeholder_idx

    def _fill_list_content(self, slide, content_item, placeholders, placeholder_idx):
        """Fill list content into placeholder"""
        if placeholder_idx < len(placeholders):
            text_frame = placeholders[placeholder_idx].text_frame
            text_frame.clear()
            for i, item in enumerate(content_item.data):
                if i > 0:
                    text_frame.add_paragraph()
                p = text_frame.paragraphs[i]
                p.text = f'• {item}'
                p.level = 0
            self._apply_text_font(text_frame)
            print(f'  - Filled list with {len(content_item.data)} items')
            return placeholder_idx + 1
        else:
            # Add new list
            top = Inches(2.5 + (placeholder_idx * 0.5))
            self._add_content_at_position(slide, content_item, Inches(1.0), top, Inches(8.0))
            return placeholder_idx

    def _add_content_at_position(self, slide, content_item: ContentItem, left, top, width):
        """Add content item at specific position"""
        if content_item.type == ContentType.TEXT:
            height = Inches(0.5)
            textbox = slide.shapes.add_textbox(left, top, width, height)
            textbox.text_frame.text = content_item.data
            self._apply_text_font(textbox.text_frame)

        elif content_item.type == ContentType.LIST:
            height = Inches(0.3) * len(content_item.data)
            textbox = slide.shapes.add_textbox(left, top, width, height)
            text_frame = textbox.text_frame

            for i, item in enumerate(content_item.data):
                if i > 0:
                    text_frame.add_paragraph()
                p = text_frame.paragraphs[i]
                p.text = f'• {item}'
            self._apply_text_font(text_frame)

    def _add_image(self, slide, content_item: ContentItem, left, top, width):
        """Add image to slide"""
        try:
            image_data = content_item.data
            image_path = self._resolve_image_path(image_data)

            if Path(image_path).exists():
                slide.shapes.add_picture(str(image_path), left, top, width=width)
                self._add_image_caption(slide, image_data, left, top, width)
            else:
                print(f'Image not found: {image_path}')
                self._add_placeholder_text(slide, f'[Image: {image_path}]', left, top, width)

        except Exception as e:
            print(f'Error adding image: {e}')

    def _resolve_image_path(self, image_data) -> str:
        """Resolve image path from data"""
        image_path = image_data.get('path') if isinstance(image_data, dict) else image_data

        if image_path and not Path(str(image_path)).is_absolute():
            # Try relative to CWD first
            rel_path = Path(str(image_path))
            if not rel_path.exists():
                # Try outputs directory
                image_path = str(Path('outputs') / str(image_path))

        return str(image_path) if image_path else ''

    def _add_image_caption(self, slide, image_data, left, top, width):
        """Add caption below image if exists"""
        if isinstance(image_data, dict) and image_data.get('caption'):
            caption_top = top + Inches(3.0)
            caption_box = slide.shapes.add_textbox(left, caption_top, width, Inches(0.3))
            caption_box.text_frame.text = image_data['caption']
            for paragraph in caption_box.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Noto Sans JP'
                    run.font.size = Pt(10)
                    run.font.italic = True

    @staticmethod
    def _set_text_preserving_format(text_frame, value: str) -> None:
        """Set a text frame's text, reusing the first existing run instead of
        `text_frame.text = value`. That setter replaces every paragraph/run
        with a single freshly-styled one, silently discarding any
        template-authored run formatting (e.g. this template's "after" title
        cells are colored/bold via their own run, not the placeholder's
        paragraph/shape defaults — resetting the run drops that styling and
        the text renders in whatever default color the theme falls back to,
        which can be invisible against the slide's own background)."""
        paragraphs = text_frame.paragraphs
        if paragraphs and paragraphs[0].runs:
            first_para = paragraphs[0]
            first_para.runs[0].text = value
            for extra_run in list(first_para.runs[1:]):
                extra_run._r.getparent().remove(extra_run._r)
            for extra_para in list(paragraphs[1:]):
                extra_para._p.getparent().remove(extra_para._p)
        else:
            text_frame.text = value
        PPTXRenderer._shrink_text_not_the_shape(text_frame)

    @staticmethod
    def _shrink_text_not_the_shape(text_frame) -> None:
        """Turn `spAutoFit` into `normAutofit` on a box we have just written to.

        The template sets many of its boxes to *resize the shape to fit the
        text* — right for the proposal it was authored for, wrong for any other,
        because this deck's designs are fixed grids. Slide 16's four point cards
        are the clearest case: a long chapter title grew its own box from 0.20"
        to 0.57", whose bottom edge (2.37") then ran straight through the
        headline box beneath it (top 2.33"), and the four cards' bodies came to
        rest at four different heights.

        The text is what has to give, not the geometry. Both directions of
        auto-fit are the renderer's, so the only question is which one, and on a
        drawn grid it is always this one.
        """
        from pptx.oxml.ns import qn

        bodyPr = text_frame._txBody.find(qn('a:bodyPr'))
        if bodyPr is None:
            return
        grow = bodyPr.findall(qn('a:spAutoFit'))
        if not grow:
            return
        for element in grow:
            bodyPr.remove(element)
        if bodyPr.find(qn('a:normAutofit')) is None:
            bodyPr.append(bodyPr.makeelement(qn('a:normAutofit'), {}))

    # A paragraph's own bullet lives in its `<a:pPr>`, as a `buChar` (a glyph the
    # template chose) or a `buAutoNum` (an auto-numbered list).
    _A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    @classmethod
    def _native_bullet(cls, text_frame) -> bool:
        """True when this box bullets its own paragraphs.

        Slide 4's two columns do: every item paragraph carries `buChar='●'`,
        with empty spacer paragraphs between them. A box like that must NOT be
        handed a glyph in its text — doing so printed `● ■ item`, one marker
        from the template and one from us, in two different shapes.
        """
        for para in text_frame.paragraphs:
            pPr = para._p.find(f'{{{cls._A_NS}}}pPr')
            if pPr is None:
                continue
            if (pPr.find(f'{{{cls._A_NS}}}buChar') is not None
                    or pPr.find(f'{{{cls._A_NS}}}buAutoNum') is not None):
                return True
        return False

    # A line that ends in a colon and is followed by others introduces them —
    # `Concrete success indicators (to be finalized with Evering):`. Bulleting it
    # states it as one of the objectives it is announcing, so the reader counts
    # five where the chapter listed four. It keeps its words and loses its
    # marker; dropping the line instead would lose content the author wrote.
    _LEAD_IN_LINE = re.compile(r'[:：]\s*$')

    @classmethod
    def _unbullet_paragraph(cls, p_element) -> None:
        """Replace a paragraph's own bullet with an explicit `buNone`."""
        ns = f'{{{cls._A_NS}}}'
        pPr = p_element.find(f'{ns}pPr')
        if pPr is None:
            return
        for tag in ('buChar', 'buAutoNum', 'buFont', 'buClr'):
            for el in pPr.findall(f'{ns}{tag}'):
                pPr.remove(el)
        if pPr.find(f'{ns}buNone') is None:
            pPr.append(pPr.makeelement(f'{ns}buNone', {}))

    @classmethod
    def _set_text_lines_preserving_format(cls, text_frame, lines: List[str]) -> None:
        """Write one PARAGRAPH per line, each keeping the template's formatting.

        `_set_text_preserving_format` writes everything into a single paragraph,
        which is right for prose and wrong for a list: a template that bullets
        its paragraphs renders one bullet for the whole block, so five items
        arrived as one bulleted run with four unmarked lines under it.

        The first paragraph is the model — its `<a:pPr>` (bullet, indent, line
        spacing) and its run's `<a:rPr>` (font, colour, weight) are cloned for
        every line, so the deck's own list styling does the marking.
        """
        paragraphs = text_frame.paragraphs
        if not paragraphs or not paragraphs[0].runs or not lines:
            cls._set_text_preserving_format(text_frame, '\n'.join(lines))
            return
        model = paragraphs[0]
        model.runs[0].text = lines[0]
        for extra_run in list(model.runs[1:]):
            extra_run._r.getparent().remove(extra_run._r)
        for extra_para in list(paragraphs[1:]):
            extra_para._p.getparent().remove(extra_para._p)
        # Cloned from a pristine copy: the model's own element is about to be
        # un-bulleted if line 0 is a lead-in, and every later line would inherit
        # that.
        blueprint = copy.deepcopy(model._p)
        anchor = model._p
        for line in lines[1:]:
            clone = copy.deepcopy(blueprint)
            anchor.addnext(clone)
            anchor = clone
            for run in clone.findall(f'{{{cls._A_NS}}}r'):
                text_el = run.find(f'{{{cls._A_NS}}}t')
                if text_el is not None:
                    text_el.text = line
                    break
            # A cloned paragraph must carry exactly one run, or the line would
            # repeat once per run the model happened to hold.
            for surplus in clone.findall(f'{{{cls._A_NS}}}r')[1:]:
                clone.remove(surplus)
        # Last pass, once every paragraph exists: a lead-in keeps its words and
        # loses its marker.
        element = model._p
        for i, line in enumerate(lines):
            if element is None:
                break
            if i < len(lines) - 1 and cls._LEAD_IN_LINE.search(line):
                cls._unbullet_paragraph(element)
            element = element.getnext()

    def _apply_text_font(self, text_frame, font_name: str = 'Noto Sans JP', font_size: int = 10):
        """Apply font and size to all text in a text frame"""
        try:
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = font_name
                    run.font.size = Pt(font_size)
        except Exception as e:
            print(f'  - Error applying font to text: {e}')

    def _apply_line_spacing(self, text_frame, spacing_pt):
        """Apply line spacing (space after) to all paragraphs in a text frame
        
        Args:
            text_frame: The text frame to apply spacing to
            spacing_pt: Spacing in points (e.g., Pt(10) for 10pt)
        """
        try:
            for paragraph in text_frame.paragraphs:
                if paragraph.text.strip():  # skip empty paragraphs (from blank lines)
                    paragraph.space_after = spacing_pt
        except Exception as e:
            print(f'  - Error applying line spacing: {e}')

    @classmethod
    def _avg_em_width(cls, sample: str = '') -> float:
        """Average per-character width of *sample*, in em.

        `sample` empty returns 1.0 — the CJK assumption every caller made before
        this existed, so an un-sampled call measures exactly as it always did.

        Passing the real string matters: a box 4.15" wide at 10pt holds 29 CJK
        characters per line but 49 latin ones. Measuring latin copy at one em
        under-estimated its capacity by ~1.7x, which truncated text that fit
        comfortably and then spent a whole appendix slide carrying the sentence
        it had just cut.
        """
        if not sample:
            return 1.0
        return sum(cls._em_width(ch) for ch in sample) / len(sample)

    def _estimate_max_chars(self, shape, font_pt: int = 8, sample: str = '') -> int:
        """Estimate max characters that fit in a shape based on its dimensions.

        Character width follows *sample*'s own script mix (`_avg_em_width`);
        omitting it keeps the historical CJK-width behavior. Shape height gives
        how many lines fit. Returns a floor of 8 so at least a short label is
        always written.
        """
        try:
            emu = self._EMU_PER_INCH
            w_inch = (shape.width or 0) / emu
            h_inch = (shape.height or 0) / emu
            char_w = font_pt * self._avg_em_width(sample) / 72
            line_h = font_pt * 1.35 / 72  # ~0.150" line height with spacing
            chars_per_line = max(5, int(w_inch / char_w))
            lines = max(1, int(h_inch / line_h))
            return max(8, chars_per_line * lines)
        except Exception:
            return 65

    def _estimate_chars_per_line(self, shape, font_pt: int = 8, sample: str = '') -> int:
        """Characters of *sample*'s script mix that fit on one line of *shape*."""
        try:
            w_inch = (shape.width or 0) / self._EMU_PER_INCH
            char_w = font_pt * self._avg_em_width(sample) / 72
            return max(5, int(w_inch / char_w))
        except Exception:
            return 5

    @staticmethod
    def _lines_needed(text: str, chars_per_line: int) -> int:
        """Lines *text* really occupies when wrapped at *chars_per_line*.

        Counted per paragraph, because a blank-line break ends a line wherever
        it falls: three short paragraphs occupy three lines, not the one line
        their combined length would suggest. Comparing a flat character total
        against a flat capacity silently assumed perfect packing and truncated
        text that had room left.
        """
        if not text.strip():
            return 1
        paragraphs = [p for p in re.split(r'\n\s*\n', text.strip()) if p.strip()]
        return sum(max(1, -(-len(p) // max(1, chars_per_line))) for p in paragraphs)

    def _estimate_extra_wrap_height(self, shape, text: str, font_pt: int = 10) -> int:
        """EMU height needed for `text` to word-wrap in `shape`'s current
        width, beyond what a single line already occupies. Used to grow a
        title box to fit its full text instead of truncating it with an
        ellipsis — a short heading cut off mid-word reads as meaningless."""
        try:
            emu = self._EMU_PER_INCH
            chars_per_line = self._estimate_chars_per_line(shape, font_pt, sample=text)
            lines_needed = -(-len(text) // chars_per_line) if text else 1  # ceil
            extra_lines = max(0, lines_needed - 1)
            line_h_inch = font_pt * 1.35 / 72
            return int(extra_lines * line_h_inch * emu)
        except Exception:
            return 0

    @staticmethod
    def _set_vertical_anchor_top(text_frame) -> None:
        """Anchor text to the top of the shape so it doesn't float to center."""
        try:
            from pptx.enum.text import MSO_ANCHOR
            text_frame.vertical_anchor = MSO_ANCHOR.TOP
        except Exception:
            pass

    @staticmethod
    def _split_at_sentence(text: str, max_chars: int = 120) -> Tuple[str, str]:
        """Split *text* into (kept, rest) at the last sentence end within budget.

        `rest` is what `kept` left behind — never a copy of the whole string.
        An appendix slide built from the full text repeated whatever had just
        been written on the slide above it; carrying only the remainder is the
        difference between a continuation and a duplicate.

        Tries Japanese sentence ends (。) then newlines then a hard cut. On a
        hard cut `kept` carries the ellipsis but `rest` starts at the cut, so
        the two still reconstruct the original.
        """
        if len(text) <= max_chars:
            return text, ''
        truncated = text[:max_chars]
        for sep in ['。', '\n', '．', '.']:
            idx = truncated.rfind(sep)
            if idx > max_chars // 2:
                return truncated[:idx + 1].rstrip(), text[idx + 1:].lstrip()
        return truncated.rstrip() + '…', text[max_chars:].lstrip()

    @classmethod
    def _truncate_at_sentence(cls, text: str, max_chars: int = 120) -> str:
        """`_split_at_sentence`'s kept half. Kept for callers that discard the rest."""
        return cls._split_at_sentence(text, max_chars)[0]

    def _apply_auto_fit(self, text_frame):
        """Enable normAutofit so text shrinks to fit the shape instead of overflowing."""
        try:
            from pptx.enum.text import MSO_AUTO_SIZE
            text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception as e:
            print(f'  - Could not apply auto-fit: {e}')

    @staticmethod
    def _cell_style_is_uniform(text_frame) -> bool:
        """True when every run in this cell is styled alike.

        The question behind it: is this cell's formatting the COLUMN's, or the
        sample sentence's? A label chip is white bold throughout, so its style
        belongs to the column and survives new text. A body cell that picked out
        one clause in red bold was styling that clause, and inheriting it would
        put a whole assumption in red.

        An empty cell counts as uniform: there is nothing to disagree.
        """
        styles = set()
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    colour = run.font.color
                    tint = str(getattr(colour, 'rgb', None)
                               or getattr(colour, 'theme_color', None))
                except (AttributeError, TypeError):
                    tint = None
                styles.add((run.font.bold, run.font.italic, tint))
        return len(styles) <= 1

    def _apply_table_cell_font(self, cell, font_name: str = 'Noto Sans JP', font_size: int = 8):
        """Apply font and size to all text in a table cell"""
        try:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = font_name
                    run.font.size = Pt(font_size)
        except Exception as e:
            print(f'  - Error applying font to table cell: {e}')

    def _add_placeholder_text(self, slide, text: str, left, top, width):
        """Add placeholder text box"""
        textbox = slide.shapes.add_textbox(left, top, width, Inches(1))
        textbox.text_frame.text = text
        self._apply_text_font(textbox.text_frame)

    def _add_table(self, slide, content_item: ContentItem, left, top, width):
        """Add table to slide"""
        try:
            table_data = content_item.data
            headers = table_data.get('headers', [])
            rows = table_data.get('rows', [])

            if not headers or not rows:
                print('Table has no headers or rows')
                return

            # Create table
            row_count = len(rows) + 1  # +1 for header
            col_count = len(headers)
            height = Inches(0.3) * row_count

            table_shape = slide.shapes.add_table(row_count, col_count, left, top, width, height)
            table = table_shape.table

            # Fill headers
            self._fill_table_headers(table, headers)

            # Fill data rows
            self._fill_table_rows(table, rows)

        except Exception as e:
            print(f'Error adding table: {e}')

    def _fill_table_headers(self, table, headers: List[str]):
        """Fill table header row"""
        for col_idx, header in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = str(header)
            # Apply font and make header bold
            self._apply_table_cell_font(cell)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

    def _fill_table_rows(self, table, rows: List[List[str]]):
        """Fill table data rows"""
        for row_idx, row_data in enumerate(rows):
            for col_idx, cell_data in enumerate(row_data):
                cell = table.cell(row_idx + 1, col_idx)  # +1 to skip header
                cell.text = str(cell_data)
                self._apply_table_cell_font(cell)

    # ==================================================================
    # Role-based rendering (new pipeline: profile.md -> ProjectProfile -> PPTX)
    # ==================================================================

    def render_from_profile(
        self,
        profile: ProjectProfile,
        template: str,
        output_name: str,
        output_dir: Optional[str] = None,
        extra_slides: Optional[List[dict]] = None,
    ) -> Path:
        """Fill template using a ProjectProfile + role-based slide configs.

        For each configured slide:
          - Resolve each role's data via source_path
          - Discover shapes via ShapeDiscovery
          - Fill shapes (text / table / image / compound pairs)
          - Clear excess shapes
        """
        print(f'Rendering from profile: {output_name}')
        self._output_dir = self._resolve_output_dir(output_dir)
        self._overflow_extras = {}
        # Profile tables already drawn into a template placeholder, so the
        # orphan-table pass does not add a second slide for the same data.
        self._tables_placed: set = set()
        # Whether this document opted into content shaping — see
        # `ProjectProfile.shape_content`. Read by `_fit_text_to_shape`, which is
        # too deep in the call stack to be handed the profile.
        self._shape_content = bool(getattr(profile, 'shape_content', False))
        # {template slide number: any role wrote content}. Drives the prune
        # pass's drop of unwritten slides in a multi-slide section block.
        self._slides_filled: Dict[int, bool] = {}
        self._detect_template_config(template)
        prs = self._load_template(template)
        # Compute safe content area from actual slide height so tables never reach the footer.
        # Footer band occupies roughly the bottom 13% of the slide; 0.87 leaves ~0.36" margin
        # above a footer at 93% (observed: 5.254" on a 5.625" slide → 93.4%).
        if prs.slide_height:
            self._SVN_SAFE_BOTTOM_EMU = int(prs.slide_height * 0.87)
            print(f'  Safe content bottom: {self._SVN_SAFE_BOTTOM_EMU / self._EMU_PER_INCH:.3f}" '
                  f'(slide is {prs.slide_height / self._EMU_PER_INCH:.3f}" tall)')
        if prs.slide_width:
            self._slide_width_emu = prs.slide_width
        # Captured before pruning: pruning deletes the trailing blank
        # end-card slides that make the safest fallback canvas for extra
        # slides, but the SlideLayout object stays valid (it's a separate
        # XML part from the slide instance) so it can still be reused after.
        # Extra slides anchored to a divider-only section (team) ARE that
        # section's body: the divider must survive pruning and be named even
        # when the markdown has no chapter for it.
        self._extra_anchor_sections = {
            e.get('anchor_section') for e in (extra_slides or []) if e.get('anchor_section')}
        blank_canvas_layout = self._find_blank_canvas_layout(prs)
        # Also captured pre-prune: the canonical chrome slide may itself be
        # pruned, but its SlideLayout stays valid afterwards (separate XML part).
        chrome_layout = self._extra_chrome_layout(prs)
        # Same reason, same timing: a donor slide is almost always one the
        # prune pass is about to delete (that is what makes it available to
        # lend), so its shapes must be copied out while it still exists.
        self._donor_library = self._capture_donor_library(prs)
        # The document's own section order, kept for anchor resolution: it is
        # the only record of where a chapter stands once its template slide has
        # been pruned.
        self._section_order = list(profile.section_order)
        spans = self._fill_slides_from_profile(prs, profile)
        last_positions = self._prune_and_reorder(prs, profile, spans)
        all_extra_slides = (list(extra_slides or [])
                            + self._build_orphan_table_extras(profile)
                            + self._build_overflow_extra_slides(set(last_positions)))
        if all_extra_slides:
            self._name_extras_from_agenda(all_extra_slides, profile.agenda.chapters)
            self._label_dropped_rows(all_extra_slides, profile)
            self._render_extra_slides(prs, all_extra_slides, last_positions,
                                      blank_canvas_layout, chrome_layout)
        # Last, so it also reaches continuation slides and anything the extra-
        # slide pass cloned from a note-bearing template slide.
        self._strip_template_notes(prs)
        return self._save_presentation(prs, output_name, output_dir)

    def _build_orphan_table_extras(self, profile: ProjectProfile) -> List[dict]:
        """Give a slide to tables whose section's template slide has no table.

        The schedule chapter is the case: slide 43 is a hand-drawn Gantt made of
        text boxes and an image placeholder, with no table shape, so a milestone
        table in the markdown had nowhere to go and was silently dropped. Rather
        than lose it, it gets its own slide right after the schedule.
        """
        extras = []
        sources = (
            ('schedule.table', profile.schedule.table, '実施スケジュール', 'schedule'),
            # The team chapter has no content slide at all (divider only), so
            # a 役割分担 table it carries is likewise given its own slide.
            ('team.rows', profile.team.rows,
             profile.team.title or profile.section_headings.get('team') or 'プロジェクト体制',
             'team'),
        )
        for source, rows, title, anchor in sources:
            if not rows or source in self._tables_placed:
                # Already drawn into a placeholder — one table, one slide.
                continue
            headers = list(rows[0].keys())
            print(f'  - {anchor} table ({len(rows)} rows) has no table shape on its '
                  f'template slide — giving it a slide of its own')
            extras.append({
                'layout': 'table',
                'title': title,
                'headers': headers,
                'rows': [[str(row.get(header, '')) for header in headers] for row in rows],
                'anchor_section': anchor,
            })
        return extras

    def _name_extras_from_agenda(self, entries: List[dict], chapters: List[str]) -> None:
        """Title each extra slide after the AGENDA chapter it belongs to.

        An extra slide is a chapter the deck has no bespoke template slide for,
        so it arrives titled with the raw markdown heading — `4. Sun* のスコープ
        — Scope`, complete with its numbering. The proposal's own AGENDA already
        names that chapter, in the wording the reader met on slide 2
        (`Sun* のスコープ・役割分担 — Scope & Responsibility Split`), so the deck
        reads as one document when the slide uses that name instead.

        Matching is on the Japanese label, allowing one to be a prefix of the
        other: a heading is routinely a shortened form of its agenda line.
        A slide matching no chapter — a revision summary, a text overflow
        appendix — keeps its own title and the `Appendix` breadcrumb.
        """
        if not chapters:
            return
        pairs = [self._split_agenda_item(item) for item in chapters]
        for entry in entries:
            title = (entry.get('title') or '').strip()
            if not title:
                continue
            label, gloss = self._agenda_match(title, pairs)
            if not label:
                continue
            if gloss:
                entry['breadcrumb'] = gloss
            # `プロジェクト体制（要件定義）` extends its chapter's name with a
            # qualifier that tells two slides of one chapter apart — keep it,
            # and take only the breadcrumb from the AGENDA.
            from lib.chapter_dividers import extends_label
            if extends_label(title, label):
                continue
            entry['title'] = label
            print(f'  - Extra slide {title!r} renamed after its AGENDA chapter: {label!r}')

    # A chapter name may be written short on a slide and long in the AGENDA
    # (`Sun*のスコープ` vs `Sun* のスコープ・役割分担`). The longer form is the
    # same chapter only when it EXTENDS the shorter one at a separator — a bare
    # prefix is not enough: `お見積り` (Pricing) is a prefix of
    # `お見積りの前提条件` (Estimated Assumptions), two different chapters, and
    # matching them renamed the 前提条件 divider to Pricing and shipped it.
    # Japanese labels are short and share prefixes constantly, so the boundary
    # is what carries the meaning.
    _AGENDA_EXTENSION_SEPARATORS = ('・', '／', '/', '＆', '&', '、', ',',
                                    '（', '(', '－', '—', '–', '-', '：', ':')

    @classmethod
    def _agenda_match(cls, title: str, pairs: List[Tuple[str, str]]) -> Tuple[str, str]:
        """Find the (label, gloss) agenda pair naming *title*'s chapter.

        Exact match wins outright. Only when nothing matches exactly is a
        separator-bounded extension accepted, so a chapter is never claimed by
        another whose name merely starts with the same characters.
        """
        from lib.section_aliases import heading_variants

        # Whitespace is stripped, not just collapsed: Japanese does not need
        # spaces, so a chapter name is written `Sun*のスコープ` on one slide and
        # `Sun* のスコープ・役割分担` in the AGENDA purely by an author's habit.
        def key_of(text: str) -> str:
            return re.sub(r'\s+', '', text.replace('　', '')).casefold()

        def extends(longer: str, shorter: str) -> bool:
            """True when *longer* is *shorter* continued at a separator."""
            if not longer.startswith(shorter) or longer == shorter:
                return False
            return longer[len(shorter):].startswith(cls._AGENDA_EXTENSION_SEPARATORS)

        variants = [key_of(variant) for variant in heading_variants(title)]
        variants = [variant for variant in variants if variant]
        if not variants:
            return '', ''
        for label, gloss in pairs:
            if key_of(label) in variants:
                return label, gloss
        for label, gloss in pairs:
            key = key_of(label)
            if any(extends(key, v) or extends(v, key) for v in variants):
                return label, gloss
        return '', ''

    def _build_overflow_extra_slides(self, kept_slides) -> List[dict]:
        """Turn queued per-item text overflow (`_queue_overflow_extra`) into
        appendix extra-slide entries, one per (section, role) group, anchored
        right after that section so the full text lands near where it was cut.

        *kept_slides* is what survived the prune pass, and a remainder from a
        slide that did not survive is discarded with it. Every configured slide
        is filled, including the ones no section claims — slide 15 draws the
        same 提案の要諦 cards as slide 16 in three slots instead of four — so
        the pruned twin was truncating text into the same queue as the slide
        that ships. The reader met an appendix slide holding `s" and "medical
        device" claims.`: a fragment cut from a sentence the deck prints in
        full two slides earlier, by a slide they never saw.
        """
        entries = []
        for (section_key, role_name), queued in self._overflow_extras.items():
            texts, seen = [], set()
            for slide_num, remainder in queued:
                if slide_num and slide_num not in kept_slides:
                    continue
                if remainder in seen:
                    continue
                seen.add(remainder)
                texts.append(remainder)
            if not texts:
                continue
            label = role_name.replace('_', ' ').title()
            entries.append({
                'layout': 'bullets',
                'title': f'{label} — Additional Detail',
                'bullets': texts,
                'anchor_section': section_key,
            })
        return entries

    # ------------------------------------------------------------------
    # Donor slides — appendix chapters borrow a real template design
    # ------------------------------------------------------------------

    def _capture_donor_library(self, prs: Presentation) -> Dict[int, tuple]:
        """Copy every donor slide's shapes out of *prs* before pruning.

        MUST run pre-prune. A donor is lendable precisely because no section
        claimed it, which is exactly the condition `_prune_and_reorder` deletes
        on — by the time the extra-slide pass runs, the slide is gone. The same
        hazard is why `blank_canvas_layout` and `chrome_layout` are captured
        here; this is the third case of it.

        Stores the shape tree deep-copied (so later edits to the original cannot
        reach it), plus the layout and the source part needed to re-bind image
        relationships when the copy is planted.
        """
        library: Dict[int, tuple] = {}
        if not self.template_config:
            return library
        try:
            from lib.templates.svn import DONOR_SLIDES
        except Exception:
            return library
        for slide_no in DONOR_SLIDES:
            if slide_no > len(prs.slides):
                continue
            slide = prs.slides[slide_no - 1]
            library[slide_no] = (copy.deepcopy(slide.shapes._spTree),
                                 slide.slide_layout, slide.part)
        if library:
            print(f'  Donor designs captured pre-prune: {sorted(library)}')
        return library

    def _insert_donor_slide(self, prs: Presentation, slide_no: int, target_idx: int):
        """Plant a captured donor's shapes as a new slide at *target_idx*."""
        spTree, layout, source_part = self._donor_library[slide_no]
        new_slide = prs.slides.add_slide(layout)
        self._ensure_unique_slide_partname(prs, new_slide)
        new_tree = new_slide.shapes._spTree
        for child in list(new_tree):
            new_tree.remove(child)
        for child in copy.deepcopy(spTree):
            new_tree.append(child)
        # Same re-binding `_duplicate_slide_at` does: without it the copy's
        # blipFills resolve their rIds against the wrong relationship table and
        # the images vanish.
        self._remap_shape_rels(source_part, new_slide.part, new_tree)
        self._move_slide_to(prs, target_idx)
        return prs.slides[target_idx]

    def _move_slide_to(self, prs: Presentation, target_idx: int) -> None:
        """Move the just-appended slide to *target_idx* in the slide order."""
        pml_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        try:
            sldIdLst = prs.part._element.find(f'{{{pml_ns}}}sldIdLst')
            if sldIdLst is None:
                return
            new_sldId = list(sldIdLst)[-1]
            sldIdLst.remove(new_sldId)
            remaining = list(sldIdLst)
            if target_idx >= len(remaining):
                sldIdLst.append(new_sldId)
            else:
                remaining[target_idx].addprevious(new_sldId)
        except Exception as e:
            print(f'  - Error positioning slide: {e}')

    def _fill_donor_slide(self, slide, slide_no: int, entry: dict) -> None:
        """Fill a planted donor with *entry*'s items, through its own roles.

        The chapter's items are bound to the attribute the donor's roles read,
        so the roles work unchanged — the same config that fills this slide from
        a profile fills it from an appendix chapter. No continuation slides: the
        classifier only offers a donor whose capacity already covers every item.
        """
        from types import SimpleNamespace

        config = self.template_config.slide_role_configs.get(slide_no)
        if config is None:
            return
        source = (entry.get('donor_source') or '').split('.')[0]
        data = SimpleNamespace(**{source: entry.get('donor_items', [])})
        discovery = ShapeDiscovery()
        for role in config.roles:
            shapes = discovery.discover(slide, role)
            items = self._resolve_source_path(data, role.source_path)
            if role.sub_roles:
                self._fill_compound_pairs(shapes, items if isinstance(items, list) else [],
                                          role, slide=slide)
            else:
                self._fill_text_shapes(shapes, self._coerce_text_items(items, role.item_field),
                                       role, slide=slide)
        self._flush_shared_fonts()
        self._apply_donor_header(slide, entry)

    def _apply_donor_header(self, slide, entry: dict) -> None:
        """Retitle a donor after the chapter borrowing it.

        Without this the slide keeps the header of the chapter it was designed
        for — the reader would meet `提案の要諦` above a chapter about something
        else. The `Appendix` lead-in matches the drawn layouts', so a borrowed
        design still reads as appendix material rather than as a bespoke slide.
        """
        title = (entry.get('title') or '').strip()
        if not title:
            return
        prefix = (entry.get('breadcrumb') or 'Appendix').strip()
        # Most donors head themselves with the deck's two placeholders. Some draw
        # their own header instead, inside a group where those names do not
        # exist — those name their shapes in the donor spec. Missing both, the
        # slide would keep the Japanese heading of the chapter it was drawn for.
        from lib.templates.svn import DONOR_SLIDES

        spec = DONOR_SLIDES.get(entry.get('donor')) or {}
        title_name, crumb_name = spec.get('header') or ('PlaceHolder 1', 'PlaceHolder 2')
        for shape in self._walk_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            if shape.name == title_name:
                self._set_text_preserving_format(shape.text_frame, title)
            elif shape.name == crumb_name:
                self._set_text_preserving_format(shape.text_frame,
                                                 f'{prefix}  ｜  {title}')

    def _find_blank_canvas_layout(self, prs: Presentation):
        """Find a clean, minimal-chrome layout to fall back on for extra slides.

        Scans from the end of the deck since the SVN template's trailing
        slides are blank end-cards with no decorative content-slide styling
        (unlike e.g. project_background's two-tone split-panel master).
        """
        slide_area = (prs.slide_width or 0) * (prs.slide_height or 0)
        for slide in reversed(prs.slides):
            if self._is_clean_master(slide.slide_layout.slide_master, slide_area):
                return slide.slide_layout
        return None

    def _fill_slides_from_profile(
        self, prs: Presentation, profile: ProjectProfile,
    ) -> Dict[int, Tuple[int, int]]:
        """Fill configured slides; return map {template_slide_num: (start_idx, end_idx)}.

        ``start_idx`` is the slide's current (0-based) index at fill time;
        ``end_idx`` is the current index of its last overflow-continuation
        slide (``end_idx == start_idx`` when the slide didn't overflow). The
        prune+reorder pass uses this block to keep overflow continuations
        contiguous with their parent.
        """
        spans: Dict[int, Tuple[int, int]] = {}
        if not self.template_config:
            print('No template config — cannot render from profile')
            return spans

        self._fill_cover_slide(prs, profile)
        self._fill_agenda_slide(prs, profile)
        self._apply_slide_text_overrides(prs)
        self._apply_slide_texts_from_profile(prs, profile)
        self._retitle_slides_from_headings(prs, profile)
        self._retitle_dividers_from_sections(prs, profile)
        self._retitle_dividers_from_agenda(prs, profile)
        self._translate_template_chrome(prs, profile)

        configs = self.template_config.slide_role_configs
        sorted_nums = sorted(configs.keys())
        slide_offset = 0
        discovery = ShapeDiscovery()

        for slide_num in sorted_nums:
            if not self.template_config.should_fill_slide(slide_num):
                print(f'Slide {slide_num} is protected — skipping')
                continue
            adj_idx = slide_num - 1 + slide_offset
            if adj_idx >= len(prs.slides):
                print(f'Slide {slide_num} not found in template (have {len(prs.slides)})')
                continue
            slide = prs.slides[adj_idx]
            role_config = configs[slide_num]
            print(f'Filling slide {slide_num} (template index {adj_idx}): '
                  f'{len(role_config.roles)} role(s)')
            self._filling_slide_num = slide_num
            extra = self._fill_slide_roles(prs, slide, role_config, profile, adj_idx, discovery)
            self._filling_slide_num = 0
            # Span covers the parent slide + any overflow continuation slides.
            spans[slide_num] = (adj_idx, adj_idx + extra)
            slide_offset += extra
        return spans

    # Rough per-character width model — CJK counts as one em, latin as a
    # fraction of one. Born as the cover's fit warning; now also the width model
    # behind `_avg_em_width` / `_estimate_max_chars`, so a body box measures its
    # text the same way the cover measures a title. One rule, one definition.
    # Calibrated against rendered output rather than nominal font metrics: the
    # cover's bold Helvetica Neue sets noticeably wider than a naive 0.55em
    # estimate, and an under-estimate is the costly direction (it ships a
    # silently overlapping cover). Over-warning only asks for a shorter title.
    # Both values are pinned by `tests/test_cover_fit_calibration.py`, which
    # holds the observed (text, pt, box_width_in, fits?) samples they must keep
    # reproducing — including the proposal title `cover-and-agenda.md` says must
    # never reach this box.
    _LATIN_EM_WIDTH = 0.60
    _COVER_FIT_MARGIN = 0.88

    @classmethod
    def _em_width(cls, ch: str) -> float:
        """Width of one character in em. The single definition of that rule."""
        return 1.0 if ord(ch) > 0x2E7F else cls._LATIN_EM_WIDTH

    def _fill_cover_slide(self, prs: Presentation, profile: ProjectProfile) -> None:
        """Overwrite the cover's title / client / date from `profile.cover`.

        Runs before the role loop, while slide 1 is still at index 0 (nothing
        has been inserted yet). Text is written in place with
        `_set_text_preserving_format` — no font normalization, no auto-fit, no
        truncation, and no geometry changes — so every box keeps the size,
        weight, color and position the template authored. A field left empty
        in the profile keeps the template's placeholder wording rather than
        blanking the box.

        Cover text is expected to fit its box on one line; the boxes are fixed
        and are deliberately not resized. Text that would wrap is still written
        as given, with a warning, since only the author can decide how to
        shorten it.
        """
        from lib.templates.svn import COVER_JP_HONORIFIC_IDX, COVER_SLIDE_LAYOUT

        if not len(prs.slides):
            return
        shapes = prs.slides[0].shapes
        for field_name, idx in COVER_SLIDE_LAYOUT.items():
            value = (getattr(profile.cover, field_name, '') or '').strip()
            if not value or idx >= len(shapes):
                continue
            shape = shapes[idx]
            if not shape.has_text_frame:
                continue
            self._set_text_preserving_format(shape.text_frame, value)
            print(f"  - Cover: {field_name} -> {value!r}")
            if not self._cover_text_fits(shape, value):
                print(f'  - Cover: WARNING {field_name} is too long for its box '
                      f'and will wrap over the line below — shorten it in the profile')

        # LAST, because removing a shape renumbers everything after it and
        # COVER_SLIDE_LAYOUT addresses the date box by the index it holds while
        # the honorific line is still there.
        if not profile.is_japanese_source and COVER_JP_HONORIFIC_IDX < len(shapes):
            self._remove_shape(shapes[COVER_JP_HONORIFIC_IDX])
            print('  - Cover: dropped the JP honorific line — the source '
                  'document is not Japanese')

    def _apply_slide_text_overrides(self, prs: Presentation) -> None:
        """Correct template text on slides re-filed into another chapter.

        See `SLIDE_TEXT_OVERRIDES` in svn.py. Written with
        `_set_text_preserving_format`, so the corrected line keeps the size,
        weight and colour the template gave it. Runs before any filling, while
        every slide is still at its template position.
        """
        from lib.templates.svn import SLIDE_TEXT_OVERRIDES

        declared, matched = set(), set()
        for slide_no, replacements in SLIDE_TEXT_OVERRIDES.items():
            declared |= {(slide_no, text) for text in replacements}
            if slide_no > len(prs.slides._sldIdLst):
                continue
            for shape in self._walk_shapes(prs.slides[slide_no - 1].shapes):
                if not shape.has_text_frame:
                    continue
                key = shape.text_frame.text.strip()
                replacement = replacements.get(key)
                if replacement is None:
                    continue
                matched.add((slide_no, key))
                self._set_text_preserving_format(shape.text_frame, replacement)
                print(f'  - Slide {slide_no}: text override -> {replacement!r}')
        self._report_unmatched('SLIDE_TEXT_OVERRIDES', declared, matched)

    # A content slide is headed by a breadcrumb, `<chapter>  ｜  <section>`, and
    # repeats the section half below it as the slide's title. Matching the
    # separator rather than a per-slide shape index means one rule covers all 17
    # configured slides and survives the template moving a box around.
    _BREADCRUMB = re.compile(r'^(?P<chapter>[^｜]*)(?P<sep>\s*｜\s*)(?P<section>.+)$')

    def _retitle_slides_from_headings(self, prs: Presentation,
                                      profile: ProjectProfile) -> None:
        """Head each content slide with the chapter title its SOURCE gave it.

        The template's headings are the ones its original proposal used —
        `プロジェクト背景`, `機能一覧` — which name the right subject only by
        luck and, on a document written in English or Vietnamese, in the wrong
        language as well. The document that is actually being rendered already
        states its own chapter names, so those win.

        Only the section half of the breadcrumb is touched. The chapter half is
        the divider's name, which `_retitle_dividers_from_agenda` owns from the
        AGENDA — rewriting it here would let two passes disagree.

        A section whose heading was never recorded (an agent-authored JSON
        profile has no source headings) leaves its slide exactly as the
        template wrote it.
        """
        from lib.templates.svn import SECTION_TO_SLIDES

        for key, slide_numbers in SECTION_TO_SLIDES.items():
            heading = (profile.section_headings.get(key) or '').strip()
            if not heading:
                continue
            for slide_no in slide_numbers:
                if slide_no > len(prs.slides._sldIdLst):
                    continue
                self._retitle_slide(prs.slides[slide_no - 1], key, heading, slide_no)

    # Longest heading the template writes. Above it, text is body copy: a
    # paragraph naming its own chapter must not be mistaken for the title.
    _TITLE_MAX_CHARS = 40

    def _retitle_slide(self, slide, key: str, heading: str, slide_no: int) -> None:
        """Rewrite one slide's breadcrumb tail and the title that repeats it."""
        shapes = [shape for shape in self._walk_shapes(slide.shapes)
                  if shape.has_text_frame]
        crumbs = [(shape, match) for shape in shapes
                  for match in [self._BREADCRUMB.match(shape.text_frame.text.strip())]
                  if match]
        if not crumbs:
            return
        old = crumbs[0][1].group('section').strip()
        for shape, match in crumbs:
            self._set_text_preserving_format(
                shape.text_frame,
                f"{match.group('chapter')}{match.group('sep')}{heading}")
        # After the breadcrumbs are rewritten none of them still reads `old`,
        # so this reaches only the standalone title boxes.
        for shape in shapes:
            if self._is_title_of(shape.text_frame.text.strip(), key, old):
                self._set_text_preserving_format(shape.text_frame, heading)
        print(f'  - Slide {slide_no} retitled from the source heading: '
              f'{old!r} -> {heading!r}')

    @classmethod
    def _is_title_of(cls, text: str, key: str, breadcrumb_tail: str) -> bool:
        """Is *text* this slide's title — the breadcrumb tail repeated below it?

        Usually it is that string exactly. Not always: slide 5's breadcrumb says
        `機能一覧` where its title says `機能一覧表`. So a name that resolves to
        the slide's OWN section counts too — same chapter, spelled the template's
        other way.

        Resolution is what makes that safe, and only EXACT resolution does:
        `機能一覧表サンプル` (the caption on the sample grid) names no section,
        and slide 34's leftover `非機能要件` names `nfr_overview`, which is not
        slide 34's section — neither is a title, and neither is touched. The
        keyword tier is deliberately off here: it reads `費用の概要です。` as the
        cost chapter, and this method's answer is a licence to OVERWRITE, so it
        must never be given for a sentence that merely mentions the subject.

        A breadcrumb is never a title either, whatever it resolves to: it names
        the chapter as well as the section, and replacing it with the section
        alone loses the chapter.
        """
        from lib.section_aliases import resolve_section_key

        if not text or len(text) > cls._TITLE_MAX_CHARS:
            return False
        if cls._BREADCRUMB.match(text):
            return False
        if breadcrumb_tail and text == breadcrumb_tail:
            return True
        return resolve_section_key(text, keywords=False) == key

    def _translate_template_chrome(self, prs: Presentation,
                                   profile: ProjectProfile) -> None:
        """English the template's fixed Japanese labels, on a non-Japanese deck.

        Runs LAST of the text passes, so it only ever sees what the earlier ones
        left behind: a label no source heading names and no role fills. See
        `TEMPLATE_CHROME_EN` in svn.py for what belongs in that table and what
        deliberately does not.

        A Japanese proposal renders exactly as before — the table is never
        consulted.
        """
        from lib.templates.svn import TEMPLATE_CHROME_EN

        if profile.is_japanese_source:
            return
        translated = 0
        for slide in prs.slides:
            for shape in self._walk_shapes(slide.shapes):
                if not shape.has_text_frame:
                    continue
                replacement = TEMPLATE_CHROME_EN.get(shape.text_frame.text.strip())
                if not replacement:
                    continue
                self._set_text_preserving_format(shape.text_frame, replacement)
                translated += 1
        if translated:
            print(f'  - Englished {translated} fixed template label(s) — the '
                  f'source document is not Japanese')

    # A chapter divider is authored as two lines: the English name, then the
    # Japanese one. Constant across all nine dividers in this template.
    _DIVIDER_EN_IDX, _DIVIDER_JP_IDX = 0, 1

    def _section_present(self, profile: ProjectProfile, key: str) -> bool:
        """`section_has_content`, plus a divider-only section whose body arrives
        as anchored extra slides (`--extra-slides` with `anchor_section`)."""
        from lib.templates.svn import DIVIDER_ONLY_SECTIONS
        if section_has_content(profile, key):
            return True
        return key in DIVIDER_ONLY_SECTIONS and key in getattr(self, '_extra_anchor_sections', set())

    @classmethod
    def _breadcrumb_with_chapter(cls, match, chapter: str) -> str:
        """*match*'s breadcrumb with its chapter half replaced by *chapter*.

        `_BREADCRUMB`'s chapter group is greedy, so it holds the spacing in
        front of the separator as well as the name. Carrying that spacing over
        is what makes the rewritten breadcrumb sit exactly where the template's
        did — dropping it prints `System Overview｜  Feature List`.
        """
        was = match.group('chapter')
        lead = was[len(was.rstrip()):]
        return f"{chapter}{lead}{match.group('sep')}{match.group('section')}"

    def _retitle_dividers_from_sections(self, prs: Presentation,
                                        profile: ProjectProfile) -> None:
        """Name a divider after the SOURCE chapter that reaches it.

        Two chapters the template has no divider for by that name: a container
        (`5. 開発スケジュールおよび成果物` holding schedule + deliverables, which
        would otherwise show the template's 実施スケジュール divider and a second
        想定成果物 one) and a divider-only section (`team` on slide 17, authored
        as Sun*のスコープ). The JP line takes the chapter's name, the EN line
        its gloss when the heading or the AGENDA carries one, and the chapter
        half of every content-slide breadcrumb in that chapter follows.
        Runs pre-fill, so template slide numbers still address the deck.
        """
        from lib.chapter_dividers import chapter_name_parts, divider_titles
        from lib.templates.svn import (CHAPTER_DIVIDERS, DIVIDER_ONLY_SECTIONS,
                                       DIVIDER_ONLY_TITLES, SECTION_TO_SLIDES)

        slide_to_divider = {c: d for d, cs in CHAPTER_DIVIDERS.items() for c in cs}
        pairs = [self._split_agenda_item(item) for item in profile.agenda.chapters]
        for divider, heading, keys in divider_titles(
                profile, slide_to_divider, SECTION_TO_SLIDES, DIVIDER_ONLY_SECTIONS,
                self._section_present, DIVIDER_ONLY_TITLES):
            jp, en = chapter_name_parts(heading)
            label, gloss = self._agenda_match(heading, pairs)
            jp, en = label or jp, gloss or en
            if divider > len(prs.slides._sldIdLst) or not jp:
                continue
            shapes = prs.slides[divider - 1].shapes
            if len(shapes) <= self._DIVIDER_JP_IDX:
                continue
            self._set_text_preserving_format(shapes[self._DIVIDER_JP_IDX].text_frame, jp)
            if en:
                self._set_text_preserving_format(shapes[self._DIVIDER_EN_IDX].text_frame, en)
            chapter_half = en or jp
            for key in keys:
                for slide_no in SECTION_TO_SLIDES.get(key, []):
                    if slide_no > len(prs.slides._sldIdLst):
                        continue
                    for shape in self._walk_shapes(prs.slides[slide_no - 1].shapes):
                        if not shape.has_text_frame:
                            continue
                        match = self._BREADCRUMB.match(shape.text_frame.text.strip())
                        if match:
                            self._set_text_preserving_format(
                                shape.text_frame,
                                self._breadcrumb_with_chapter(match, chapter_half))
            print(f'  - Divider {divider} named from the source chapter: '
                  f'{en or "—"} / {jp} (sections: {keys})')

    def _retitle_dividers_from_agenda(self, prs: Presentation,
                                      profile: ProjectProfile) -> None:
        """Name each chapter divider the way the source document's AGENDA does.

        The template's dividers carry the wording of the proposal they were
        authored for, which drifts from the document actually being rendered:
        slide 37 says `The Expected Output / 想定作成物` where this proposal's
        AGENDA says `想定成果物 — Expected Deliverables`. The AGENDA is the
        contents page the reader just met on slide 2, so a divider disagreeing
        with it reads as two different documents stapled together.

        Matched on the divider's Japanese line, which is the stable half — the
        English is a gloss and varies most. A divider matching no agenda line is
        left exactly as the template wrote it, except on a non-Japanese deck,
        where its Japanese line falls back to its own English one rather than
        standing as the only Japanese left in the deck.
        """
        from lib.templates.svn import CHAPTER_DIVIDER_SLIDES

        if not profile.agenda.chapters:
            return
        pairs = [self._split_agenda_item(item) for item in profile.agenda.chapters]
        for slide_no in CHAPTER_DIVIDER_SLIDES:
            if slide_no > len(prs.slides._sldIdLst):
                continue
            shapes = prs.slides[slide_no - 1].shapes
            if len(shapes) <= self._DIVIDER_JP_IDX:
                continue
            jp_shape = shapes[self._DIVIDER_JP_IDX]
            en_shape = shapes[self._DIVIDER_EN_IDX]
            if not (jp_shape.has_text_frame and en_shape.has_text_frame):
                continue
            en_text = en_shape.text_frame.text.strip()
            # Structure first, wording second. This divider introduces known
            # content slides, those slides belong to a section, and that section
            # came out of one chapter of the document — a chain that cannot be
            # wrong. Matching the template's own words against the AGENDA can:
            # divider 27 reads `System Architecture`, the AGENDA lists a chapter
            # of that name, and the two matched — while the slides under it were
            # filled by a DIFFERENT chapter, `5. Estimation Assumptions`.
            chapter = self._chapter_of_divider(slide_no, profile)
            label, gloss = self._agenda_match(chapter, pairs) if chapter else ('', '')
            if not label:
                label, gloss = self._agenda_match(jp_shape.text_frame.text.strip(), pairs)
            if not label:
                # An English or Vietnamese agenda has no Japanese to match on,
                # so the divider's own English line is the only handle left.
                label, gloss = self._agenda_match(en_text, pairs)
            if not label and chapter:
                # The AGENDA does not list this chapter under any wording, and
                # neither of the divider's own lines matched either. The
                # chapter heading is then the only name that came from the
                # document rather than from the template.
                label, gloss = chapter, ''
            if not label and not profile.is_japanese_source:
                # Still nothing: this chapter is worded differently in the
                # AGENDA than on the template's divider. The name is then
                # unknowable, but the LANGUAGE is not — copying the divider's
                # English line over its Japanese one beats leaving a lone
                # Japanese heading in a deck that dropped Japanese everywhere
                # else. A Japanese deck keeps both lines, as authored.
                label, gloss = en_text, ''
            if not label:
                continue
            if label != jp_shape.text_frame.text.strip():
                self._set_text_preserving_format(jp_shape.text_frame, label)
            # A Japanese deck keeps its two distinct lines. Everywhere else the
            # divider is one name written twice, and leaving the English line
            # alone kept the TEMPLATE's chapter name visible right above the
            # document's — two names for one chapter, on one slide.
            english = gloss or ('' if profile.is_japanese_source else label)
            if english and english != en_shape.text_frame.text.strip():
                self._set_text_preserving_format(en_shape.text_frame, english)
            self._rename_chapter_in_breadcrumbs(prs, slide_no, english or label)
            print(f'  - Divider {slide_no} named from the AGENDA: '
                  f'{gloss or "—"} / {label}')

    def _rename_chapter_in_breadcrumbs(self, prs: Presentation, divider: int,
                                       chapter: str) -> None:
        """Carry a divider's new name onto the breadcrumbs of its own slides.

        Every content slide heads itself `<chapter>  ｜  <section>`.
        `_retitle_slides_from_headings` deliberately rewrites only the section
        half and leaves the chapter half to this pass — but this pass used to
        rename the divider alone, so a deck ended up with a divider reading
        `Benefits of the Proposed Solution` and, two slides later, a breadcrumb
        still reading `Advantage of system introduction`. One chapter, two
        names, and the reader has no way to tell they are the same one.
        """
        from lib.templates.svn import CHAPTER_DIVIDERS

        for slide_no in CHAPTER_DIVIDERS.get(divider) or []:
            if slide_no > len(prs.slides._sldIdLst):
                continue
            for shape in self._walk_shapes(prs.slides[slide_no - 1].shapes):
                if not shape.has_text_frame:
                    continue
                match = self._BREADCRUMB.match(shape.text_frame.text.strip())
                if not match or match.group('chapter').strip() == chapter:
                    continue
                self._set_text_preserving_format(
                    shape.text_frame, self._breadcrumb_with_chapter(match, chapter))

    @staticmethod
    def _chapter_of_divider(divider: int, profile: ProjectProfile) -> str:
        """The document chapter this divider introduces, or ''.

        Walks the two structural maps the template already declares —
        `CHAPTER_DIVIDERS` (divider → the content slides it announces) and
        `SECTION_TO_SLIDES` (section → its slides) — to reach the section this
        divider covers, then reads that section's chapter out of the profile.

        This is the connection `_agenda_match` was standing in for. A proposal
        titled `2. Benefits of the Proposed Solution` against a template divider
        reading `Advantage of system introduction` shares no words at all, and
        the old fallback printed the TEMPLATE's chapter name into a deck whose
        AGENDA said something else. The structure knows; the strings never did.
        """
        from lib.templates.svn import CHAPTER_DIVIDERS, SECTION_TO_SLIDES

        content_slides = set(CHAPTER_DIVIDERS.get(divider) or [])
        if not content_slides:
            return ''
        # In the document's own order, so a divider covering two sections is
        # named after whichever chapter the document introduced first.
        for key in profile.section_order:
            if key in SECTION_TO_SLIDES and content_slides.intersection(
                    SECTION_TO_SLIDES[key]):
                chapter = profile.section_chapters.get(key)
                if chapter:
                    return chapter
        return ''

    def _apply_slide_texts_from_profile(self, prs: Presentation,
                                        profile: ProjectProfile) -> None:
        """Replace template wording with the source document's own.

        See `SLIDE_TEXT_FROM_PROFILE` in svn.py. A path resolving to an empty
        value leaves the template's line alone — better a heading written for
        another proposal than a blank one.
        """
        from lib.templates.svn import SLIDE_TEXT_FROM_PROFILE

        declared, matched = set(), set()
        for slide_no, sources in SLIDE_TEXT_FROM_PROFILE.items():
            declared |= {(slide_no, text) for text in sources}
            if slide_no > len(prs.slides._sldIdLst):
                continue
            for shape in self._walk_shapes(prs.slides[slide_no - 1].shapes):
                if not shape.has_text_frame:
                    continue
                key = shape.text_frame.text.strip()
                path = sources.get(key)
                if not path:
                    continue
                # Found in the template — the contract holds. Whether the
                # profile has a value for it is a separate question.
                matched.add((slide_no, key))
                value = str(self._resolve_source_path(profile, path) or '').strip()
                if not value:
                    continue
                self._set_text_preserving_format(shape.text_frame, value)
                print(f'  - Slide {slide_no}: {path} -> {value[:48]!r}')
        self._report_unmatched('SLIDE_TEXT_FROM_PROFILE', declared, matched)

    @classmethod
    def _walk_shapes(cls, shapes):
        """Yield every shape, descending into groups.

        The template keeps a slide's title and breadcrumb inside a group on some
        slides and at the top level on others, so anything addressing them by
        content has to look in both places.
        """
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from cls._walk_shapes(shape.shapes)
            else:
                yield shape

    @classmethod
    def _removable_text_shapes(cls, slide) -> list:
        """Every text shape on *slide*, groups included, as a stable list.

        Nested as well as top-level: the sample lead-ins on slides 33 and 36 sit
        inside groups, and a top-level-only sweep left exactly those behind.
        Materialized before anything is deleted — removing from the tree while
        walking it skips the next sibling.
        """
        return [shape for shape in cls._walk_shapes(slide.shapes)
                if shape.has_text_frame]

    @staticmethod
    def _normalize_note_text(text: str) -> str:
        """Whitespace-normalized form used to match template note text.

        Line-level strip + blank-line drop, so a stray trailing space or an
        empty paragraph in the template doesn't defeat the match.
        """
        lines = [line.strip() for line in (text or '').splitlines()]
        return '\n'.join(line for line in lines if line)

    # Separator between a chapter label and its gloss in a profile agenda item.
    # The source proposals write `システムの概要 — System Overview`; en dash and
    # hyphen are accepted too, since which one a document uses is arbitrary.
    _AGENDA_DASH = re.compile(r'\s+[—–-]\s+')

    # Agenda lines are left-aligned in a box with small internal margins, so
    # they can use nearly its full width before wrapping onto the line below.
    _AGENDA_FIT_MARGIN = 0.95

    def _fill_agenda_slide(self, prs: Presentation, profile: ProjectProfile) -> None:
        """Rewrite slide 2's agenda from `profile.agenda`.

        The template ships each agenda line as `<JP label>: <Vietnamese note to
        the proposal writer>` — authoring guidance, never deck content. The
        source proposal's own AGENDA block is what the deck should say, so each
        mapped box is rewritten from the profile's lines, in the profile's order.

        What happens to a box the profile has NO lines for depends on whether
        the profile has an agenda at all:

        - **It does** (this is the supplied-document path): the box is emptied.
          The source's agenda is the whole truth about this slide, so a box it
          says nothing about must show nothing. Leaving the template's own
          `開発費用 / インフラコスト` there would put two lines on the deck that
          the document never wrote — right for this proposal by luck, wrong for
          the next one.
        - **It doesn't** (an older profile with no `## Agenda`): the template's
          labels stay and are only glossed to English (`_gloss_agenda_box`), so
          such a profile still never ships the Vietnamese notes.

        Runs right after the cover, while slide 2 is still at index 1 (the fill
        pass hasn't inserted any overflow slides yet).
        """
        from lib.templates.svn import AGENDA_SLIDE_LAYOUT

        if len(prs.slides) < 2:
            return
        agenda_supplied = bool(profile.agenda.chapters or profile.agenda.cost_breakdown)
        shapes = prs.slides[1].shapes
        for field_name, idx in AGENDA_SLIDE_LAYOUT.items():
            if idx >= len(shapes) or not shapes[idx].has_text_frame:
                continue
            items = [item.strip()
                     for item in (getattr(profile.agenda, field_name, None) or [])
                     if item.strip()]
            if items:
                self._write_agenda_box(shapes[idx], items, field_name)
            elif agenda_supplied:
                # Emptied, not deleted: the boxes are addressed by index, so
                # removing a shape would shift the one still to be written.
                shapes[idx].text_frame.clear()
                print(f'  - Agenda: {field_name} not in the profile — box emptied '
                      f'rather than left showing the template\'s own lines')
            else:
                self._gloss_agenda_box(shapes[idx])

    def _write_agenda_box(self, shape, items: List[str], field_name: str) -> None:
        """Rewrite one agenda box's lines from `items`, one line per item.

        Each paragraph is written in place, so it keeps the template's bullet,
        indent, spacing and its two-tone runs (bold black label + grey 9pt
        gloss). Surplus template lines are deleted — leaving them would ship a
        chapter the proposal doesn't have, still carrying its Vietnamese note.
        A profile with more lines than the template clones the last paragraph,
        which grows the box's content downward: allowed, but warned about,
        since only the author can decide what to cut.
        """
        text_frame = shape.text_frame
        template_lines = len(text_frame.paragraphs)
        for _ in range(len(items) - template_lines):
            last = text_frame.paragraphs[-1]._p
            last.getparent().append(copy.deepcopy(last))

        print(f'  - Agenda: {field_name} <- {len(items)} line(s) from the profile')
        box_width_in = (shape.width or 0) / self._EMU_PER_INCH
        paragraphs = text_frame.paragraphs
        for para, item in zip(paragraphs, items):
            label, gloss = self._split_agenda_item(item)
            self._write_agenda_line(para, label, gloss)
            if box_width_in and not self._agenda_line_fits(para, box_width_in):
                print(f'  - Agenda: WARNING {field_name} line {item!r} is too '
                      f'long for its box and will wrap — shorten it in the profile')
        for leftover in list(paragraphs[len(items):]):
            leftover._p.getparent().remove(leftover._p)

        if len(items) > template_lines:
            # Verified by rendering: the 10th chapter line already collides with
            # the 費用 sub-item box below, and past ~11 the box clips the rest.
            print(f'  - Agenda: WARNING {field_name} has {len(items)} lines but the '
                  f'template box holds {template_lines} — the extra line(s) run '
                  f'past it and collide with what sits below')

    @classmethod
    def _split_agenda_item(cls, item: str) -> Tuple[str, str]:
        """Split `<label> — <gloss>` into its two parts (gloss may be empty)."""
        parts = cls._AGENDA_DASH.split(item, maxsplit=1)
        return parts[0].strip(), (parts[1].strip() if len(parts) > 1 else '')

    def _write_agenda_line(self, para, label: str, gloss: str) -> None:
        """Write one `<label> — <gloss>` line into an existing paragraph.

        The template builds each line from a label run (10pt bold black), a
        `:` separator and a note run (9pt grey). Reusing them keeps that
        two-tone look. The note is taken as the LAST run rather than the third:
        a label can be split across several runs of its own (slide 2's
        'Sun*' + 'のスコープ'), and writing the gloss into the third run there
        would set it in the label's bold black. The separator is restyled from
        the note run so the dash doesn't set in bold black either.
        """
        runs = para.runs
        if not runs:
            return
        runs[0].text = label
        if not gloss:
            self._drop_runs(runs[1:])
            return
        if len(runs) == 1:
            # A template revision merged the runs — one flat line is better
            # than dropping the gloss.
            runs[0].text = f'{label} — {gloss}'
            return
        note_run = runs[-1]
        if len(runs) >= 3:
            separator = runs[1]
            self._copy_run_style(note_run, separator)
            separator.text = ' — '
            note_run.text = gloss
            self._drop_runs(runs[2:-1])
        else:
            note_run.text = f' — {gloss}'

    @staticmethod
    def _drop_runs(runs) -> None:
        """Delete the given runs from their paragraph."""
        for run in list(runs):
            run._r.getparent().remove(run._r)

    @staticmethod
    def _copy_run_style(src, dst) -> None:
        """Give run `dst` run `src`'s character formatting (size/weight/color).

        Copies the whole `a:rPr` element rather than the handful of properties
        python-pptx's font API exposes, so theme colors and latin/east-asian
        typeface pairs survive too.
        """
        from pptx.oxml.ns import qn

        src_rPr = src._r.find(qn('a:rPr'))
        if src_rPr is None:
            return
        dst_rPr = dst._r.find(qn('a:rPr'))
        if dst_rPr is not None:
            dst._r.remove(dst_rPr)
        dst._r.insert(0, copy.deepcopy(src_rPr))

    @classmethod
    def _agenda_line_fits(cls, para, box_width_in: float) -> bool:
        """Estimate whether a written agenda paragraph stays on one line.

        Same approximate width model as the cover's check, summed per run
        because a line mixes a 10pt label with a 9pt gloss. Drives a warning
        only — never a change to the deck.
        """
        width = sum(cls._text_width_in(run.text, run.font.size.pt if run.font.size else 10)
                    for run in para.runs)
        return width <= box_width_in * cls._AGENDA_FIT_MARGIN

    def _gloss_agenda_box(self, shape) -> None:
        """Fallback for a box the profile has no lines for: replace each
        template line's Vietnamese note with its English gloss.

        A line is only touched when the text before its first colon is a known
        chapter label — anything else (the 'AGENDA' heading, a label a future
        template revision renames) is left exactly as authored.
        """
        from lib.templates.svn import AGENDA_GLOSSES

        rewritten = 0
        for para in shape.text_frame.paragraphs:
            if not para.runs or ':' not in para.text:
                continue
            label = para.text.split(':', 1)[0].strip()
            english = AGENDA_GLOSSES.get(label)
            if not english:
                continue
            self._write_agenda_line(para, label, english)
            rewritten += 1
        if rewritten:
            print(f'  - Agenda: no profile lines — glossed {rewritten} template '
                  f'line(s) to English')

    def _strip_template_notes(self, prs: Presentation) -> int:
        """Delete every shape holding text the TEMPLATE wrote and nothing replaced.

        Two kinds, listed separately in svn.py because they are separate
        mistakes to ship:

        * `TEMPLATE_NOTE_TEXTS` — the template's notes to the proposal writer
          (Vietnamese "what goes on this slide" blurbs) and its editors' own JP
          revision notes. Never deck content.
        * `TEMPLATE_SAMPLE_TEXTS` — the SAMPLE COPY of the proposal this deck
          was cloned from: the lead-in above the infrastructure table, the
          claim that the client supplied an NFR document, the sample Gantt's
          own phase blocks. No role writes these shapes, so on any other deck
          they are another client's sentences — in Japanese, under an English
          heading, and this pass is what stops them shipping.

        Both match on exact (whitespace-normalized) text rather than index or
        position — see svn.py. Running LAST is what makes the sample case work:
        a shape a profile field replaced (slide 56's 諸経費 note) no longer
        matches, so supplying the content still wins over deleting the sample.

        Returns the number of shapes removed.
        """
        from lib.templates.svn import TEMPLATE_NOTE_TEXTS, TEMPLATE_SAMPLE_TEXTS

        removed = 0
        for slide_no, slide in enumerate(prs.slides, 1):
            for shape in self._removable_text_shapes(slide):
                text = self._normalize_note_text(shape.text_frame.text)
                if text in TEMPLATE_NOTE_TEXTS:
                    kind = 'note'
                elif text in TEMPLATE_SAMPLE_TEXTS:
                    kind = 'sample copy'
                else:
                    continue
                shape._element.getparent().remove(shape._element)
                removed += 1
                print(f'  - Removed template {kind} shape from slide {slide_no}')
        if removed:
            print(f'  - Stripped {removed} template note/sample shape(s)')
        # Deliberately NOT reported through `_report_unmatched`: this pass runs
        # last, after prune has deleted the slides a section had no content for,
        # so most entries legitimately match nothing on any given deck. Coverage
        # for this map is checked against the FULL template by
        # `scripts/verify-template-contract.py`, which is the only place the
        # question is answerable. A warning that fires on every healthy run
        # teaches people to ignore warnings.
        return removed

    @staticmethod
    def _report_unmatched(map_name: str, declared: set, matched: set) -> None:
        """Warn about mapping entries that found nothing in this template.

        These maps address the template by exact string. That precision is
        deliberate — see svn.py, where a positional or diacritic rule would
        delete real content — but it means a template edit silently disables an
        entry rather than erroring. An unmatched entry is therefore always a
        defect: either the template changed under the map, or the entry was
        wrong when it was written.
        """
        missing = declared - matched
        if not missing:
            return
        shown = sorted(str(item) for item in missing)[:3]
        tail = ', …' if len(missing) > 3 else ''
        print(f'  - WARNING: {map_name} has {len(missing)} entry/entries that '
              f'matched nothing in this template — it may have changed under '
              f'the map: {shown}{tail}')

    def _cover_text_fits(self, shape, value: str) -> bool:
        """Estimate whether `value` stays on one line inside `shape`.

        Approximate by design (no real font metrics available here), so it
        only ever drives a warning — never a change to the deck.
        """
        font_pt = next((run.font.size.pt
                        for para in shape.text_frame.paragraphs
                        for run in para.runs if run.font.size), None)
        if not font_pt or not shape.width:
            return True
        return self._cover_fits(value, font_pt, shape.width / self._EMU_PER_INCH)

    @classmethod
    def _cover_fits(cls, value: str, font_pt: float, box_width_in: float) -> bool:
        """Pure width check, split out so the calibration test can exercise it."""
        return cls._text_width_in(value, font_pt) <= box_width_in * cls._COVER_FIT_MARGIN

    @classmethod
    def _text_width_in(cls, text: str, font_pt: float) -> float:
        """Approximate rendered width of one run, in inches."""
        return sum(cls._em_width(ch) for ch in text) * font_pt / 72

    # ------------------------------------------------------------------
    # Prune + reorder pass — drop template slides not backed by content,
    # keep chapter dividers only when their chapter has content, and order
    # the survivors by `profile.section_order`.
    # ------------------------------------------------------------------

    def _prune_and_reorder(
        self, prs: Presentation, profile: ProjectProfile, spans: Dict[int, Tuple[int, int]],
    ) -> Dict[int, int]:
        """Delete unwanted template slides and reorder survivors.

        Returns ``new_last_positions``: {template_slide_num: new_0based_index}
        for every configured slide that survived — the fresh position map
        ``_render_extra_slides`` (and its chrome/canvas pickers) must use,
        since positions computed before this pass are stale afterward.

        No-ops (returns positions unchanged) when there's no recognized
        template config — this pass hardcodes SVN's slide-picking semantics
        (``CHAPTER_DIVIDERS`` etc.), so it must never run against a template
        `_detect_template_config` didn't recognize, or it would destructively
        delete slides using slide numbers that mean nothing for that deck.
        """
        if not self.template_config:
            return {n: e for n, (s, e) in spans.items()}

        from lib.templates.svn import ALWAYS_KEEP_SLIDES, CHAPTER_DIVIDERS, SECTION_TO_SLIDES

        pml_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        r_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        r_id_attr = f'{{{r_ns}}}id'

        # Reverse lookup: content slide number -> the chapter divider that introduces it.
        slide_to_divider = {c: d for d, cs in CHAPTER_DIVIDERS.items() for c in cs}

        def _current_index(orig: int) -> int:
            """Current (post-fill, pre-prune) 0-based index of an un-configured
            slide number (cover/agenda/divider), accounting for overflow slides
            inserted before it by earlier configured sections."""
            return self._current_slide_index(orig, spans)

        # --- (a)+(b) build wanted_orig: always-kept + content sections' slides,
        #     strictly in `profile.section_order` (i.e. markdown heading order).
        #     The user owns the presentation order they write in the markdown —
        #     sections are never reshuffled into template/canonical order. Each
        #     chapter divider is inserted once, right before the first sibling
        #     of that chapter to appear in the markdown; later-appearing
        #     siblings of the same chapter are emitted in place (no duplicate
        #     divider) since the template only has one divider slide per chapter.
        order_source = list(profile.section_order or SECTION_TO_SLIDES.keys())
        # A divider-only section fed purely by anchored extras has no markdown
        # chapter, so it is missing from section_order — give it the last spot.
        from lib.templates.svn import DIVIDER_ONLY_SECTIONS as _DIVIDER_ONLY
        for only_key in _DIVIDER_ONLY:
            if only_key not in order_source and self._section_present(profile, only_key):
                order_source.append(only_key)
        # Siblings lifted out of one container chapter share the first
        # sibling's divider; a divider-only section (team) contributes its
        # divider and nothing else — see `chapter_dividers.py`.
        from lib.chapter_dividers import chapter_divider_map
        from lib.templates.svn import DIVIDER_ONLY_SECTIONS
        chapter_dividers = chapter_divider_map(
            profile, slide_to_divider, SECTION_TO_SLIDES, DIVIDER_ONLY_SECTIONS)

        wanted_orig: list = list(ALWAYS_KEEP_SLIDES)
        emitted_keys: set = set()
        dividers_emitted: set = set()
        for key in order_source:
            if key in emitted_keys:
                continue
            slide_nums = SECTION_TO_SLIDES.get(key)
            if not slide_nums and key not in DIVIDER_ONLY_SECTIONS:
                continue
            if not self._section_present(profile, key):
                continue
            divider = chapter_dividers.get(key)
            if divider is None:
                divider = slide_to_divider.get(slide_nums[0]) if slide_nums else None
            if divider is not None and divider not in dividers_emitted:
                wanted_orig.append(divider)
                dividers_emitted.add(divider)
            if slide_nums:
                wanted_orig.extend(self._filled_slides_of(key, slide_nums))
            emitted_keys.add(key)

        self._keep_dividers_for_extra_chapters(
            prs, profile, wanted_orig, dividers_emitted, SECTION_TO_SLIDES, spans)

        # --- (c) expand each wanted original slide number to its current-index
        #     block: configured slides expand to their full overflow block;
        #     everything else (cover/agenda/dividers) resolves to one index.
        wanted_indices_ordered: list = []
        seen_idx = set()
        for orig in wanted_orig:
            if orig in spans:
                start, end = spans[orig]
                idx_block = range(start, end + 1)
            else:
                idx_block = [_current_index(orig)]
            for idx in idx_block:
                if idx in seen_idx:
                    print(f'  - _prune_and_reorder: duplicate index {idx} '
                          f'(slide {orig}) — skipping')
                    continue
                wanted_indices_ordered.append(idx)
                seen_idx.add(idx)

        keep_set = set(wanted_indices_ordered)

        # --- (d) sldIdLst surgery: delete unwanted slides (drop_rel so no
        #     orphaned parts remain), reorder survivors to match final order.
        sldIdLst = prs.part._element.find(f'{{{pml_ns}}}sldIdLst')
        if sldIdLst is None:
            print('  - _prune_and_reorder: no <p:sldIdLst> found — skipping prune')
            return {n: e for n, (s, e) in spans.items()}

        snapshot = list(sldIdLst)
        final_elems = [snapshot[i] for i in wanted_indices_ordered]

        for child in list(sldIdLst):
            sldIdLst.remove(child)
        for elem in final_elems:
            sldIdLst.append(elem)

        dropped = 0
        for i, elem in enumerate(snapshot):
            if i in keep_set:
                continue
            rid = elem.get(r_id_attr)
            if rid:
                try:
                    prs.part.drop_rel(rid)
                    dropped += 1
                except KeyError:
                    pass

        print(f'  - _prune_and_reorder: kept {len(final_elems)}/{len(snapshot)} slides '
              f'(dropped {dropped} relationship(s)); kept template slides in order: '
              f'{wanted_orig}')

        # --- (e) fresh positions for surviving configured slides.
        new_last_positions: Dict[int, int] = {}
        # Where each surviving section STARTS, kept alongside where it ends.
        # An appendix slide belonging before its chapter's section needs the
        # first index, and a span's end says nothing about its beginning.
        self._first_positions = {}
        for n, (s, e) in spans.items():
            if e in keep_set:
                new_last_positions[n] = wanted_indices_ordered.index(e)
            if s in keep_set:
                self._first_positions[n] = wanted_indices_ordered.index(s)
        # Dividers too, so an appendix slide pinned to one (see
        # `_keep_dividers_for_extra_chapters`) resolves to it instead of falling
        # back to the end of the deck.
        for divider in dividers_emitted:
            idx = _current_index(divider)
            if idx in keep_set:
                new_last_positions.setdefault(divider,
                                              wanted_indices_ordered.index(idx))
        return new_last_positions

    def _fill_slide_roles(
        self, prs: Presentation, slide, role_config: SlideRoleConfig,
        profile: ProjectProfile, slide_index: int, discovery: ShapeDiscovery,
    ) -> int:
        """Fill all roles on a slide. Returns extra slides inserted (overflow continuations).

        TABLE roles keep their existing independent overflow — no slide in this
        template mixes a table with other overflow-prone roles. TEXT-kind roles
        (collections / compound title+body pairs) instead share ONE set of
        continuation slides per base slide, sized to the largest role's chunk
        count — so multiple overflowing roles on the same slide (e.g. business
        process's `category_labels` + `after_blocks`) land on the same
        continuation pages instead of each duplicating the slide independently,
        which would corrupt non-overflowing sibling roles on the resulting
        chain of duplicates.
        """
        if role_config.slide_number == 11:
            split_extra = self._maybe_split_business_process_slide(
                prs, slide, role_config, profile, slide_index, discovery)
            if split_extra is not None:
                return split_extra

        extra_total = 0
        text_chunks = []  # [(role, [chunk0, chunk1, ...]), ...]
        for role in role_config.roles:
            data = self._resolve_source_path(profile, role.source_path)
            # Record whether this slide received anything at all. A section that
            # owns several template slides (assumptions owns 23/24/25 via
            # `assumptions[0:5]`, `[5:8]`, `[8:9]`) leaves the surplus slides
            # unwritten when the profile has fewer items — and an unwritten
            # slide still carries the template's own SAMPLE content, which then
            # ships as if it were this proposal's. The prune pass drops them.
            if data:
                self._slides_filled[role_config.slide_number] = True
            self._slides_filled.setdefault(role_config.slide_number, False)
            shapes = discovery.discover(slide, role)
            if role.kind == ContentKind.TABLE:
                prose_roles = [other for other in role_config.roles
                               if other is not role and not other.sub_roles
                               and other.kind == ContentKind.TEXT]
                if self._is_labelled_table_set(data):
                    extra_total += self._fill_grouped_table_role(
                        prs, slide, slide_index + extra_total, shapes, data, role,
                        lead_in_roles=prose_roles)
                else:
                    extra_total += self._fill_table_role(prs, slide, slide_index + extra_total,
                                                        shapes, data, role,
                                                        lead_in_roles=prose_roles)
            elif role.kind == ContentKind.IMAGE:
                self._fill_image_role(slide, shapes, data, role, profile)
            elif role.sub_roles:  # compound (title+body pairs)
                items = data if isinstance(data, list) else []
                chunks = self._chunk_list(items, len(shapes))
                self._fill_compound_pairs(shapes, chunks[0], role, slide=slide)
                text_chunks.append((role, chunks))
            else:  # simple text / text collection
                items = self._coerce_text_items(data, role.item_field)
                chunks = self._chunk_list(items, len(shapes))
                self._fill_text_shapes(shapes, chunks[0], role, slide=slide)
                text_chunks.append((role, chunks))
        # Every role of this slide is written, so a font family that spans two of
        # them can finally be sized.
        self._flush_shared_fonts()

        max_chunks = max((len(chunks) for _, chunks in text_chunks), default=1)
        if max_chunks > 1:
            overflowing = [r.name for r, chunks in text_chunks if len(chunks) > 1]
            print(f'  - Text overflow: {max_chunks} slide(s) needed (role(s): {overflowing})')
            current_index = slide_index + extra_total
            for k in range(1, max_chunks):
                new_slide = self._duplicate_slide(prs, current_index)
                extra_total += 1
                current_index += 1
                for role, chunks in text_chunks:
                    # Roles with fewer chunks than max_chunks have nothing new to
                    # show on this continuation — clear them rather than repeating
                    # their last chunk verbatim, which reads as a mistake (the
                    # same items appearing twice) rather than intentional content.
                    chunk = chunks[k] if k < len(chunks) else []
                    new_shapes = ShapeDiscovery().discover(new_slide, role)
                    if role.sub_roles:
                        self._fill_compound_pairs(new_shapes, chunk, role, slide=new_slide)
                    else:
                        self._fill_text_shapes(new_shapes, chunk, role, slide=new_slide)
                self._flush_shared_fonts()
                print(f'  - Filled continuation {k}/{max_chunks - 1}')
        return extra_total

    @staticmethod
    def _chunk_list(items: list, capacity: int) -> list:
        """Split items into chunks of `capacity` size each (always >=1 chunk)."""
        if capacity <= 0:
            return [items]
        return [items[i:i + capacity] for i in range(0, len(items), capacity)] or [[]]

    def _maybe_split_business_process_slide(
        self, prs: Presentation, slide, role_config: SlideRoleConfig,
        profile: ProjectProfile, slide_index: int, discovery: ShapeDiscovery,
    ) -> Optional[int]:
        """Bespoke fallback for template slide 11 (business process, horizontal
        view): duplicating the whole slide for `after_blocks` overflow would
        repeat the unrelated `before_steps` row pointlessly on the
        continuation. Instead split into a before-only slide (trimmed to the
        real step count) and an after-only slide with the overflow items laid
        out as a second row, reusing the existing column shapes — closer to
        what a human editor would do by hand.

        Returns None (caller falls back to the generic per-role handling)
        when `after_blocks` isn't overflowing, or when it overflows by more
        than a second row can hold (2 rows of `capacity` each) — the generic
        duplicate-and-paginate path is the safe fallback for that case.
        """
        before_role = next((r for r in role_config.roles if r.name == 'before_steps'), None)
        after_role = next((r for r in role_config.roles if r.name == 'after_blocks'), None)
        if before_role is None or after_role is None:
            return None

        after_items = self._resolve_source_path(profile, after_role.source_path)
        after_items = after_items if isinstance(after_items, list) else []
        row1_pairs = discovery.discover(slide, after_role)
        capacity = len(row1_pairs)
        if capacity == 0 or not (capacity < len(after_items) <= capacity * 2):
            return None

        from lib.templates.svn import BUSINESS_PROCESS_SLIDE11_LAYOUT as L

        def shape_at(target_slide, idx):
            return ShapeDiscovery._safe_get_by_index(target_slide, idx)

        print(f'  - Role "after_blocks": {len(after_items)} items > {capacity} slots — '
              f'splitting into before-only + after-only (2 rows) slides')

        before_items = self._coerce_text_items(
            self._resolve_source_path(profile, before_role.source_path))
        before_boxes = discovery.discover(slide, before_role)
        n_before = len(before_items)
        n_row2 = len(after_items) - capacity

        # Resolve every shape reference we'll need — on BOTH the original
        # slide and its as-yet-uncreated duplicate — before any deletion.
        # Removing a shape shifts the live index of every shape after it, so
        # re-looking-up by index mid-deletion would silently grab the wrong
        # shape; every index lookup below happens against untouched slides.
        before_connectors = [shape_at(slide, i) for i in L['before_connectors']]
        before_bg = shape_at(slide, L['before_bg'])
        before_bg_top = before_bg.top
        after_bg = shape_at(slide, L['after_bg'])
        after_bg_top = after_bg.top
        after_label_top = shape_at(slide, L['after_label']).top
        after_card_bg_top = shape_at(slide, L['after_card_bg'][0]).top
        after_title_top = shape_at(slide, L['after_title_boxes'][0]).top
        after_body_top = shape_at(slide, L['after_body_boxes'][0]).top
        # Connectors sit left-to-right between consecutive columns; the raw
        # index order in the template isn't left-to-right, so sort by position.
        after_connectors_sorted = sorted(
            (shape_at(slide, i) for i in L['after_connectors']), key=lambda s: s.left)
        after_connector_top = after_connectors_sorted[0].top if after_connectors_sorted else after_body_top
        row2_label_top = before_bg_top + (after_label_top - after_bg_top)
        row2_card_bg_top = before_bg_top + (after_card_bg_top - after_bg_top)
        row2_title_top = before_bg_top + (after_title_top - after_bg_top)
        row2_body_top = before_bg_top + (after_body_top - after_bg_top)
        row2_connector_top = before_bg_top + (after_connector_top - after_bg_top)
        after_side_shapes = [shape_at(slide, idx) for idx in [
            L['after_bg'], L['after_label'], L['row_arrow'], *L['after_card_bg'],
            *L['after_title_boxes'], *L['after_body_boxes'], *L['after_connectors'],
        ]]

        # 1. Duplicate the pristine slide first — it becomes the after-only slide.
        after_slide = self._duplicate_slide(prs, slide_index)

        # Resolve the duplicate's own shape references too, before touching it.
        # `row_arrow` is kept (not deleted) on this slide: originally it showed
        # before-state flowing into after-state, and with 2 after-rows stacked
        # it now reads just as naturally as "row 2 flows into row 1" — same
        # gap between panels, no repositioning needed.
        before_side_shapes_on_dup = [shape_at(after_slide, idx) for idx in [
            L['before_label'], *L['before_boxes'], *L['before_connectors'],
        ]]
        before_bg_on_dup = shape_at(after_slide, L['before_bg'])
        after_label_on_dup = shape_at(after_slide, L['after_label'])
        after_card_bg_on_dup = [shape_at(after_slide, i) for i in L['after_card_bg']]
        after_connectors_on_dup = sorted(
            (shape_at(after_slide, i) for i in L['after_connectors']), key=lambda s: s.left)
        row1_after_pairs = discovery.discover(after_slide, after_role)

        # Drop before-side shapes and replace the (differently-styled)
        # before_bg panel with a clone of after_bg repositioned into its spot
        # — both rows share the same panel style instead of one row looking
        # like a leftover "before" panel. This MUST happen before row 2's
        # content is cloned below: `_clone_shape` appends to the end of the
        # shape tree, so cloning the panel afterwards would stack it on top
        # of (hiding) row 2's title/body/card_bg.
        for shp in before_side_shapes_on_dup:
            self._remove_shape(shp)
        self._remove_shape(before_bg_on_dup)
        self._clone_shape(after_slide, after_bg, top=before_bg_top)

        # Normalize row 1's own body cells to their largest height AND width
        # before filling/cloning — several of this template's 4 body cells are
        # much smaller than the others (as little as 0.13" tall / 1.21" wide),
        # which forces heavy truncation. Widening in place (left unchanged) is
        # safe here — checked against this exact template's column spacing,
        # the widened cells still clear the next column's left edge. Row 2
        # clones inherit the fix since it copies row 1 *after* this runs.
        row1_bodies = [b for _, b in row1_after_pairs if b is not None]
        if row1_bodies:
            max_body_h = max(b.height for b in row1_bodies)
            max_body_w = max(b.width for b in row1_bodies)
            for b in row1_bodies:
                b.height = max_body_h
                b.width = max_body_w

        row2_pairs = []
        for i in range(n_row2):
            title_src, body_src = row1_after_pairs[i]
            # Card backdrop cloned first so it stays behind title/body in z-order,
            # matching row 1's own stacking (card_bg shapes precede title/body in
            # the template's original shape order).
            self._clone_shape(after_slide, after_card_bg_on_dup[i], top=row2_card_bg_top)
            new_title = self._clone_shape(after_slide, title_src, top=row2_title_top)
            new_body = self._clone_shape(after_slide, body_src, top=row2_body_top)
            row2_pairs.append((new_title, new_body))
        for i in range(max(0, n_row2 - 1)):
            self._clone_shape(after_slide, after_connectors_on_dup[i], top=row2_connector_top)

        # 2. Original slide -> before-only: fill + trim + widen before_steps to
        #    fill the freed-up width instead of staying bunched at their
        #    original (8-box) narrow size, drop after-side shapes.
        self._fill_text_shapes(before_boxes, before_items, before_role)
        kept_boxes = before_boxes[:n_before]
        kept_connectors = before_connectors[:max(0, n_before - 1)]
        for shp in before_boxes[n_before:]:
            self._remove_shape(shp)
        for shp in before_connectors[max(0, n_before - 1):]:
            self._remove_shape(shp)
        if kept_boxes:
            span_left = before_boxes[0].left
            span_right = before_boxes[-1].left + before_boxes[-1].width
            gap = before_connectors[0].width if before_connectors else 0
            box_w = int((span_right - span_left - gap * (n_before - 1)) / n_before)
            x = span_left
            for i, box in enumerate(kept_boxes):
                box.left, box.width = int(x), box_w
                x += box_w
                if i < len(kept_connectors):
                    kept_connectors[i].left = int(x)
                    x += gap
        for shp in after_side_shapes:
            self._remove_shape(shp)

        # 3. Duplicated slide -> after-only, 2 rows: fill row 1 (existing
        #    columns) + row 2 (cloned columns from step 1), and move the
        #    single "after" label above both rows.
        if after_label_on_dup is not None:
            after_label_on_dup.top = row2_label_top
        self._fill_compound_pairs(row1_after_pairs, after_items[:capacity], after_role)
        self._fill_compound_pairs(row2_pairs, after_items[capacity:], after_role)

        return 1

    # ------------------------------------------------------------------
    # Source-path resolution: dot notation + optional [start:end] slice
    # ------------------------------------------------------------------

    def _resolve_source_path(self, profile: ProjectProfile, path: str):
        """Resolve `a.b.c` or `a[0:2]` or `a.b[2:4]` into a value on the profile.

        Returns None if any segment is missing.
        """
        if not path:
            return None
        # Split off trailing slice like benefits[0:2]
        slice_spec = None
        if path.endswith(']') and '[' in path:
            base, slc = path.rsplit('[', 1)
            slice_spec = slc[:-1]  # strip ']'
            path = base
        cur = profile
        for part in path.split('.'):
            if cur is None:
                return None
            cur = getattr(cur, part, None)
        if slice_spec is not None and isinstance(cur, list):
            try:
                a, b = slice_spec.split(':')
                start = int(a) if a else None
                end = int(b) if b else None
                return cur[start:end]
            except Exception:
                return cur
        return cur

    # ------------------------------------------------------------------
    # Role fillers
    # ------------------------------------------------------------------

    def _fill_text_shapes(self, shapes: list, items: list, role: ShapeRole, slide=None):
        """Fill one page's worth of text into shapes; clears unused slots.

        An item too long for its fixed box is truncated in place at a
        sentence boundary — same policy `_fill_compound_pairs` uses for body
        text. These boxes (e.g. business-process step/category labels) sit
        in diagram layouts next to sibling shapes/connectors, so they can't
        safely grow like a title/body pair can; the full text is queued for
        an auto-generated appendix extra-slide instead (`_queue_overflow_extra`)
        so nothing is lost, and the slide itself is never duplicated just to
        make room for one long item (that would repeat every OTHER role on
        the same slide for no reason).

        When `slide` is given, an unused slot's decoration (icon/background —
        not tracked by the role, so blanking text alone would leave an
        empty-looking box) is removed via `_remove_decorations_around`.
        Callers with their own bespoke unused-slot handling (e.g.
        `_maybe_split_business_process_slide`) omit `slide` to skip this.
        """
        filled: list = []
        for i, shape in enumerate(shapes):
            if shape is None:
                continue
            if i >= len(items):
                if shape.has_text_frame:
                    self._set_text_preserving_format(shape.text_frame, '')
                if slide is not None:
                    self._remove_decorations_around(slide, [shape])
                continue
            value = items[i]
            if shape.has_text_frame:
                shape.text_frame.word_wrap = True
                display_value, truncated = self._fit_text_to_shape(shape, value, role)
                lines = [ln for ln in display_value.split('\n') if ln.strip()]
                if len(lines) > 1 and self._native_bullet(shape.text_frame):
                    # A list in a box that bullets its own paragraphs — give each
                    # item a paragraph so the template marks every one of them.
                    self._set_text_lines_preserving_format(shape.text_frame, lines)
                else:
                    self._set_text_preserving_format(shape.text_frame, display_value)
                self._apply_text_font(shape.text_frame)
                # Matches _fill_compound_pairs's body handling: once text is
                # truncated to the estimated capacity, skip auto-fit — on
                # some renderers (LibreOffice) normAutofit's fit calculation
                # can still shrink already-truncated text to near-invisible
                # sizes. Untruncated text (either it fit already, or the box
                # was too small to truncate meaningfully) keeps auto-fit as
                # a safety net — unless this role sizes itself with its
                # siblings, where one shared size is written explicitly below
                # and any auto-fit would undo it per box.
                if not truncated and not role.font_group:
                    self._apply_auto_fit(shape.text_frame)
                filled.append((shape, display_value))
                if role.name == 'current_issues' or role.name == 'objectives':
                    self._apply_line_spacing(shape.text_frame, Pt(10))
        self._queue_shared_font(role, filled, slide)

    def _estimate_max_lines(self, shape, font_pt: int = 10) -> int:
        """How many lines of *font_pt* text fit in *shape*'s height."""
        try:
            h_inch = (shape.height or 0) / self._EMU_PER_INCH
            return max(1, int(h_inch / (font_pt * 1.35 / 72)))
        except Exception:
            return 1

    def _fit_text_to_shape(self, shape, value: str, role: ShapeRole,
                           font_pt: int = 10, can_autofit: bool = True) -> Tuple[str, bool]:
        """Fit *value* to *shape*; return (text to write, was it truncated).

        Shared by `_fill_text_shapes` and `_fill_compound_pairs` so a body box
        and a compound body answer the same question the same way.

        The ladder, least destructive rung first:
          1. a box too small to truncate meaningfully keeps the full text;
          2. text that FITS — measured in real wrapped lines, in its own
             script's width — is never touched;
          3. a LOSSLESS reshape: six challenges set as six bullet lines say
             everything the paragraph said and read better doing it, so it wins
             outright whenever it fits and drops nothing;
          4. text that still overflows only slightly keeps ALL its words and
             lets normAutofit shrink them — a chapter over by one clause is
             better set a point smaller than cut down to its headline;
          5. a LOSSY reshape: headlines only, then points off the end. The
             caller is told how many were left off;
          6. what even that cannot fit is split at a sentence, and the
             remainder goes to an appendix slide.

        Rungs 3 and 5 exist only for a document that opted into shaping;
        `can_autofit` says whether rung 4 exists for THIS shape. Compound bodies
        pass False: this template's cells carry ~0.23" of inset against boxes as
        short as 0.13", so normAutofit goes deeply negative there and some
        renderers shrink to invisibility.
        """
        max_chars = self._estimate_max_chars(shape, font_pt=font_pt, sample=value)
        if max_chars <= self._MIN_MEANINGFUL_TRUNCATION_CHARS:
            return value, False
        chars_per_line = self._estimate_chars_per_line(shape, font_pt, sample=value)
        max_lines = self._estimate_max_lines(shape, font_pt)

        def fits(text: str) -> bool:
            return self._lines_needed(text, chars_per_line) <= max_lines

        if fits(value):
            return value, False

        shaped = ''
        if self._shape_content:
            # The glyph is the SHAPE's business, not the condenser's: a box that
            # bullets its own paragraphs gets plain lines and marks them itself.
            bullet = '' if self._native_bullet(shape.text_frame) else DEFAULT_BULLET
            shaped = condense(value, max_chars=max_chars, bullet=bullet,
                              chars_per_line=chars_per_line, max_lines=max_lines)
            dropped = count_items(value) - emitted_items(shaped)
            if shaped and shaped != value and dropped <= 0 and fits(shaped):
                print(f"  - Role '{role.name}': reshaped to its box — "
                      f'{describe(value, shaped)} ({len(value)} -> {len(shaped)} chars)')
                return shaped, False

        overflow = self._split_at_sentence(value, max_chars=max_chars)[1]
        if can_autofit and len(overflow) < self._MIN_OVERFLOW_APPENDIX_CHARS:
            print(f"  - Role '{role.name}': overflows its box by {len(overflow)} chars — "
                  f'small enough to keep whole and let auto-fit shrink it')
            return value, False

        if shaped and shaped != value and fits(shaped):
            print(f"  - Role '{role.name}': condensed to its box — "
                  f'{describe(value, shaped)} ({len(value)} -> {len(shaped)} chars)')
            return shaped, False

        kept, rest = self._split_at_sentence(value, max_chars=max_chars)
        if not rest:
            return value, False
        self._queue_overflow_extra(role, rest)
        return kept, True

    def _queue_overflow_extra(self, role: ShapeRole, remainder: str) -> None:
        """Record the REMAINDER an item's box could not hold, for an appendix slide.

        The remainder, not the whole item: an appendix built from the full text
        reprinted whatever had just been written on the slide above it, so a
        451-character paragraph that overflowed by one sentence produced a
        second slide repeating 71% of the first.

        Grouped by (section, role name) so e.g. multiple long `before_steps`
        items land on one shared appendix slide rather than one each.
        Rendered by `render_from_profile` after the main fill+prune pass —
        see `_overflow_extras`.
        """
        source_root = (role.source_path or '').split('.')[0].split('[')[0] or role.name
        # The source attribute is usually the section's own name; where it is
        # not, the template says so. An unrecognized section makes the anchor
        # unresolvable and drops the slide at the end of the deck — see
        # `SOURCE_SECTION_ALIASES`.
        try:
            from lib.templates.svn import SOURCE_SECTION_ALIASES
        except ImportError:
            SOURCE_SECTION_ALIASES = {}
        section_key = SOURCE_SECTION_ALIASES.get(source_root, source_root)
        bucket = self._overflow_extras.setdefault((section_key, role.name), [])
        entry = (self._filling_slide_num, remainder)
        if entry not in bucket:
            bucket.append(entry)
        print(f"  - Role '{role.name}': item too long for its box — truncated in place, "
              f'the remaining {len(remainder)} chars queued for an appendix slide')

    def _fill_compound_pairs(self, pairs: list, items: list, role: ShapeRole, slide=None):
        """Fill a single slide's worth of (title, body) pairs; clears unused slots.

        When `slide` is given, an unused pair's decoration (card background,
        icon badge, icon picture — not tracked by the role, so blanking
        title/body text alone would leave an empty-looking card) is removed
        via `_remove_decorations_around`. Callers with their own bespoke
        unused-slot handling (e.g. `_maybe_split_business_process_slide`)
        omit `slide` to skip this.
        """
        titles_filled: list = []
        bodies_filled: list = []
        for i, pair in enumerate(pairs):
            if i >= len(items):
                # Clear unused pair shapes
                if isinstance(pair, tuple):
                    for shp in pair:
                        ShapeDiscovery.clear_text(shp)
                    if slide is not None:
                        self._remove_decorations_around(slide, list(pair))
                continue
            item = items[i]
            title = getattr(item, 'title', None) or (item.get('title') if isinstance(item, dict) else '')
            body = getattr(item, 'body', None) or getattr(item, 'content', None) \
                or (item.get('body') if isinstance(item, dict) else '') \
                or (item.get('content') if isinstance(item, dict) else '')
            if not isinstance(pair, tuple):
                continue
            title_shape, body_shape = pair
            if title_shape is not None and title_shape.has_text_frame:
                # Full text, never truncated with an ellipsis — a heading cut
                # off mid-word ("Referral trac…") reads as meaningless. The type
                # gives instead of the box: growing the box was safe only where
                # a title has exactly one slot beneath it to push down, and
                # slide 16's cards have TWO. A long chapter title grew its box
                # from 0.20" to 0.57" and ran through the headline slot at
                # 2.33", while the body it did push left the four cards' text
                # starting at four different heights. `_equalize_autofit` below
                # shrinks the words instead, by one shared amount, so the grid
                # the template drew survives contact with a real proposal.
                title_text = title or ''
                # Grow into the empty space under the title, and not one EMU
                # further. The template leaves a little room there on purpose;
                # taking it keeps the type at its authored size, which no amount
                # of shrinking can give back. What the gap cannot absorb is
                # handled by `_equalize_autofit`.
                wanted = self._estimate_extra_wrap_height(title_shape, title_text,
                                                          font_pt=10)
                room = self._free_height_below(slide, title_shape)
                if wanted > 0 and room > 0:
                    title_shape.height += min(wanted, room)
                titles_filled.append((title_shape, title_text))
                title_shape.text_frame.word_wrap = True
                self._set_text_preserving_format(title_shape.text_frame, title_text)
                self._apply_text_font(title_shape.text_frame)
                self._set_vertical_anchor_top(title_shape.text_frame)
            if body_shape is not None and body_shape.has_text_frame:
                # font_pt matches _apply_text_font's default (10pt) below — otherwise
                # the char-count estimate (previously assuming 8pt) under-truncates
                # relative to what actually gets rendered, and on template cells with
                # very little height to begin with (some of this template's cells are
                # barely 0.13" tall) the overflow renders as overlapping text lines.
                # No auto-fit here: this template's cells carry ~0.23" top+bottom
                # text insets against boxes as short as 0.13-0.34" tall, so
                # normAutofit's fit calculation goes deeply negative and some
                # renderers (LibreOffice) respond by shrinking the text to
                # invisibility rather than clamping to a sane minimum. Truncating
                # to the estimated capacity is enough to keep text on-shape.
                body_text, _ = self._fit_text_to_shape(
                    body_shape, body or '', role, font_pt=10, can_autofit=False)
                body_shape.text_frame.word_wrap = True
                self._set_text_preserving_format(body_shape.text_frame, body_text)
                self._apply_text_font(body_shape.text_frame)
                self._set_vertical_anchor_top(body_shape.text_frame)
                bodies_filled.append((body_shape, body_text))
        self._equalize_card_title_heights(titles_filled, slide)
        self._equalize_autofit(titles_filled, font_pt=10)
        # The bodies get the same treatment, in a family the sibling roles of
        # this same card can join (slide 16's headline slot is a role of its
        # own). Left on `normAutofit` they were sized by the renderer, one box
        # at a time: the bold headline came out several points below the plain
        # body directly beneath it, on a card where both are authored at the
        # same size.
        self._queue_shared_font(role, bodies_filled, slide)

    def _equalize_card_title_heights(self, filled: list, slide) -> None:
        """Give a row of card titles ONE box height — the tallest they can share.

        Each title box was grown on its own above, by exactly what its own text
        wanted. Four titles of four lengths therefore ended at four heights
        (0.20" / 0.42" / 0.39" / 0.39" on the essence row), and since they share
        a top, the row's titles stopped sharing a baseline — the misalignment
        reads as a mistake on a slide whose whole design is that the cards match.

        Growing the short ones to match the tall one also buys the row line
        space, so `_equalize_autofit` has less to shrink and the titles stay
        nearer the size the template drew them at.

        Only grows. A box that cannot reach the target — something sits too
        close beneath it — is left as tall as it got, because overlapping the
        shape below is worse than a row one box short.
        """
        boxes = [shape for shape, _ in filled if shape is not None]
        if len(boxes) < 2:
            return
        # The tallest height every box in the row can reach — not merely the
        # tallest one already is. Taking the shared ceiling hands the row all
        # the daylight the design left above the next band, which is what lets
        # a long title stay at a readable size instead of being shrunk into the
        # box it happens to have.
        room = {id(shape): self._free_height_below(slide, shape) for shape in boxes}
        ceiling = min(shape.height + room[id(shape)] for shape in boxes)
        target = max(max(shape.height for shape in boxes), ceiling)
        grown = 0
        for shape in boxes:
            short = target - shape.height
            if short <= 0:
                continue
            shape.height += min(short, room[id(shape)])
            grown += 1
        if grown:
            print(f'  - {grown} card title box(es) grown to the row\'s shared '
                  f'height so the titles sit on one line')

    def _free_height_below(self, slide, shape) -> int:
        """EMU of empty slide between *shape*'s bottom and whatever is under it.

        "Under it" means directly under: only shapes whose horizontal span
        overlaps this one can be collided with, so a card in the next column
        never constrains this one. Everything is measured in slide coordinates,
        which is what makes a shape nested in a group comparable with one that
        is not.

        Returns 0 when nothing is below (the caller then leaves the box alone
        rather than growing it off the slide) — the safe answer either way.
        """
        if slide is None:
            return 0
        try:
            left, right = shape.left, shape.left + shape.width
            bottom = shape.top + shape.height
        except TypeError:                     # a shape without explicit geometry
            return 0
        nearest = None
        for other in self._walk_shapes(slide.shapes):
            if other is shape:
                continue
            try:
                o_left, o_top = other.left, other.top
                o_right = o_left + other.width
            except TypeError:
                continue
            if o_top < bottom or o_right <= left or o_left >= right:
                continue
            nearest = o_top if nearest is None else min(nearest, o_top)
        if nearest is None:
            return 0
        return max(0, nearest - bottom - self._SLOT_CLEARANCE_EMU)

    # A hairline of daylight kept between a grown box and the shape beneath it,
    # so "just fits" never renders as "just touches".
    _SLOT_CLEARANCE_EMU = 27432          # 0.03"

    def _queue_shared_font(self, role: ShapeRole, filled: list, slide=None) -> None:
        """Hold *filled* until every role of its font family has been written.

        A role opts in by naming a `font_group`; anything else is a no-op, so
        the boxes that were happy with the renderer's auto-fit keep it.
        """
        if not role.font_group or not filled:
            return
        family = self._font_families.setdefault(role.font_group, {'slide': slide, 'entries': []})
        family['slide'] = family['slide'] or slide
        family['entries'].extend((shape, text, role) for shape, text in filled)

    def _flush_shared_fonts(self) -> None:
        """Set each family's boxes at one size, then forget them.

        Called once per slide, after its roles are filled. Slide 16's card is
        why the decision waits: its bold headline and the plain body under it
        are two separate roles over the same four cards, and sizing each role
        alone put the headline several points below the body it introduces.
        """
        for name, family in self._font_families.items():
            self._fit_font_family(name, family['entries'], family['slide'])
        self._font_families = {}

    # Sizes a card family is allowed to be set at, largest first. 10pt is what
    # `_apply_text_font` writes and what this template authored its card text
    # at; 7pt is the smallest that still reads on a projected slide, and a
    # family that does not fit even there is trimmed instead of shrunk further.
    _FAMILY_FONT_STEPS = (10.0, 9.5, 9.0, 8.5, 8.0, 7.5, 7.0)

    def _fit_font_family(self, name: str, entries: list, slide=None) -> None:
        """Give one card row a single font size — and the room to hold it.

        Two things have to agree for a row of cards to read as a row: the type
        size, and where each block of text starts. Auto-fit gave neither. It is
        computed per box by the renderer, so four cards holding four lengths of
        prose came out in four sizes — the bold headline several points under
        the plain body it introduces on the very same card.

        So the size is chosen here, ONCE: the largest step at which every box in
        the family fits. And fitting is measured against boxes this method is
        free to re-divide. A card's headline slot and its body slot share a band
        of the card; how that band was split was decided by the template for the
        sample text it shipped with, and a real proposal's headline is routinely
        longer than that sample's while its body is shorter. Re-dividing the
        band — the SAME division for every card in the row, never growing past
        what the two boxes already covered — is what lets the row hold real copy
        at a readable size instead of shrinking to fit a split nobody chose.

        What still does not fit at the smallest step is trimmed at a sentence
        boundary, with the remainder queued for the appendix: the contract every
        other overflowing box in this deck follows.
        """
        levels = self._font_family_levels(entries)
        geometry = {id(shape): (self._top_of(shape), shape.height or 0)
                    for shape, _text, _role in entries}
        plan, size_pt = None, self._FAMILY_FONT_STEPS[-1]
        for candidate in self._FAMILY_FONT_STEPS:
            plan = self._plan_row(levels, geometry, candidate) if levels else None
            if plan or (levels is None and self._boxes_hold(entries, candidate)):
                size_pt = candidate
                break
        if plan:
            self._apply_row_plan(levels, geometry, plan, slide)
        for shape, text, role in entries:
            final, _ = self._fit_text_to_shape(shape, text, role,
                                               font_pt=int(size_pt), can_autofit=False)
            if final != text:
                self._set_text_preserving_format(shape.text_frame, final)
            # Top, not the centre the template anchored these to. Centring is
            # right for boxes sized to their own text and wrong for a row that
            # now shares one height: a four-line headline centred beside a
            # six-line one starts a line lower, and the row stops reading as a
            # row — the same misalignment the shared size was chosen to end.
            self._set_vertical_anchor_top(shape.text_frame)
            self._freeze_font_size(shape, size_pt)
        detail = f', its band re-divided into {len(plan)} row(s)' if plan else ''
        print(f'  - Font family "{name}": {len(entries)} box(es) set at '
              f'{size_pt:g}pt, one size for the row{detail}')

    # Sizes a card family is allowed to be set at, largest first. 10pt is what
    # `_apply_text_font` writes and what this template authored its card text
    # at; 7pt is the smallest that still reads on a projected slide, and text
    # that does not fit even there is trimmed rather than shrunk further.
    _FAMILY_FONT_STEPS = (10.0, 9.5, 9.0, 8.5, 8.0, 7.5, 7.0)

    def _font_family_levels(self, entries: list) -> Optional[list]:
        """The family's boxes as ROWS across the cards, top row first.

        Cards are columns — two boxes belong to the same card when their
        horizontal spans overlap — and the n-th box down each card is one row of
        the design. Rows are what get re-divided, not individual cards: giving
        one card's headline more room than its neighbour's would drop that
        card's body text below the others, and the four cards' text starting at
        four heights is exactly the misalignment this design cannot survive.

        Returns None when the cards do not stack alike (different box counts),
        where there is no row to speak of and the geometry is left alone.
        """
        columns: list = []
        for entry in sorted(entries, key=lambda e: self._left_of(e[0])):
            shape = entry[0]
            left, right = self._left_of(shape), self._left_of(shape) + (shape.width or 0)
            for column in columns:
                other = column[0][0]
                o_left = self._left_of(other)
                if left < o_left + (other.width or 0) and o_left < right:
                    column.append(entry)
                    break
            else:
                columns.append([entry])
        if len(columns) < 2:
            return None
        depth = len(columns[0])
        if depth < 2 or any(len(column) != depth for column in columns):
            return None
        for column in columns:
            column.sort(key=lambda e: self._top_of(e[0]))
        return [[column[level] for column in columns] for level in range(depth)]

    @staticmethod
    def _left_of(shape) -> int:
        return shape.left or 0

    @staticmethod
    def _top_of(shape) -> int:
        return shape.top or 0

    def _plan_row(self, levels: list, geometry: dict, font_pt: float) -> Optional[list]:
        """Heights for the family's rows at *font_pt*, or None if they overrun.

        Each row is as tall as its longest text needs, so a row of short bodies
        hands its unused lines to the row of long headlines above it. The band
        is what the boxes already covered — the first row's top to the shallowest
        card's last row bottom — with the gaps the template drew kept intact, so
        no card can end up taller than the card it is drawn inside.
        """
        line_h = int(font_pt * self._FAMILY_LINE_EM / 72 * self._EMU_PER_INCH)
        heights = []
        for row in levels:
            tallest = 0
            for shape, text, _role in row:
                per_line = self._estimate_chars_per_line(shape, int(font_pt), sample=text)
                needed = max(1, self._lines_needed(text or '', per_line))
                tallest = max(tallest, needed * line_h)
            heights.append(tallest)
        band = self._row_band(levels, geometry)
        if band is None or sum(heights) + sum(self._row_gaps(levels, geometry)) > band:
            return None
        return heights

    # How much vertical space one line of type really takes, in em. Measured
    # against this template's own font (Noto Sans JP, whose line metrics run
    # well above a latin face's) as rendered: at 1.35 — what the older
    # box-capacity estimates assume — a six-line block came out taller than the
    # box planned for it and pushed a line through the rule above.
    _FAMILY_LINE_EM = 1.6

    def _boxes_hold(self, entries: list, font_pt: float) -> bool:
        """True when every box of a NON-row family holds its text at *font_pt*.

        A continuation slide can carry a single card, where there is no row to
        re-divide and each box keeps the height the template gave it. Measured
        the same way a row is, so the last card of a chapter is not set smaller
        than the first four for want of a neighbour to compare with.
        """
        line_h = int(font_pt * self._FAMILY_LINE_EM / 72 * self._EMU_PER_INCH)
        for shape, text, _role in entries:
            per_line = self._estimate_chars_per_line(shape, int(font_pt), sample=text)
            needed = max(1, self._lines_needed(text or '', per_line))
            if needed * line_h > (shape.height or 0):
                return False
        return True

    def _row_band(self, levels: list, geometry: dict) -> Optional[int]:
        """Vertical room the rows may share, as the template itself marked it out.

        The floor is the DEEPEST bottom the last row reaches across the cards.
        The template drew that row at four different depths for four different
        sample sentences (0.96" to 1.34" on slide 16) inside cards of one height,
        so the deepest is the design's own statement of how far text may go —
        and taking it for every card is what puts them back on one baseline.
        """
        top = min(geometry[id(shape)][0] for shape, _t, _r in levels[0])
        floor = max(self._authored_bottom(geometry, shape) for shape, _t, _r in levels[-1])
        return floor - top if floor > top else None

    def _row_gaps(self, levels: list, geometry: dict) -> list:
        """The daylight the template drew between consecutive rows, kept as-is."""
        gaps = []
        for upper, lower in zip(levels, levels[1:]):
            above = max(self._authored_bottom(geometry, shape) for shape, _t, _r in upper)
            below = min(geometry[id(shape)][0] for shape, _t, _r in lower)
            gaps.append(max(0, below - above))
        return gaps

    @staticmethod
    def _authored_bottom(geometry: dict, shape) -> int:
        """Where *shape*'s bottom edge was BEFORE any row was re-divided."""
        top, height = geometry[id(shape)]
        return top + height

    def _apply_row_plan(self, levels: list, geometry: dict, heights: list, slide) -> None:
        """Re-divide the band per *heights*, rules and all.

        Anything the template drew between two rows — the dotted hairline under
        slide 16's headline — moves with the boundary it marks, so re-dividing
        never leaves a rule cutting through a line of text.
        """
        gaps = self._row_gaps(levels, geometry)
        top = min(geometry[id(shape)][0] for shape, _t, _r in levels[0])
        for level, (row, height) in enumerate(zip(levels, heights)):
            if level:
                old_gap_top = max(self._authored_bottom(geometry, shape)
                                  for shape, _t, _r in levels[level - 1])
                shift = top - old_gap_top
                for rule in self._rules_in_gap(levels[level - 1], levels[level],
                                               geometry, slide):
                    rule.top = (rule.top or 0) + shift
                top += gaps[level - 1]
            for shape, _text, _role in row:
                shape.top, shape.height = top, height
            top += height

    def _rules_in_gap(self, upper: list, lower: list, geometry: dict, slide) -> list:
        """The template's own hairlines drawn between two rows of boxes.

        Found by geometry rather than by name: a rule is a shape with no height
        of its own (this template draws them as flat freeforms) sitting in the
        gap and overlapping one of the row's boxes horizontally.
        """
        if slide is None:
            return []
        gap_top = max(self._authored_bottom(geometry, shape) for shape, _t, _r in upper)
        gap_bottom = min(geometry[id(shape)][0] for shape, _t, _r in lower)
        spans = [(self._left_of(shape), self._left_of(shape) + (shape.width or 0))
                 for shape, _t, _r in upper]
        boxes = {id(shape) for row in (upper, lower) for shape, _t, _r in row}
        rules = []
        for shape in self._walk_shapes(slide.shapes):
            if id(shape) in boxes:
                continue
            try:
                if (shape.height or 0) > self._RULE_MAX_HEIGHT_EMU:
                    continue
                s_left, s_top = self._left_of(shape), self._top_of(shape)
            except TypeError:
                continue
            if not gap_top <= s_top <= gap_bottom:
                continue
            s_right = s_left + (shape.width or 0)
            if any(s_left < right and left < s_right for left, right in spans):
                rules.append(shape)
        return rules

    # Thicker than this and it is a box, not a rule — and a box must not be
    # dragged around by a row boundary moving.
    _RULE_MAX_HEIGHT_EMU = 9144            # 0.01"

    def _freeze_font_size(self, shape, font_pt: float) -> None:
        """Write *font_pt* onto every run and turn auto-fit OFF.

        The size is WRITTEN, not left to a `fontScale` hint. PowerPoint honours
        the hint; LibreOffice recomputes auto-fit from scratch and sizes each box
        to its own text, which is exactly the per-box difference the caller is
        removing. An explicit size means every renderer agrees.
        """
        from pptx.oxml.ns import qn

        size = Pt(font_pt)
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = size
        bodyPr = shape.text_frame._txBody.find(qn('a:bodyPr'))
        if bodyPr is None:
            return
        for tag in ('a:normAutofit', 'a:spAutoFit', 'a:noAutofit'):
            for existing in bodyPr.findall(qn(tag)):
                bodyPr.remove(existing)
        bodyPr.append(bodyPr.makeelement(qn('a:noAutofit'), {}))

    def _equalize_autofit(self, filled: list, font_pt: int = 10) -> None:
        """Give a row of sibling boxes ONE font scale — the tightest they need.

        `normAutofit` with no `fontScale` leaves the amount to the renderer,
        which decides per box: four cards holding four titles of four lengths
        came out in four different type sizes, on a slide whose whole design is
        that the cards match. Writing the scale explicitly settles it, and
        writing the SAME scale on all of them is what keeps them a set.

        Floored, because a title shrunk past legibility helps nobody; a title
        that still overflows at the floor is left to overflow, which is visible
        and fixable, where invisible text is neither.
        """
        ratios = []
        for shape, text in filled:
            if not (text or '').strip():
                continue
            per_line = self._estimate_chars_per_line(shape, font_pt, sample=text)
            max_lines = self._estimate_max_lines(shape, font_pt)
            needed = self._lines_needed(text, per_line)
            if needed > max_lines >= 1:
                ratios.append(max_lines / needed)
        if not ratios:
            return
        scale = max(self._MIN_TITLE_FONT_SCALE, min(ratios))
        for shape, _ in filled:
            self._freeze_font_size(shape, round(font_pt * scale, 1))
        print(f'  - {len(filled)} card title(s) set at one shared scale '
              f'({scale:.0%}) so the row stays a row')
        # The floor can be reached before the text fits. Say so: a heading two
        # words too long for a chip is fixed by shortening the heading, and the
        # operator can only do that if they are told which slide to look at.
        if scale <= self._MIN_TITLE_FONT_SCALE and min(ratios) < scale:
            longest = max((t for _, t in filled if t), key=len, default='')
            print(f'  - Card title still overruns its box at the {scale:.0%} '
                  f'floor ({len(longest)} chars: "{longest[:60]}") — shorten '
                  f'the source heading if the row must stay inside its rule')

    # Below this the type is too small to read on a projected slide, and an
    # overflow the reader can see beats text they cannot.
    #
    # 0.85, not the 0.62 this started at: at 0.62 a single long English heading
    # ("AI-Driven Development (AIDD) as a First-Class Methodology") dragged the
    # whole essence row down to 6.7pt beside its own 10pt body text, where the
    # template drew those titles at 12pt. The floor is paired with
    # `_equalize_card_title_heights`, which gives the row the line space to hold
    # a long title at this size instead of shrinking to reach it.
    _MIN_TITLE_FONT_SCALE = 0.85

    @staticmethod
    def _is_labelled_table_set(data) -> bool:
        """True when *data* is a list of `{label, table}` groups, not table rows.

        Both shapes are lists of dicts, so the discriminator is the keys: a group
        carries a `table` of its own.
        """
        return bool(data) and isinstance(data, list) and all(
            isinstance(entry, dict) and 'table' in entry for entry in data)

    def _fill_grouped_table_role(
        self, prs: Presentation, slide, slide_index: int,
        shapes: list, groups: List[dict], role: ShapeRole,
        lead_in_roles: Optional[list] = None,
    ) -> int:
        """Give each labelled table its own slide, named by its own label.

        A feature chapter written as `#### 1.4.1 Screen groups` / `#### 1.4.2
        Full feature breakdown` holds several tables with different columns. Fed
        to the single-table path they merged: the second table's header row
        rendered as a data row of the first, and its surplus column was cut to
        fit the first table's grid. Every continuation slide then repeated one
        label over six different tables.

        So each group gets a slide, and the label pill above the grid says which
        group it is — which is what that pill was drawn for. Overflow WITHIN a
        group still splits through `_fill_table_role`, so a long table behaves
        exactly as it always did.
        """
        extra = 0
        current_slide, current_index = slide, slide_index
        current_shapes = shapes
        for i, group in enumerate(groups):
            if i:
                current_slide = self._duplicate_slide(prs, current_index)
                extra += 1
                current_index += 1
                current_shapes = ShapeDiscovery().discover(current_slide, role)
            self._set_table_label(current_slide, role, group.get('label', ''))
            spill = self._fill_table_role(prs, current_slide, current_index,
                                          current_shapes, group.get('table') or [], role,
                                          lead_in_roles=lead_in_roles,
                                          label=group.get('label', ''))
            extra += spill
            current_index += spill
        if len(groups) > 1:
            print(f'  - Role "{role.name}": {len(groups)} labelled table(s), '
                  f'one slide each: ' + ', '.join(g.get('label', '?') for g in groups))
        return extra

    def _set_table_label(self, slide, role: ShapeRole, label: str) -> None:
        """Write *label* into the pill that names this table, if the role has one."""
        pill = self._table_label_shape(slide, role)
        if pill is None or not label:
            return
        self._set_text_preserving_format(pill.text_frame, label)
        self._fit_label_to_pill(pill, label, slide)

    def _table_label_shape(self, slide, role: ShapeRole):
        """The pill shape that names this role's table, or None."""
        if not role.label_shape_name:
            return None
        for shape in slide.shapes:
            if shape.name == role.label_shape_name and shape.has_text_frame:
                return shape
        return None

    def _fit_label_to_pill(self, pill, label: str, slide=None) -> None:
        """Make the chip hold *label* on ONE line — widen it first, shrink after.

        The pill is a drawn chip exactly one line tall, sized for the single word
        the template put in it (`機能一覧表サンプル`). `Full feature breakdown`
        wrapped to two lines and the second was set outside the chip — dark text
        on the slide's own background, reading as a rendering fault.

        Widening comes first because the chip sits alone on its line with the
        whole card to its right, and a label at the size the design chose beats
        the same label shrunk to fit a width nothing needed it to keep.
        """
        # The size the TEMPLATE drew this chip at, not a number picked here: the
        # pill is 8pt on slide 5, and fitting labels at an assumed 10pt made
        # every numbered page's chip larger than the unnumbered one beside it.
        base_pt = self._authored_font_pt(pill, default=self._PILL_FONT_PT)
        usable = (pill.width or 0) / self._EMU_PER_INCH - self._TABLE_CELL_SIDE_PADDING_IN
        wanted = self._text_width_in(label, base_pt)
        if usable <= 0 or wanted <= usable:
            return
        short_by = int((wanted - usable) * self._EMU_PER_INCH)
        room = min(self._free_width_right(slide, pill),
                   int(self._PILL_MAX_WIDTH_IN * self._EMU_PER_INCH) - (pill.width or 0))
        grow = max(0, min(short_by, room))
        if grow:
            pill.width += grow
            usable += grow / self._EMU_PER_INCH
        if wanted > usable:
            self._freeze_font_size(
                pill, max(self._PILL_MIN_FONT_PT,
                          round(base_pt * usable / wanted, 1)))

    def _free_width_right(self, slide, shape) -> int:
        """EMU of empty slide between *shape*'s right edge and what is beside it.

        The horizontal twin of `_free_height_below`, and the same rule: only a
        shape whose VERTICAL span overlaps this one can be collided with, so the
        table under the pill never limits how wide the pill may grow. A shape
        that CONTAINS this one — the card the pill is drawn on — is ignored, or
        the chip could never leave the width it was authored at.
        """
        if slide is None:
            return 0
        try:
            top, bottom = shape.top, shape.top + shape.height
            right = shape.left + shape.width
        except TypeError:
            return 0
        nearest = None
        for other in self._walk_shapes(slide.shapes):
            if other is shape:
                continue
            try:
                o_left, o_top = other.left, other.top
                o_bottom = o_top + other.height
            except TypeError:
                continue
            if o_left < right or o_bottom <= top or o_top >= bottom:
                continue
            nearest = o_left if nearest is None else min(nearest, o_left)
        if nearest is None:
            nearest = self._slide_width_emu or 0
        return max(0, nearest - right - self._SLOT_CLEARANCE_EMU)

    @staticmethod
    def _authored_font_pt(shape, default: float) -> float:
        """The size the template set this shape's text at, or *default*.

        Read from the first run that states one, then the paragraph's own
        default — the two places a template puts it.
        """
        from pptx.oxml.ns import qn

        body = shape.text_frame._txBody
        for tag in ('a:rPr', 'a:defRPr'):
            for element in body.iter(qn(tag)):
                size = element.get('sz')
                if size:
                    return int(size) / 100
        return default

    # Fallback size when the chip states none, the smallest it may be set at
    # before a shorter label is the only real fix, and the widest a chip should
    # ever be drawn.
    _PILL_FONT_PT: int = 10
    _PILL_MIN_FONT_PT: float = 7.0
    _PILL_MAX_WIDTH_IN: float = 3.0

    def _number_table_pages(self, slide, role: ShapeRole, base: str,
                            page: int, pages: int) -> None:
        """Name a page of a split table `label (2/7)`.

        A table over seven slides wore one label seven times, so a reader who
        turned two pages had no way of telling whether they were looking at the
        same grid or a different one.
        """
        pill = self._table_label_shape(slide, role)
        if pill is None or not base:
            return
        label = f'{base} ({page}/{pages})'
        self._set_text_preserving_format(pill.text_frame, label)
        self._fit_label_to_pill(pill, label, slide)

    def _fill_table_role(
        self, prs: Presentation, slide, slide_index: int,
        shapes: list, data, role: ShapeRole,
        lead_in_roles: Optional[list] = None,
        label: Optional[str] = None,
    ) -> int:
        """Fill a table role. Reuses overflow logic by converting to markdown.

        `lead_in_roles` are the slide's own prose roles, cleared on every
        continuation page this creates — see `_blank_lead_in`.

        `label` is the name the document gave this table, when it gave one. Pages
        are numbered `label (2/7)` only then: the alternative is numbering the
        template's own `機能一覧表サンプル`, which advertises five pages of a
        sample rather than naming the grid.
        """
        if not shapes or shapes[0] is None:
            print(f"  - Table shape not found for role '{role.name}'")
            return 0
        shape = shapes[0]
        rows = data if isinstance(data, list) else []
        if not rows:
            return 0
        md_table = self._rows_to_md_table(rows, fill_cols=role.fill_cols)
        target = _LegacyTarget(content_key=role.name, fill_cols=role.fill_cols)
        parsed_rows = self._parse_md_table_rows(md_table)
        # Reshape the grid to the data, then normalize widths, both BEFORE the
        # row estimate below: that estimate reads column widths to predict how
        # the text wraps, so it has to see the widths the fill will actually
        # use. Estimating against the template's 6-column grid is what split a
        # 2-column table over three slides. Both calls are idempotent, so the
        # later ones inside `_fill_table` are no-ops.
        if shape.has_table:
            # Unmerge first: a merged grid blocks the column reshape below and
            # swallows text written into the hidden cells.
            self._unmerge_table_cells(shape.table)
            self._fit_table_cols_to_data(shape.table, parsed_rows, role.fill_cols)
            if role.fill_cols:
                # A fixed label column keeps the template's proportions; only
                # the flat floor applies, and on these grids it is a no-op.
                self._normalize_table_col_widths(shape.table)
            else:
                # Widths come from the WHOLE table, once, and every page of it
                # keeps them. Measured per page instead, the pages disagreed:
                # the same grid was drawn 1.28" / 2.97" / 2.65" on one slide and
                # 1.48" / 3.23" / 2.38" on the next, so a table split over seven
                # slides read as seven different tables. The last page — one
                # `Total` row — spread its four columns evenly and lined none of
                # them up with the six rows above it.
                self._spread_col_widths(
                    shape.table, parsed_rows,
                    sum(col.width or 0 for col in shape.table.columns))
        # Use shape-position-aware, content-aware row limit so tables never
        # overflow into the footer (estimated from this exact data, not the
        # template's placeholder text). fill_cols mode emits row-only markdown
        # (no header line), so every parsed row is data; otherwise the first
        # parsed row is the header.
        has_header = not role.fill_cols
        if role.fill_cols:
            data_rows = parsed_rows
        else:
            data_rows = parsed_rows[1:] if len(parsed_rows) > 1 else []
        # How many rows will SHARE a slide, measured before anything is cut.
        # The trim divides one slide's height between the rows standing on it,
        # and a 33-row feature list does not stand on one slide — charging every
        # row 1/33 of the height cut each cell to a single line and threw away
        # 1,490 characters to make room the split pass was going to make anyway.
        rows_per_slide = max(1, min(
            len(data_rows),
            self._max_rows_for_shape(shape, rows_data=data_rows,
                                     fill_cols=role.fill_cols)))
        # Trim before measuring again: a cell cut to its share of the height
        # changes how many rows fit, and trimming is what keeps a table the
        # template sized for one slide on that one slide.
        if self._trim_table_cells_to_fit(shape, parsed_rows, role, has_header,
                                         rows_per_slide=rows_per_slide):
            md_table = self._md_table_from_rows(parsed_rows, has_header)
        max_rows = self._max_rows_for_shape(shape, rows_data=data_rows,
                                            fill_cols=role.fill_cols)
        chunks = self._split_table_data(md_table, max_rows, has_header=has_header)
        if len(chunks) <= 1:
            self._fill_table(shape, md_table, target)
            return 0

        print(f'  - Role "{role.name}": table split ({len(chunks)} slides, max {max_rows} rows/slide)')
        base = (label or '').strip()
        self._fill_table(shape, chunks[0], target)
        self._number_table_pages(slide, role, base, 1, len(chunks))
        extra = 0
        current_index = slide_index
        for i, chunk in enumerate(chunks[1:], 1):
            new_slide = self._duplicate_slide(prs, current_index)
            extra += 1
            current_index += 1
            # Re-discover same shape on duplicate, then fill
            new_shape = ShapeDiscovery().discover(new_slide, role)
            if new_shape and new_shape[0] is not None:
                self._fill_table(new_shape[0], chunk, target)
                self._number_table_pages(new_slide, role, base, i + 1, len(chunks))
                self._blank_lead_in(new_slide, lead_in_roles)
                print(f'  - Filled continuation {i}/{len(chunks)-1}')
        return extra

    def _blank_lead_in(self, slide, roles: Optional[list]) -> None:
        """Clear the chapter's own prose on a table's continuation page.

        The lead-in introduces the table — `The MVP encompasses 36 screens…` —
        and a duplicated slide inherits it, so a table over nine pages said it
        nine times. Once is the point of it; nine times is the deck reading like
        a mail merge.
        """
        for role in roles or []:
            for shape in ShapeDiscovery().discover(slide, role):
                if shape is not None and shape.has_text_frame:
                    self._set_text_preserving_format(shape.text_frame, '')

    # Image roles whose empty placeholder a table may take over, and where that
    # table comes from in the profile.
    _PLACEHOLDER_TABLE_SOURCES = {'schedule_image': 'schedule.table'}

    def _fill_placeholder_with_table(self, slide, placeholder, role: ShapeRole,
                                     profile: Optional[ProjectProfile]) -> None:
        """Draw a profile table inside an unused image placeholder, then remove it."""
        source = self._PLACEHOLDER_TABLE_SOURCES.get(role.name)
        if not source or profile is None:
            return
        rows = self._resolve_source_path(profile, source) or []
        if not rows:
            return
        from lib.extra_slide_layouts import draw_table_in_box
        from lib.timeline_data import build_chart
        from lib.timeline_layouts import draw_chart_in_box

        headers = list(rows[0].keys())
        # A schedule slide's placeholder was drawn for a PICTURE of the plan, so
        # the chart belongs in it ahead of the grid: dated milestones and a
        # month-by-month plan are what the template's own slides 42 and 46 draw.
        chart = build_chart(rows)
        if chart is not None:
            draw_chart_in_box(slide, placeholder.left, placeholder.top,
                              placeholder.width, placeholder.height, chart)
            placeholder._element.getparent().remove(placeholder._element)
            self._tables_placed.add(source)
            print(f"  - Role '{role.name}': no image — drew the "
                  f'{len(rows)}-row {source} as a {chart.kind} timeline')
            return
        draw_table_in_box(
            slide, placeholder.left, placeholder.top,
            placeholder.width, placeholder.height, headers,
            [[str(row.get(header, '')) for header in headers] for row in rows],
        )
        placeholder._element.getparent().remove(placeholder._element)
        self._tables_placed.add(source)
        print(f"  - Role '{role.name}': no image — drew the {len(rows)}-row "
              f'{source} in its place')

    def _fill_image_role(self, slide, shapes: list, data, role: ShapeRole,
                         profile: Optional[ProjectProfile] = None):
        """Fill an image role. data: str (image path or data URI)."""
        if not shapes or shapes[0] is None:
            print(f"  - Image shape not found for role '{role.name}'")
            return
        if not data:
            if role.name == 'schedule_image' and profile is not None:
                from lib.schedule_gantt_fill import fill_gantt_placeholder
                if fill_gantt_placeholder(slide, shapes[0], profile.schedule):
                    return
            # No image for this placeholder. The schedule slide's Gantt frame is
            # a large empty rectangle when the source has no chart, and that same
            # source often carries a milestone table with nowhere to go — so the
            # table takes the space the picture was going to occupy, rather than
            # leaving a blank box on one slide and adding another slide for it.
            self._fill_placeholder_with_table(slide, shapes[0], role, profile)
            return
        if getattr(role, 'fit_to_slide', False):
            self._place_image_fit_slide(slide, data, role.name)
        else:
            target = _LegacyTarget(content_key=role.name)
            self._replace_shape_with_image(slide, shapes[0], data, target)

    def _place_image_fit_slide(self, slide, image_path: str, role_name: str):
        """Scale image to fill slide content area (below title), preserving aspect ratio.

        Uses Pillow to read native image dimensions.  Falls back to placeholder-based
        placement when the file cannot be opened.
        """
        image_path = image_path.strip()

        # Resolve relative path
        if not Path(image_path).is_absolute():
            for candidate in [Path(image_path),
                               Path('outputs') / image_path,
                               (self._output_dir / image_path) if self._output_dir else None]:
                if candidate and candidate.exists():
                    image_path = str(candidate)
                    break

        if not Path(image_path).exists():
            print(f'  - Image not found for fit-to-slide: {image_path}')
            return

        try:
            from PIL import Image as _PILImage
            with _PILImage.open(image_path) as img:
                img_w, img_h = img.size
        except Exception as e:
            print(f'  - Could not read image dimensions ({e}); skipping fit-to-slide')
            return

        # Use actual slide dimensions recorded during render_from_profile.
        # Fall back to common 10" width if not yet set.
        slide_w = self._slide_width_emu or 9144000
        # Use the dynamically computed safe-bottom (above footer) as the effective
        # lower boundary so the image never overlaps the footer band.
        safe_bottom = self._SVN_SAFE_BOTTOM_EMU
        margin_l = int(0.5 * self._EMU_PER_INCH)   # 0.5"
        margin_r = int(0.5 * self._EMU_PER_INCH)   # 0.5"
        margin_t = int(1.1 * self._EMU_PER_INCH)   # 1.1" (below title bar)

        avail_w = slide_w - margin_l - margin_r
        avail_h = max(0, safe_bottom - margin_t)

        scale = min(avail_w / img_w, avail_h / img_h)
        new_w = int(img_w * scale)
        new_h = int(img_h * scale)

        # Center within the content area
        left = margin_l + (avail_w - new_w) // 2
        top = margin_t + (avail_h - new_h) // 2

        try:
            slide.shapes.add_picture(image_path, left, top, width=new_w, height=new_h)
            print(f"  Placed '{role_name}' fit-to-slide: "
                  f"{new_w / 914400:.2f}\" × {new_h / 914400:.2f}\" "
                  f"at ({left / 914400:.2f}\", {top / 914400:.2f}\")")
        except Exception as e:
            print(f'  - Error placing fit-to-slide image: {e}')

    # ------------------------------------------------------------------
    # Helpers for role-based fill
    # ------------------------------------------------------------------

    @staticmethod
    def _coerce_text_items(data, item_field: Optional[str] = None) -> list:
        """Normalize text-role data into a list of strings.

        `item_field` projects one key out of each entry first, for a role bound
        to one slot of a multi-slot card. Without it a list of dicts stringifies
        to `{'title': ...}` reprs, so the projection has to happen HERE — by the
        time the filler sees the items they are already strings.
        """
        if data is None:
            return []
        if isinstance(data, str):
            return [data]
        if isinstance(data, list):
            if item_field:
                data = [item.get(item_field, '') if isinstance(item, dict)
                        else getattr(item, item_field, '') for item in data]
            return [str(item) if item is not None else '' for item in data]
        return [str(data)]

    @staticmethod
    def _rows_to_md_table(rows: list, fill_cols=None) -> str:
        """Convert list[dict | dataclass] to markdown pipe-table string.

        - When fill_cols is None: emit standard header + data rows.
        - When fill_cols is set: emit row-only format (no header/separator), and
          emit only the *last* len(fill_cols) fields per row (template provides
          fixed columns for the rest, e.g. assumption labels).
        """
        if not rows:
            return ''
        # Normalize rows to list[dict]
        import dataclasses as _dc
        rows = [r if isinstance(r, dict) else _dc.asdict(r) if _dc.is_dataclass(r) else dict(vars(r))
                for r in rows]
        headers = list(rows[0].keys())

        def esc(v):
            return str(v).replace('|', '&#124;').replace('\n', ' ')

        if fill_cols:
            n = len(fill_cols)
            target_keys = headers[-n:] if n <= len(headers) else headers
            lines = []
            for r in rows:
                cells = [esc(r.get(k, '')) for k in target_keys]
                lines.append('| ' + ' | '.join(cells) + ' |')
            return '\n'.join(lines)

        lines = [
            '| ' + ' | '.join(headers) + ' |',
            '|' + '|'.join(['---'] * len(headers)) + '|',
        ]
        for r in rows:
            lines.append('| ' + ' | '.join(esc(r.get(h, '')) for h in headers) + ' |')
        return '\n'.join(lines)


    # ==================================================================
    # Extra-slide rendering — AI-generated, generate-slide-style layouts
    # ==================================================================

    def _extra_chrome_layout(self, prs: Presentation):
        """Layout of the one template slide whose chrome every extra slide uses.

        Fixed by `EXTRA_SLIDE_CHROME_SLIDE` rather than derived from the deck,
        so the footer is byte-identical across decks regardless of which
        sections had content. Returns None for a non-SVN template.
        """
        if not self.template_config:
            return None
        try:
            from lib.templates.svn import EXTRA_SLIDE_CHROME_SLIDE
        except ImportError:
            return None
        idx = EXTRA_SLIDE_CHROME_SLIDE - 1
        if 0 <= idx < len(prs.slides):
            return prs.slides[idx].slide_layout
        return None

    @staticmethod
    def _label_dropped_rows(extras: List[dict], profile: ProjectProfile) -> None:
        """Tell each table layout how to say "and N more rows" in this deck.

        The wording belongs to the DECK's language, and only the profile knows
        which that is. Left to the layout it was Vietnamese, so an English
        proposal carried one Vietnamese line — in a table, where it read as a row
        of data.
        """
        from lib.extra_slide_layouts import _MORE_ROWS_EN, _MORE_ROWS_JA

        wording = _MORE_ROWS_JA if profile.is_japanese_source else _MORE_ROWS_EN
        for entry in extras:
            if entry.get('layout') == 'table':
                entry.setdefault('more_label', wording)

    def _render_extra_slides(
        self, prs: Presentation, extras: List[dict], last_positions: Dict[int, int],
        blank_canvas_layout=None, chrome_layout=None,
    ) -> None:
        """Insert AI-generated slides rendered with generate-slide-style layouts.

        Each entry of *extras* is a dict with keys:
          - layout:          str  (numbered_points|card_grid|comparison_2|... ; default bullets)
          - title:           str  (slide title)
          - <layout content> (e.g. points / items / columns / steps / message)
          - anchor_section:  str | None (profile section name; resolves to anchor slide)
          - anchor_slide:    int | None (explicit slide number; overrides anchor_section)

        New slides are created from a CONTENT slide's layout (so they inherit the
        deck's white background + master footer/page-number/logo) and inserted
        right after the LAST variant of the anchor slide (accounting for overflow
        inserts). Multiple extras sharing an anchor keep their given order.
        """
        from lib.templates.svn import SECTION_TO_SLIDES

        w_emu, h_emu = prs.slide_width, prs.slide_height

        # Prefer a chrome reference slide — a configured CONTENT slide that
        # has slide-level chrome shapes (logo / footer in the bottom strip).
        # Cloning that slide makes the chrome real slide-level shapes (editable)
        # rather than master-only chrome (which appears "captured" / locked).
        # The fixed chrome layout wins when available: cloning whichever
        # content slide survived pruning gave each deck a different footer, and
        # stripping that clone's content also deleted the header band.
        chrome_ref_idx = None
        canvas_layout = chrome_layout
        if canvas_layout is None:
            chrome_ref_idx = self._chrome_canvas_idx(prs, last_positions)
            if chrome_ref_idx is None:
                # Fallback: master-chrome only (legacy behavior).
                canvas_layout = self._content_canvas_layout(prs, last_positions, blank_canvas_layout)
                if canvas_layout is None:
                    print('No content-slide layout available for extra slides')
                    return

        # Group by anchor (preserve given order within a group)
        groups: Dict[int, List[dict]] = {}
        for e in extras:
            anchor = self._resolve_extra_anchor(e, last_positions, SECTION_TO_SLIDES, prs)
            groups.setdefault(anchor, []).append(e)

        # Insertions shift later positions; process in ascending anchor order
        # and track cumulative offset.
        offset = 0
        for anchor in sorted(groups.keys()):
            for i, extra in enumerate(groups[anchor]):
                target = anchor + offset + 1  # 0-indexed insertion point
                donor = extra.get('donor')
                if donor in self._donor_library:
                    new_slide = self._insert_donor_slide(prs, donor, target)
                    self._fill_donor_slide(new_slide, donor, extra)
                    if chrome_ref_idx is not None and chrome_ref_idx >= target:
                        chrome_ref_idx += 1
                    offset += 1
                    print(f'  - Inserted extra slide "{extra.get("title", "")[:40]}"'
                          f' [borrowed the design of template slide {donor}]'
                          f' at index {target}')
                    continue
                if chrome_ref_idx is not None:
                    # Source index for chrome_ref shifts right with every prior
                    # insert that landed at or before it.
                    src_idx = chrome_ref_idx
                    new_slide = self._duplicate_slide_at(prs, src_idx, target)
                    self._strip_content_keep_chrome(new_slide, h_emu)
                    self._fill_extra_slide(new_slide, extra, w_emu, h_emu, clear=False)
                    if src_idx >= target:
                        chrome_ref_idx += 1  # our insert shifted the ref slide
                else:
                    new_slide = self._add_slide_at(prs, canvas_layout, target)
                    self._fill_extra_slide(new_slide, extra, w_emu, h_emu)
                offset += 1
                print(f'  - Inserted extra slide "{extra.get("title", "")[:40]}"'
                      f' [{extra.get("layout", "bullets")}] at index {target}')

    def _content_canvas_layout(self, prs: Presentation, last_positions: Dict[int, int],
                               blank_canvas_layout=None):
        """Pick a CLEAN content-slide layout (light bg + standard small footer logo).

        Some configured slides (e.g. slide 4 in the SVN deck) use a unique master
        with a large decorative side panel and an oversized logo. Using that
        layout for AI-generated extra slides makes them visually inconsistent
        with the rest of the deck. We skip those and prefer the layout of a
        configured content slide whose master only contains small chrome
        (footer text, page number, small logo).
        """
        if not last_positions:
            return blank_canvas_layout or (prs.slides[0].slide_layout if len(prs.slides) else None)

        # Iterate configured slides in template-order; pick the first clean one.
        slide_w = prs.slide_width or 0
        slide_h = prs.slide_height or 0
        slide_area = slide_w * slide_h
        for sn in sorted(last_positions):
            idx = last_positions[sn]
            if not (0 <= idx < len(prs.slides)):
                continue
            layout = prs.slides[idx].slide_layout
            if self._is_clean_master(layout.slide_master, slide_area):
                return layout

        # None of the surviving configured slides has a clean master (e.g. only
        # a two-tone divider-style slide like project_background survived
        # pruning). Reusing its layout would bake that decorative background
        # into every extra slide, so prefer the blank end-card layout captured
        # before pruning instead.
        if blank_canvas_layout is not None:
            return blank_canvas_layout

        # Last resort: earliest configured slide's layout (legacy behavior)
        idx = last_positions[min(last_positions)]
        if 0 <= idx < len(prs.slides):
            return prs.slides[idx].slide_layout
        return prs.slides[0].slide_layout if len(prs.slides) else None

    @staticmethod
    def _is_clean_master(master, slide_area: int) -> bool:
        """A master is 'clean' if it has no large solid-filled decorative shapes.

        Large = >= 25% of the slide area. This filters out section-divider style
        masters (e.g. SVN slide 4 with a pink right-half panel) while keeping
        standard content-slide masters that only carry footer / logo chrome.
        """
        if slide_area <= 0:
            return True
        threshold = slide_area * 0.25
        try:
            for sh in master.shapes:
                if not (sh.width and sh.height):
                    continue
                if sh.width * sh.height < threshold:
                    continue
                # Large shape — clean only if it's not a solid colored fill
                try:
                    fill = sh.fill
                    if fill.type == 1:  # MSO_FILL.SOLID
                        return False
                except Exception:
                    pass
        except Exception:
            return True
        return True

    def _add_slide_at(self, prs: Presentation, layout, target_index: int):
        """Add a fresh slide from *layout* and move it to target_index (0-based)."""
        new_slide = prs.slides.add_slide(layout)
        self._ensure_unique_slide_partname(prs, new_slide)
        pml_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        try:
            sldIdLst = prs.part._element.find(f'{{{pml_ns}}}sldIdLst')
            if sldIdLst is not None:
                new_sldId = list(sldIdLst)[-1]  # just appended → at end
                sldIdLst.remove(new_sldId)
                remaining = list(sldIdLst)
                if target_index >= len(remaining):
                    sldIdLst.append(new_sldId)
                else:
                    remaining[target_index].addprevious(new_sldId)
        except Exception as e:
            print(f'  - Error positioning extra slide: {e}')
        return prs.slides[target_index]

    def _filled_slides_of(self, key: str, slide_nums: list) -> list:
        """Drop the slides of section *key* that the fill pass never wrote to.

        A section owning several template slides splits its items across them
        (`assumptions[0:5]` / `[5:8]` / `[8:9]`). A profile with fewer items
        leaves the surplus slides untouched — and untouched means still holding
        the TEMPLATE'S OWN SAMPLE CONTENT, which then ships as if the client had
        agreed to it. For 前提条件 that is a set of contractual preconditions
        nobody wrote.

        A section never disappears entirely: when NOTHING in its group was
        filled, the first slide is kept as its primary, however thin.

        That fallback used to be written as `position == 0 or filled`, which
        kept the first slide even when the fill pass had explicitly reported it
        empty. It held for a section whose slides are one list cut into pages
        (`assumptions[0:5]`/`[5:8]`/`[8:9]` — the first page always has rows),
        and broke for a section whose slides are ALTERNATIVE DESIGNS for the
        same chapter: `approach_comparison` maps to slide 16 (point cards) and
        slide 21 (comparison matrix), filled from two different fields. A
        chapter written as a matrix only left slide 16 unwritten, and it still
        shipped — a page carrying nothing but its own header.
        """
        if len(slide_nums) < 2:
            return list(slide_nums)
        kept, dropped = [], []
        for number in slide_nums:
            if self._slides_filled.get(number, True):
                kept.append(number)
            else:
                dropped.append(number)
        if not kept:
            # Every slide came back empty. The section still has content the
            # parser recognized, so it keeps its primary rather than vanishing.
            return [slide_nums[0]]
        if dropped:
            print(f'  - Section \'{key}\' filled {len(kept)} of {len(slide_nums)} '
                  f'template slide(s); dropped {dropped} rather than ship the '
                  f"template's own sample content")
        return kept

    @staticmethod
    def _current_slide_index(orig: int, spans: Dict[int, Tuple[int, int]]) -> int:
        """0-based deck index of template slide *orig* after the fill pass.

        Overflow slides inserted by earlier configured sections push everything
        after them down; `spans` records how far. Shared by the prune pass and
        by `_dividers_by_label`, which must read the same slide the prune pass
        will later keep — addressing a divider by its template number after the
        fill inserted continuation slides reads a different slide entirely, and
        the deck then ships a chapter divider the document never had.
        """
        return (orig - 1) + sum(e - s for m, (s, e) in spans.items() if m < orig)

    def _keep_dividers_for_extra_chapters(
        self, prs: Presentation, profile: ProjectProfile, wanted_orig: list,
        dividers_emitted: set, section_map: Dict[str, list],
        spans: Dict[int, Tuple[int, int]],
    ) -> None:
        """Give a divider back to a chapter that reached the deck as appendix only.

        A chapter the schema has no section for still reaches the deck — as
        appendix slides carrying its tables and lists. But the divider that
        announces it was dropped, because the prune pass finds dividers through
        sections and this chapter has none. `4. Scope of Work` was the case: its
        four subsections were in the deck while its chapter heading was not, so
        the reader met a responsibility matrix with no idea which chapter had
        begun.

        Two things changed here from the version that matched by name:

        * **Which chapters qualify** is decided by what the emitted dividers
          already ANNOUNCE (`_chapter_of_divider`), not by comparing the
          chapter's name against the template's own wording. The old test
          compared the AGENDA's `Scope of Work` with the template's
          `Sun*'s Scope`, found nothing, and the rescue never once ran on an
          English proposal.
        * **Which divider it gets** is any unused one. The template's nine
          dividers are the same slide with different words on them, and this
          deck already retitles every one of them from the document — so the
          spare that is free is as good as the one that happens to be named
          for this chapter, and unlike it, it exists.

        The divider is placed exactly where the chapter's first appendix slide
        was going to land, and every slide of the chapter is pinned to it, so
        the heading and its content cannot drift apart.
        """
        if not (profile.auto_extra_slides and profile.agenda.chapters):
            return
        if not dividers_emitted:
            # No chapter reached a template section, so the deck is appendix
            # slides end to end. There is no structure here for a chapter to be
            # missing FROM — rescuing every one of them would put a full-bleed
            # chapter slide in front of every single page and double the deck.
            return
        from lib.templates.svn import CHAPTER_DIVIDER_SLIDES

        pairs = [self._split_agenda_item(item) for item in profile.agenda.chapters]
        announced = {self._chapter_of_divider(d, profile) for d in dividers_emitted}
        announced.discard('')

        # Chapter -> its appendix slides, in the order the document wrote them.
        # A chapter routed whole (no subsections) records no `chapter`, and then
        # its own title IS the chapter name.
        # Keyed by the AGENDA's name for the chapter, not by the raw string.
        # One chapter reaches here under two spellings: the slide routed WHOLE
        # records no `chapter` and is keyed by its title, which still wears the
        # source numbering (`4. Scope of Work`), while its siblings carry the
        # stripped `chapter` field (`Scope of Work`). Keyed raw, that is two
        # buckets for one chapter — the first spent the only spare divider on
        # the chapter's minor slides and pinned them to it, the second found
        # none, and the chapter's main slide ended up on the far side of its own
        # heading. `_name_extras_from_agenda` normalizes the titles, but it runs
        # after the prune pass, so the resolution has to happen here too.
        chapters: Dict[str, list] = {}
        for entry in profile.auto_extra_slides:
            name = (entry.get('chapter') or entry.get('title') or '').strip()
            if not name:
                continue
            label, _ = self._agenda_match(name, pairs)
            chapters.setdefault(label or name, []).append(entry)

        spare = [d for d in CHAPTER_DIVIDER_SLIDES if d not in dividers_emitted]
        # Where the NEXT front-of-deck divider goes. The first one lands ahead of
        # every section the deck fills; the ones after it follow their
        # predecessor, so six chapters that all sit before the first filled
        # section still come out in the order the document wrote them.
        front_cursor = None
        for name, entries in chapters.items():
            if name in announced:
                continue
            # Only a chapter the AGENDA announces gets a divider. A stray
            # heading (`Appendices`, `Executive Summary`) is appendix material,
            # and giving it a chapter divider would promise a chapter the
            # contents page never listed.
            label, gloss = self._agenda_match(name, pairs)
            if not label:
                continue
            if not spare:
                print(f'  - No spare divider left for "{name}" — its slides '
                      f'stay in the deck without a chapter heading')
                continue
            divider = spare.pop(0)
            idx = self._current_slide_index(divider, spans)
            if not (0 <= idx < len(prs.slides)):
                continue
            position = self._divider_insert_position(entries[0], wanted_orig, section_map)
            if (entries[0].get('anchor_section') or '') in self._DECK_FRONT_MATTER:
                position = front_cursor if front_cursor is not None else position
                front_cursor = position + 1
            wanted_orig.insert(position, divider)
            dividers_emitted.add(divider)
            announced.add(name)
            self._write_divider_lines(prs.slides[idx], label, gloss)
            for entry in entries:
                entry['anchor_slide'] = divider
            print(f'  - Kept divider {divider} for "{label}", a chapter that '
                  f'only has appendix slides ({len(entries)} of them)')

    @classmethod
    def _divider_insert_position(cls, entry: dict, wanted_orig: list,
                                 section_map: Dict[str, list]) -> int:
        """Where in *wanted_orig* the rescued divider goes.

        The same anchor the chapter's first appendix slide resolves to, read one
        step earlier: the divider has to come out on the near side of it, or the
        reader passes the content before the heading that names it.

        A chapter anchored to the deck's front matter goes ahead of every filled
        section instead of at the end. Its anchor is the AGENDA, which owns no
        content slide of its own, so there is nothing to come after — and read as
        "nowhere" it went last, which put a lone resolved chapter (a cost table,
        say) in front of the six chapters that lead up to it.
        """
        anchor = entry.get('anchor_section') or ''
        slides = section_map.get(anchor, [])
        present = [wanted_orig.index(n) for n in slides if n in wanted_orig]
        if not present:
            if anchor in cls._DECK_FRONT_MATTER:
                opening = cls._first_chapter_position(wanted_orig, section_map)
                if opening is not None:
                    return opening
            return len(wanted_orig)
        if (entry.get('anchor_edge') or 'after') == 'before':
            return min(present)
        return max(present) + 1

    def _write_divider_lines(self, slide, label: str, gloss: str) -> None:
        """Write a rescued divider's two lines.

        `_retitle_dividers_from_agenda` has already run by now and left this
        slide wearing the template's own chapter name — the name of a chapter
        this deck does not have. It is rewritten here rather than there because
        only the prune pass knows which spare divider was handed to which
        chapter.
        """
        shapes = slide.shapes
        if len(shapes) <= max(self._DIVIDER_JP_IDX, self._DIVIDER_EN_IDX):
            return
        jp, en = shapes[self._DIVIDER_JP_IDX], shapes[self._DIVIDER_EN_IDX]
        if jp.has_text_frame:
            self._set_text_preserving_format(jp.text_frame, label)
        if en.has_text_frame:
            self._set_text_preserving_format(en.text_frame, gloss or label)

    def _dividers_by_label(self, prs: Presentation,
                           spans: Dict[int, Tuple[int, int]]) -> Dict[str, int]:
        """Map each chapter divider's Japanese line to its TEMPLATE slide number.

        Runs inside the prune pass, so the deck already holds the continuation
        slides the fill inserted: each divider is READ at its current index
        (`_current_slide_index`) but RECORDED under its template number, which
        is the numbering `wanted_orig` and `dividers_emitted` speak.
        """
        from lib.templates.svn import CHAPTER_DIVIDER_SLIDES

        labels: Dict[str, int] = {}
        for slide_no in CHAPTER_DIVIDER_SLIDES:
            # Bound-check the RESOLVED index, not the template number: after the
            # fill the deck is longer, and it is the resolved index that is
            # about to be dereferenced.
            idx = self._current_slide_index(slide_no, spans)
            if not (0 <= idx < len(prs.slides)):
                continue
            shapes = prs.slides[idx].shapes
            if len(shapes) <= self._DIVIDER_JP_IDX:
                continue
            shape = shapes[self._DIVIDER_JP_IDX]
            if shape.has_text_frame and shape.text_frame.text.strip():
                labels[shape.text_frame.text.strip()] = slide_no
        return labels

    def _resolve_extra_anchor(
        self, extra: dict, last_positions: Dict[int, int],
        section_map: Dict[str, list], prs: Presentation,
    ) -> int:
        """Return 0-indexed slide position to insert AFTER.

        ``last_positions`` reflects POST-prune positions — a requested anchor
        (explicit slide number or section) may no longer exist if it was
        pruned for having no content. In that case fall back to appending
        after the last surviving slide rather than using a stale/incorrect
        index (there is no valid "pre-prune index" once slides are deleted).
        """
        fallback = max(last_positions.values()) if last_positions else len(prs.slides) - 2

        explicit = extra.get('anchor_slide')
        if isinstance(explicit, int) and explicit > 0:
            return last_positions.get(explicit, fallback)
        section = extra.get('anchor_section')
        from lib.templates.svn import DIVIDER_ONLY_SECTIONS
        if section in DIVIDER_ONLY_SECTIONS:
            # No content slide to follow — the extras ARE the chapter body, so
            # they go straight after its divider.
            return last_positions.get(DIVIDER_ONLY_SECTIONS[section], fallback)
        if section and section in section_map:
            resolved = self._anchor_within_section(section, section_map, last_positions,
                                                   extra.get('anchor_edge'))
            if resolved is not None:
                return resolved
            # The section reached no surviving slide at all — usually because
            # its chapter's content had the wrong shape for the box (a list
            # where the template wants a table), so the slide was pruned. The
            # nearest EARLIER section that did survive is where this chapter
            # sits in the document; the end of the deck is not.
            neighbour = self._anchor_from_earlier_section(section, section_map,
                                                          last_positions)
            if neighbour is not None:
                return neighbour
            return fallback
        if ((extra.get('anchor_edge') or '') == 'before'
                or section in self._DECK_FRONT_MATTER):
            # A chapter standing ahead of every section the deck fills — an
            # executive summary, a preface. It has no anchor section because the
            # document had not reached one yet, and appending it to the end put
            # the summary of the proposal after the terms and conditions.
            #
            # Anchored to the AGENDA is the same case wearing a different label:
            # the agenda owns no content slide, so a chapter that follows it and
            # nothing else is ahead of every section too. Read as an unknown
            # anchor it went to the end of the deck instead, which on a proposal
            # where only the cost chapter resolved printed the price first and
            # the six chapters that justify it afterwards.
            opening = self._anchor_before_first_section(section_map, last_positions)
            if opening is not None:
                return opening
        # Default: append before the extra-slide template (i.e. after last filled slide)
        if last_positions:
            return max(last_positions.values())
        return len(prs.slides) - 2  # before the template slide itself

    # Sections that head the DECK rather than sit in it: neither owns a content
    # slide, so anything anchored to one belongs at the front.
    _DECK_FRONT_MATTER = ('cover', 'agenda')

    @staticmethod
    def _first_chapter_position(wanted_orig: list,
                                section_map: Dict[str, list]) -> Optional[int]:
        """Index of the first FILLED chapter in *wanted_orig* — divider included.

        A chapter begins at its divider, not at its content slide, so a slide
        going "ahead of every chapter this deck fills" has to land ahead of that
        divider too. Measured against the content slides alone, an insertion
        landed between the `Pricing` divider and the cost table it announces.
        """
        from lib.templates.svn import CHAPTER_DIVIDERS

        filled = {n for numbers in section_map.values() for n in numbers
                  if n in wanted_orig}
        starts = set(filled)
        for divider, contents in CHAPTER_DIVIDERS.items():
            if divider in wanted_orig and filled.intersection(contents):
                starts.add(divider)
        positions = [wanted_orig.index(n) for n in starts]
        return min(positions) if positions else None

    def _anchor_within_section(self, section: str, section_map: Dict[str, list],
                               last_positions: Dict[int, int],
                               edge: Optional[str]) -> Optional[int]:
        """Insertion point for a slide anchored to *section*, or None if it is gone.

        A section owns several template slides (`approach_comparison` owns 16
        and 21) and the prune pass keeps only the ones its content filled.
        Reading `[-1]` and giving up when that one slide is missing sent eight
        appendix slides — a whole chapter among them — to the end of the deck.
        Every slide of the section is tried, so the anchor survives as long as
        any part of the section does.
        """
        slides = section_map.get(section) or []
        if (edge or 'after') == 'before':
            firsts = getattr(self, '_first_positions', {})
            starts = [firsts[n] for n in slides if n in firsts]
            # One before the section's first slide: the insertion loop places
            # the new slide AFTER the index it is handed.
            return min(starts) - 1 if starts else None
        for slide_num in reversed(slides):
            if slide_num in last_positions:
                return last_positions[slide_num]
        return None

    def _anchor_before_first_section(self, section_map: Dict[str, list],
                                     last_positions: Dict[int, int]) -> Optional[int]:
        """Insertion point just ahead of the deck's first content section."""
        for key in getattr(self, '_section_order', []) or []:
            found = self._anchor_within_section(key, section_map, last_positions,
                                                'before')
            if found is not None:
                return found
        return None

    def _anchor_from_earlier_section(self, section: str, section_map: Dict[str, list],
                                     last_positions: Dict[int, int]) -> Optional[int]:
        """Insertion point after the nearest earlier section that survived.

        `section_order` is the order the DOCUMENT introduced its sections, so
        walking back through it lands the orphaned slide where its chapter
        actually stands rather than after everything else in the deck.
        """
        order = getattr(self, '_section_order', []) or []
        if section not in order:
            return None
        for earlier in reversed(order[:order.index(section)]):
            found = self._anchor_within_section(earlier, section_map,
                                                last_positions, 'after')
            if found is not None:
                return found
        return None

    def _fill_extra_slide(self, slide, extra: dict, w_emu: int, h_emu: int,
                          clear: bool = True) -> None:
        """Render a generate-slide-style layout onto the cloned canvas.

        When *clear* is True the slide is cleared first (use for slides created
        from a layout). When False, existing shapes are kept — used after the
        slide was duplicated from an anchor and its content shapes stripped, so
        only the chrome (logo / footer) remains and the layout is drawn on top.

        See lib/extra_slide_layouts.py for supported layouts.
        """
        from lib.extra_slide_layouts import render_extra_layout
        render_extra_layout(slide, extra, w_emu, h_emu, clear=clear)

    # ------------------------------------------------------------------
    # Chrome-aware slide cloning helpers (used for extra AI-generated slides)
    # ------------------------------------------------------------------

    def _chrome_canvas_idx(self, prs: Presentation,
                           last_positions: Dict[int, int]) -> Optional[int]:
        """Find a configured content slide that has slide-level chrome shapes.

        Chrome shapes are small elements parked in the bottom strip of the
        slide (logo, footer text). Cloning such a slide preserves the chrome as
        real slide-level shapes — editable in PowerPoint, rather than appearing
        as locked master-only chrome.

        Returns the slide index, or None when no suitable slide is found.
        """
        if not last_positions:
            return None
        slide_w = prs.slide_width or 0
        slide_h = prs.slide_height or 0
        if slide_h <= 0:
            return None
        slide_area = slide_w * slide_h
        threshold_top = int(slide_h * 0.88)  # shapes anchored in bottom ~12%
        # Iterate configured slides in template order, prefer clean masters.
        for sn in sorted(last_positions):
            idx = last_positions[sn]
            if not (0 <= idx < len(prs.slides)):
                continue
            slide = prs.slides[idx]
            if not self._is_clean_master(slide.slide_layout.slide_master, slide_area):
                continue
            for sh in slide.shapes:
                if sh.top is not None and sh.top >= threshold_top:
                    return idx
        return None

    def _duplicate_slide_at(self, prs: Presentation, source_idx: int,
                            target_idx: int):
        """Duplicate slide at *source_idx* and insert the copy at *target_idx*.

        Differs from :meth:`_duplicate_slide` (which always inserts immediately
        after the source) in that it places the copy at an arbitrary position.
        Also re-binds any relationship references (``r:embed`` / ``r:link``
        / ``r:id``) inside the cloned shapes so embedded images survive the
        cross-slide copy — otherwise the new slide's blipFills would resolve
        their rIds against the wrong relationship table.
        """
        source_slide = prs.slides[source_idx]

        new_slide = prs.slides.add_slide(source_slide.slide_layout)
        self._ensure_unique_slide_partname(prs, new_slide)
        src_tree = source_slide.shapes._spTree
        new_tree = new_slide.shapes._spTree
        for child in list(new_tree):
            new_tree.remove(child)
        for child in src_tree:
            new_tree.append(copy.deepcopy(child))

        self._remap_shape_rels(source_slide.part, new_slide.part, new_tree)

        pml_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        try:
            sldIdLst = prs.part._element.find(f'{{{pml_ns}}}sldIdLst')
            if sldIdLst is not None:
                new_sldId = list(sldIdLst)[-1]  # just appended at end
                sldIdLst.remove(new_sldId)
                remaining = list(sldIdLst)
                if target_idx >= len(remaining):
                    sldIdLst.append(new_sldId)
                else:
                    remaining[target_idx].addprevious(new_sldId)
        except Exception as e:
            print(f'  - Error positioning extra slide: {e}')
        return prs.slides[target_idx]

    @staticmethod
    def _remap_shape_rels(source_part, target_part, new_spTree) -> None:
        """Re-bind ``r:embed`` / ``r:link`` / ``r:id`` references in *new_spTree*
        to relationships on *target_part*.

        Cross-slide deep-copies preserve the textual rId values (e.g. ``rId1``)
        but those IDs are scoped to the originating slide's relationship table.
        Without remapping, an image blipFill copied from another slide ends up
        pointing at whatever the destination slide's ``rId1`` happens to be
        (typically the slideLayout), causing the image to vanish.
        """
        r_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        src_rels = source_part.rels
        rid_map: Dict[str, str] = {}
        for el in new_spTree.iter():
            for attr_name in list(el.attrib):
                if not attr_name.startswith(f'{{{r_ns}}}'):
                    continue
                local = attr_name.split('}', 1)[1]
                if local not in ('embed', 'link', 'id'):
                    continue
                old_rid = el.attrib[attr_name]
                if not old_rid:
                    continue
                if old_rid in rid_map:
                    el.attrib[attr_name] = rid_map[old_rid]
                    continue
                try:
                    rel = src_rels[old_rid]
                except KeyError:
                    continue
                try:
                    if rel.is_external:
                        new_rid = target_part.relate_to(
                            rel.target_ref, rel.reltype, is_external=True)
                    else:
                        new_rid = target_part.relate_to(
                            rel.target_part, rel.reltype)
                except Exception as exc:
                    print(f'  - Warn: could not rebind {old_rid} ({rel.reltype}): {exc}')
                    continue
                rid_map[old_rid] = new_rid
                el.attrib[attr_name] = new_rid

    @staticmethod
    def _strip_content_keep_chrome(slide, slide_h_emu: int) -> None:
        """Remove every shape whose top edge sits above the bottom chrome strip.

        Anything anchored in the bottom ~12% of the slide is treated as chrome
        (logo / footer text) and retained so the layout renderer can draw the
        new content on top while the chrome remains intact and editable.
        """
        if slide_h_emu <= 0:
            return
        threshold = int(slide_h_emu * 0.88)
        for sh in list(slide.shapes):
            if sh.top is None or sh.top < threshold:
                sh._element.getparent().remove(sh._element)


class _LegacyTarget:
    """Lightweight stand-in for legacy ShapeTarget used by _fill_table / _replace_shape_with_image."""

    def __init__(self, content_key: str, fill_cols=None):
        self.content_key = content_key
        self.fill_cols = fill_cols
        self.shape_index = None
        self.shape_name = None
