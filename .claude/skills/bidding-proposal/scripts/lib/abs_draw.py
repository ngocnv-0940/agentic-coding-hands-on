"""Absolute-inch drawing primitives + the slide-43 legend palette.

`extra_slide_layouts._Canvas` scales the generate-slide canvas to the slide;
the schedule/team layouts instead draw in TEMPLATE inches (10 x 5.625) because
the Gantt must also land inside a fixed frame on template slide 43. The owner
palette mirrors that slide's own legend (その他 / 貴社 / Sun* / 次フェーズ) so a
bar means the same on slide 43 and on an extra slide.
"""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from lib.brand import (
    BORDER_GREY, EMU_IN as _EMU_IN, FONT, GEN_H, GEN_W, LIGHT_GREY, SUN_GOLD, SUN_RED,
    TEXT_DARK, TEXT_GREY, WHITE,
)

# Slide 43's legend colours, read off the template: 貴社 blue, Sun* red, その他
# grey; 次フェーズ is a white bar with a dashed red outline.
CLIENT_BLUE = RGBColor(0x02, 0x8C, 0xD7)
NEUTRAL_GREY = RGBColor(0x9E, 0x9E, 0x9E)
OWNER_COLORS = {'sun': SUN_RED, 'client': CLIENT_BLUE, 'other': NEUTRAL_GREY}
OWNER_LABELS = {'other': 'その他', 'client': '貴社', 'sun': 'Sun*', 'next': '次フェーズ'}
LANE_ACCENTS = {'red': SUN_RED, 'gold': SUN_GOLD, 'blue': CLIENT_BLUE, 'grey': NEUTRAL_GREY}

# Content band of an extra slide, in template inches: under the two-line
# header, above the footer strip (0.87 x slide height).
CONTENT_BAND = (0.47, 1.15, 9.05, 3.72)


def content_band(cv) -> tuple:
    """CONTENT_BAND scaled from the 10 x 5.625 template to *cv*'s slide size."""
    sw, sh = cv.sx * GEN_W, cv.sy * GEN_H
    l, t, w, h = CONTENT_BAND
    return l * sw / 10, t * sh / 5.625, w * sw / 10, h * sh / 5.625


def _emu(inch: float) -> Emu:
    return Emu(int(inch * _EMU_IN))


def _rect(slide, l, t, w, h, fill=None, line=None, dash=False,
          shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, _emu(l), _emu(t), _emu(w), _emu(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
        if dash:
            sp.line.dash_style = MSO_LINE.DASH
    sp.shadow.inherit = False
    return sp


def _text(slide, l, t, w, h, value, pt, color, *, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(_emu(l), _emu(t), _emu(w), _emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value if isinstance(value, list) else [value]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = str(line)
        run.font.name = FONT
        run.font.size = Pt(max(6.5, pt))
        run.font.bold = bold
        run.font.color.rgb = color
    return tb
