"""Render generate-slide-style layouts onto a python-pptx slide (SVN deck canvas).

Ports a subset of the Sun* slide layout library
(.claude/skills/generate-slide/references/layouts.md) to python-pptx so that
AI-authored "extra" slides match the polished generate-slide look while living
inside the SVN proposal deck.

Coordinates in this module are authored for the generate-slide LAYOUT_WIDE
canvas (13.333 x 7.5 in) and scaled automatically to the target slide size
(SVN template is 10 x 5.625 in). Brand constants mirror
generate-slide/assets/helpers.js.

Supported layouts (entry['layout']):
  - bullets        (default) : {bullets: [str, ...]}
  - numbered_points          : {points: [{header, body}, ...]}  (3-5)
  - card_grid                : {items: [{title, body, icon?}, ...]} (up to 4)
  - comparison_2             : {columns: [{title, items: [str]}, {title, items: [str]}]}
  - process_flow             : {steps: [str, ...]} (3-5)
  - hero                     : {message: str}
  - gantt / org_chart        : see `gantt_layout.py` / `org_chart_layout.py`
"""
from __future__ import annotations

import math

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# Brand constants live in `brand.py` (shared with the absolute-inch layouts);
# re-exported here so existing imports keep working.
from lib.brand import (  # noqa: F401
    BORDER_GREY, EMU_IN as _EMU_IN, FONT, GEN_H, GEN_W, LIGHT_GREY, SUN_DARK_RED,
    SUN_GOLD, SUN_RED, TEXT_DARK, TEXT_GREY, WHITE,
)

# --- SVN template header chrome -------------------------------------------
# Absolute inches on the template's own 10 x 5.625 slide, copied verbatim from
# its content slides (PlaceHolder 1 / PlaceHolder 2 on e.g. slides 5, 6, 21).
# These are NOT scaled from the generate-slide canvas: the header must land in
# exactly the same place as on every hand-authored slide, so it is pinned to
# the template's coordinates and type sizes.
HEADER_GREY = RGBColor(0x99, 0x99, 0x99)
HEADER_BLACK = RGBColor(0x00, 0x00, 0x00)  # template titles are pure black, not TEXT_DARK
HDR_BREADCRUMB = {'left': 0.213, 'top': 0.108, 'width': 9.05, 'height': 0.22, 'size': 8}
HDR_TITLE = {'left': 0.472, 'top': 0.355, 'width': 9.05, 'height': 0.50, 'size': 16}
# Grey lead-in of the breadcrumb. Template slides use their chapter name here;
# AI-generated slides are appendix material, so this is fixed.
HDR_BREADCRUMB_PREFIX = 'Appendix'

# --- SVN template table style ---------------------------------------------
# Read off the template's OWN tables, not chosen: slides 5 (機能一覧表), 6 and 36
# agree on every value below, and all seven tables in the file share the header
# fill. An appendix table drawn in the generate-slide palette (SUN_GOLD #B69256
# with #DDDDDD hairlines) landed a visible shade off every hand-authored table
# in the same deck — near-miss reads as a mistake where a clean difference
# would not.
#
# Borders are the one value the template does NOT hold constant: 5/6/36 use
# #BFBFBF at 12240 EMU, 33/39 use #9E9E9E at 9360, and 23/56 use white gaps at
# 28080. The 機能一覧表 style is taken as canonical — it is the most common and
# the table a reader meets first.
TBL_HEAD_FILL = RGBColor(0xD3, 0xB0, 0x3A)
TBL_HEAD_TEXT = WHITE
TBL_BODY_FILL = WHITE
TBL_BODY_TEXT = RGBColor(0x00, 0x00, 0x00)
TBL_BORDER = RGBColor(0xBF, 0xBF, 0xBF)
TBL_BORDER_PT = 0.964            # 12240 EMU, the template's own hairline
TBL_MARGIN_X_IN = 0.075          # 68400 EMU
TBL_MARGIN_Y_IN = 0.05           # template leaves marT/marB at the OOXML default


class _Canvas:
    """Coordinate/font scaler + shape helpers bound to one slide."""

    def __init__(self, slide, w_emu: int, h_emu: int):
        self.slide = slide
        self.sx = (w_emu / _EMU_IN) / GEN_W
        self.sy = (h_emu / _EMU_IN) / GEN_H
        self.fs = self.sy  # font scale follows vertical scale

    def _x(self, inch: float) -> Emu:
        return Emu(int(inch * self.sx * _EMU_IN))

    def _y(self, inch: float) -> Emu:
        return Emu(int(inch * self.sy * _EMU_IN))

    def rect(self, x, y, w, h, color: RGBColor):
        sp = self.slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, self._x(x), self._y(y), self._x(w), self._y(h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        return sp

    def oval(self, x, y, w, h, color: RGBColor):
        sp = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, self._x(x), self._y(y), self._x(w), self._y(h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        return sp

    def text(self, x, y, w, h, value, size, color, *, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        tb = self.slide.shapes.add_textbox(self._x(x), self._y(y), self._x(w), self._y(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = str(line)
            run.font.name = FONT
            run.font.size = Pt(max(7, size * self.fs))
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
        return tb

    def abs_text(self, box: dict, runs) -> None:
        """Text pinned to absolute template inches, bypassing canvas scaling.

        `runs` is a list of (text, color, bold) tuples sharing `box['size']`,
        so a single line can mix colors the way the template's breadcrumb does.
        """
        tb = self.slide.shapes.add_textbox(
            Emu(int(box['left'] * _EMU_IN)), Emu(int(box['top'] * _EMU_IN)),
            Emu(int(box['width'] * _EMU_IN)), Emu(int(box['height'] * _EMU_IN)))
        tf = tb.text_frame
        # Must wrap: with wrap="none" the auto-fit shrinks the box to the text
        # and re-centers it, which drags the header away from its fixed left edge.
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT  # a fresh textbox would otherwise inherit centering
        for text, color, bold in runs:
            run = p.add_run()
            run.text = str(text)
            run.font.name = FONT
            run.font.size = Pt(box['size'])
            run.font.bold = bold
            run.font.color.rgb = color


def _draw_template_header(cv: _Canvas, title: str, prefix: str = '') -> None:
    """Reproduce the template's two-line header (breadcrumb + title).

    Matches the template's content slides exactly in position, font, size and
    color so an AI-generated slide is indistinguishable from a hand-authored
    one. The old hand-drawn 26pt title and red underline bar were neither.

    `prefix` is the grey lead-in of the breadcrumb. A template slide puts its
    chapter's English name there, so a generated slide belonging to a chapter
    the AGENDA lists passes that name in and reads like the rest of the deck.
    Everything else falls back to `HDR_BREADCRUMB_PREFIX` ('Appendix'), which
    is honest about the slide not belonging to any chapter.
    """
    cv.abs_text(HDR_BREADCRUMB, [
        (f'{prefix or HDR_BREADCRUMB_PREFIX}  ', HEADER_GREY, True),
        ('｜  ', HEADER_GREY, True),
        (title, SUN_RED, True),
    ])
    cv.abs_text(HDR_TITLE, [(title, HEADER_BLACK, True)])


def render_extra_layout(slide, entry: dict, w_emu: int, h_emu: int,
                        clear: bool = True) -> None:
    """Render the chosen layout for one extra slide.

    When *clear* is True (default), the slide's own shapes are removed first —
    suitable for a fresh slide created from a content-slide layout where chrome
    comes from the master.

    When *clear* is False, existing shapes are kept (caller is expected to have
    already left only the chrome shapes on the slide, e.g. via duplicating an
    anchor slide then stripping the content shapes). This allows the chrome
    (logo / footer) to remain as editable slide-level shapes rather than only
    living on the master.
    """
    if clear:
        _clear_shapes(slide)
    cv = _Canvas(slide, w_emu, h_emu)

    title = (entry.get('title') or '').strip()
    layout = (entry.get('layout') or 'bullets').strip().lower()

    if title:
        _draw_template_header(cv, title, (entry.get('breadcrumb') or '').strip())

    renderer = _LAYOUTS.get(layout, _bullets)
    renderer(cv, entry)


def _clear_shapes(slide) -> None:
    """Remove the canvas slide's own shapes (master chrome is unaffected)."""
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


# ---------------------------------------------------------------------------
# Layout renderers — coordinates from layouts.md (13.333 x 7.5 canvas)
# ---------------------------------------------------------------------------

def _bullets(cv: _Canvas, entry: dict) -> None:
    """A plain list, drawn in one of the template's own list designs.

    The layout name stays `bullets` — it is the contract an --extra-slides JSON
    and every caller already writes — but what it draws is now picked from the
    list's shape by `list_layouts`. `variant` overrides that choice when the
    classifier (or an author) has already made it.
    """
    from lib.list_layouts import render_list

    bullets = entry.get('bullets') or entry.get('content', {}).get('bullets') or []
    if isinstance(bullets, str):
        bullets = [b.strip() for b in bullets.splitlines() if b.strip()]
    render_list(cv, bullets, entry.get('variant') or '')


def _numbered_points(cv: _Canvas, entry: dict) -> None:
    points = _get(entry, 'points')
    n = max(1, len(points))
    step = 1.4 if n <= 4 else 0.95
    for i, p in enumerate(points[:5]):
        y = 1.55 + i * step
        cv.oval(0.7, y, 0.6, 0.6, SUN_RED)
        cv.text(0.7, y, 0.6, 0.6, str(i + 1), 20, WHITE, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(1.5, y - 0.05, 11.0, 0.45, _f(p, 'header'), 16, TEXT_DARK,
                bold=True, anchor=MSO_ANCHOR.MIDDLE)
        body = _f(p, 'body')
        if body:
            cv.text(1.5, y + 0.45, 11.0, step - 0.5, body, 12, TEXT_GREY)


def _card_grid(cv: _Canvas, entry: dict) -> None:
    items = _get(entry, 'items')[:4]
    for i, item in enumerate(items):
        col, row = i % 2, i // 2
        x = 0.6 + col * 6.2
        y = 1.5 + row * 2.35
        cv.rect(x, y, 6.0, 2.1, LIGHT_GREY)
        cv.rect(x, y, 0.1, 2.1, SUN_RED)
        icon = _f(item, 'icon')
        tx = x + (0.95 if icon else 0.3)
        if icon:
            cv.text(x + 0.25, y + 0.15, 0.7, 0.7, icon, 30, TEXT_DARK,
                    anchor=MSO_ANCHOR.MIDDLE)
        cv.text(tx, y + 0.2, 4.8, 0.45, _f(item, 'title'), 15, TEXT_DARK,
                bold=True, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.3, y + 0.95, 5.4, 1.0, _f(item, 'body'), 11, TEXT_GREY)


def _comparison_2(cv: _Canvas, entry: dict) -> None:
    cols = _get(entry, 'columns')[:2]
    accents = [SUN_GOLD, SUN_RED]
    bgs = [RGBColor(0xFA, 0xF5, 0xEC), RGBColor(0xFF, 0xEE, 0xEC)]
    col_w, start_x, start_y, col_h = 6.0, 0.6, 1.55, 4.9
    for i, col in enumerate(cols):
        x = start_x + i * (col_w + 0.3)
        cv.rect(x, start_y, col_w, col_h, bgs[i])
        cv.rect(x, start_y, 0.1, col_h, accents[i])
        cv.text(x + 0.35, start_y + 0.2, col_w - 0.6, 0.55, _f(col, 'title'),
                19, accents[i], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        items = col.get('items') or []
        lines = [f'■  {it}' for it in items]
        cv.text(x + 0.35, start_y + 1.0, col_w - 0.6, col_h - 1.2, lines, 12, TEXT_DARK)


def _process_flow(cv: _Canvas, entry: dict) -> None:
    steps = _get(entry, 'steps')[:5]
    if isinstance(steps, str):
        steps = [s.strip() for s in steps.splitlines() if s.strip()]
    n = max(1, len(steps))
    gap = 0.2
    total = GEN_W - 1.2
    step_w = (total - gap * (n - 1)) / n
    start_x = 0.6
    y = 3.0
    for i, s in enumerate(steps):
        x = start_x + i * (step_w + gap)
        cv.rect(x, y, step_w, 1.5, LIGHT_GREY)
        cv.rect(x, y, 0.1, 1.5, SUN_RED)
        cv.text(x + 0.25, y + 0.15, step_w - 0.4, 0.5, str(i + 1), 18, SUN_RED,
                bold=True)
        cv.text(x + 0.25, y + 0.65, step_w - 0.45, 0.75, _f2(s), 11, TEXT_DARK)


# --- table layout ---------------------------------------------------------
# Canvas inches. The band runs from under the template header down to just
# above the footer, matching where the bespoke slides put their own tables.
_TBL_LEFT, _TBL_TOP, _TBL_WIDTH = 0.9, 1.6, 11.5
_TBL_BOTTOM = 6.6
_TBL_FONT_PT = 10.67       # exactly 8pt once scaled to the template (× 0.75)
_TBL_LINE_RATIO = 1.4     # line height as a multiple of the font size
_TBL_PAD_RATIO = 0.7      # cell padding, likewise
_LATIN_EM = 0.55          # latin character width, in ems
_TBL_MIN_COL = 0.9

# What the last visible row says when the box could not hold every row. Callers
# in a Japanese deck pass `_MORE_ROWS_JA`; the source document decides.
_MORE_ROWS_EN = '… and {count} more row(s) — see the source document'
_MORE_ROWS_JA = '… ほか{count}行（元資料に記載）'


def _table(cv: _Canvas, entry: dict) -> None:
    """Render a real table — the layout a markdown table deserves.

    Without this, `extra_slide_classifier` had nowhere to send a table but
    `bullets`, which joins each row's cells with an em dash and throws the
    column structure away. A responsibility-split matrix rendered that way is
    unreadable.

    A SCHEDULE table is drawn as the chart it describes instead: phases, a
    role-by-month matrix and dated milestones all read at a glance as a timeline
    and one column at a time as a grid. `timeline_data.build_chart` returns None
    for everything else, so an ordinary table still comes through here.
    """
    headers = [str(head) for head in _get(entry, 'headers')]
    rows = [[str(cell) for cell in row] for row in _get(entry, 'rows')]
    if _draw_as_chart(cv, headers, rows):
        return
    draw_table_in_box(
        cv.slide, cv._x(_TBL_LEFT), cv._y(_TBL_TOP),
        cv._x(_TBL_WIDTH), cv._y(_TBL_BOTTOM - _TBL_TOP),
        headers, rows,
        font_pt=_TBL_FONT_PT * cv.fs,
        more_label=entry.get('more_label'),
    )


def _draw_as_chart(cv: _Canvas, headers: list, rows: list) -> bool:
    """Draw a schedule table as a timeline; True when it was one."""
    from lib.timeline_data import build_chart
    from lib.timeline_layouts import draw_chart_in_box

    chart = build_chart([dict(zip(headers, row)) for row in rows]) if headers else None
    if chart is None:
        return False
    draw_chart_in_box(
        cv.slide, cv._x(_TBL_LEFT), cv._y(_TBL_TOP),
        cv._x(_TBL_WIDTH), cv._y(_TBL_BOTTOM - _TBL_TOP),
        chart, font_pt=_TBL_FONT_PT * cv.fs)
    return True


def draw_table_in_box(slide, left, top, width, height, headers, rows,
                      font_pt: float = 8.0, more_label: str = None) -> None:
    """Draw a table filling the EMU box (*left*, *top*, *width*, *height*).

    Shared by the `table` extra-slide layout and by the renderer when a table
    takes over an image placeholder the source had no picture for, so both look
    the same wherever they land.

    Rows that do not fit the box are not silently lost: the last visible row
    says how many were left out, and the source markdown still holds them all.
    """
    headers = list(headers or [])
    rows = [list(row) for row in (rows or [])]
    if not rows and not headers:
        return
    n_cols = max([len(headers)] + [len(row) for row in rows]) or 1
    headers = (headers + [''] * n_cols)[:n_cols]
    rows = [(row + [''] * n_cols)[:n_cols] for row in rows]

    box_w = width / _EMU_IN
    box_h = height / _EMU_IN
    # Line height and cell padding scale with the type, so a table drawn at
    # the layout's 11pt canvas size and one drawn at 8pt inside a template
    # placeholder pack their rows the same way.
    line_h = font_pt * _TBL_LINE_RATIO / 72
    widths = _column_widths(headers, rows, n_cols, box_w)
    visible, dropped = _rows_that_fit(headers, rows, widths, box_h, font_pt, line_h)
    if dropped:
        # The deck's own language, not this file's: the marker was written in
        # Vietnamese and printed as `… và 2 hàng nữa` on an English proposal,
        # the one line of the deck that was not in the client's language.
        template = more_label or _MORE_ROWS_EN
        visible = visible[:-1] + [[template.format(count=dropped + 1)]
                                  + [''] * (n_cols - 1)]

    has_head = any(h.strip() for h in headers)
    body = ([headers] if has_head else []) + visible
    # Row heights come from the SAME formula that decided which rows fit, so the
    # stored table is as tall as the fitting pass assumed. Sizing it at bare
    # `line_h` per row (the old total) ignored the padding that formula already
    # counted, and stored 0.16" rows against the template's 0.35" — a table that
    # read as squeezed next to every hand-authored one.
    row_heights = [_row_height(row, widths, font_pt, line_h) for row in body]
    table = slide.shapes.add_table(
        len(body), n_cols, int(left), int(top), int(width),
        int(sum(row_heights) * _EMU_IN),
    ).table
    _clear_table_style(table)
    for col, col_w in zip(table.columns, widths):
        col.width = Emu(int(col_w * _EMU_IN))
    for row_obj, row_h in zip(table.rows, row_heights):
        row_obj.height = Emu(int(row_h * _EMU_IN))

    for r, row in enumerate(body):
        head_row = has_head and r == 0
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TBL_HEAD_FILL if head_row else TBL_BODY_FILL
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _set_cell_borders(cell, TBL_BORDER, TBL_BORDER_PT)
            cell.margin_left = cell.margin_right = Emu(int(TBL_MARGIN_X_IN * _EMU_IN))
            cell.margin_top = cell.margin_bottom = Emu(int(TBL_MARGIN_Y_IN * _EMU_IN))
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.LEFT
                for run in para.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_pt)
                    run.font.bold = head_row
                    run.font.color.rgb = TBL_HEAD_TEXT if head_row else TBL_BODY_TEXT


def _clear_table_style(table) -> None:
    """Strip the theme table style python-pptx attaches to a new table.

    `add_table` stamps `firstRow="1" bandRow="1"` plus a `tableStyleId`, so
    PowerPoint paints its own header band and zebra striping UNDER the explicit
    per-cell fills set below. Every table in the SVN template ships a bare
    `a:tblPr` and carries its look in the cells themselves — matching that is
    what keeps the appendix table from picking up banding no template table has.
    """
    from pptx.oxml.ns import qn

    tblPr = table._tbl.find(qn('a:tblPr'))
    if tblPr is None:
        return
    for attr in ('firstRow', 'firstCol', 'lastRow', 'lastCol', 'bandRow', 'bandCol'):
        tblPr.attrib.pop(attr, None)
    for style_id in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(style_id)


def _set_cell_borders(cell, color: RGBColor, width_pt: float = 0.75) -> None:
    """Draw all four borders of a table cell.

    python-pptx exposes cell fill but not cell borders, and a table's default
    theme style leaves them off — which is why the drawn tables came out as
    floating columns of text with only the header band visible. The four
    `a:lnL/lnR/lnT/lnB` elements have to be written directly, in that order:
    `a:tcPr` is an ordered sequence and PowerPoint rejects the file otherwise.
    """
    from pptx.oxml.ns import qn

    props = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnB', 'a:lnT', 'a:lnR', 'a:lnL'):  # reversed: each inserts at 0
        for existing in props.findall(qn(tag)):
            props.remove(existing)
        line = props.makeelement(qn(tag), {'w': str(int(width_pt * 12700)),
                                           'cap': 'flat', 'cmpd': 'sng',
                                           'algn': 'ctr'})
        fill = line.makeelement(qn('a:solidFill'), {})
        srgb = fill.makeelement(qn('a:srgbClr'), {'val': f'{color}'})
        fill.append(srgb)
        line.append(fill)
        props.insert(0, line)


def _column_widths(headers: list, rows: list, n_cols: int, box_w: float) -> list:
    """Split *box_w* inches by how much text each column carries."""
    min_col = min(_TBL_MIN_COL, box_w / n_cols)
    demands = []
    for c in range(n_cols):
        longest = max([len(headers[c])] + [len(row[c]) for row in rows] or [1])
        demands.append(max(float(longest), 4.0))
    scale = box_w / sum(demands)
    widths = [max(min_col, d * scale) for d in demands]
    # The floor can push the total over the box; take the excess back off the
    # widest column, which has the most room to give.
    excess = sum(widths) - box_w
    if excess > 0:
        widest = max(range(n_cols), key=lambda i: widths[i])
        widths[widest] = max(min_col, widths[widest] - excess)
    return widths


def _rows_that_fit(headers: list, rows: list, widths: list, box_h: float,
                   font_pt: float, line_h: float) -> tuple:
    """Return (rows that fit *box_h* inches, how many were dropped)."""
    used = (_row_height(headers, widths, font_pt, line_h)
            if any(h.strip() for h in headers) else 0.0)
    kept = []
    for row in rows:
        height = _row_height(row, widths, font_pt, line_h)
        if kept and used + height > box_h:
            return kept, len(rows) - len(kept)
        used += height
        kept.append(row)
    return kept, 0


def _row_height(cells: list, widths: list, font_pt: float, line_h: float) -> float:
    """Inches one row needs, from its widest wrapping cell."""
    em = font_pt / 72
    lines = 1
    for cell, width in zip(cells, widths):
        # CJK sets about one em per character, latin roughly half. Counting
        # every character as full width — the earlier estimate — made a mostly
        # latin cell look twice as tall as it renders, and cost the table the
        # rows it had room for.
        text_in = sum(em if ord(ch) > 0x2E7F else em * _LATIN_EM
                      for ch in str(cell))
        lines = max(lines, math.ceil(text_in / width) if width else 1)
    return lines * line_h + font_pt * _TBL_PAD_RATIO / 72


def _hero(cv: _Canvas, entry: dict) -> None:
    msg = entry.get('message') or entry.get('content', {}).get('message') or ''
    cv.rect(0.7, 3.0, 12.0, 1.6, SUN_RED)
    cv.text(0.9, 3.0, 11.6, 1.6, msg, 22, WHITE, bold=True, italic=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _diagram(cv: _Canvas, entry: dict) -> None:
    """A fenced drawing, set as a picture on the template's own white card.

    The lead-in prose sits above it when the chapter wrote any, and the picture
    takes whatever height is left, centred. When no monospace face is available
    the block falls back to the list layouts rather than to nothing — worse, but
    still the author's words.
    """
    from lib.diagram_image import render, render_boxes
    from lib.list_layouts import BAND_BOTTOM, BAND_LEFT, BAND_TOP, BAND_WIDTH, _card

    text = entry.get('diagram') or (entry.get('content') or {}).get('diagram') or ''
    lead = entry.get('lead_in') or []
    if not text.strip():
        return _bullets(cv, {'bullets': lead})

    top = BAND_TOP
    if lead:
        # One condensed line per point, so the diagram keeps the room it needs.
        used = min(1.6, 0.34 * len(lead) + 0.2)
        cv.text(BAND_LEFT + 0.25, top, BAND_WIDTH - 0.5, used,
                [f'■  {line}' for line in lead[:4]], 12, TEXT_GREY)
        top += used + 0.15

    box_w, box_h = BAND_WIDTH, BAND_BOTTOM - top
    # The canvas is authored at 13.333in wide and scaled to the real slide, so
    # the picture is MEASURED against the box in real inches and then placed
    # back on the canvas by dividing the scale out again.
    # Box art first: recovering the rectangles the author drew and redrawing
    # them properly beats photographing the characters. A block that is not box
    # art — or one the parser cannot read with confidence — falls back to the
    # monospace picture, which is always faithful even when it is plain.
    stream, drawn_w, drawn_h = render_boxes(text, box_w * cv.sx, box_h * cv.sy)
    if stream is None:
        stream, drawn_w, drawn_h = render(text, box_w * cv.sx, box_h * cv.sy)
    if stream is None:
        return _bullets(cv, {'bullets': lead + [' '.join(text.split())]})

    # The card is drawn round the picture rather than round the band. A diagram
    # keeps its own proportions, so a band-sized card leaves a broad blank
    # margin down one axis — the card then reads as an empty panel that happens
    # to contain a drawing, instead of as the drawing's own mount.
    pad = 0.18
    card_w = min(box_w, drawn_w / cv.sx + 2 * pad)
    card_h = min(box_h, drawn_h / cv.sy + 2 * pad)
    card_x = BAND_LEFT + (box_w - card_w) / 2
    card_y = top + (box_h - card_h) / 2
    _card(cv, card_x, card_y, card_w, card_h)
    left = Emu(int((card_x * cv.sx + (card_w * cv.sx - drawn_w) / 2) * _EMU_IN))
    top_emu = Emu(int((card_y * cv.sy + (card_h * cv.sy - drawn_h) / 2) * _EMU_IN))
    cv.slide.shapes.add_picture(stream, left, top_emu,
                                Emu(int(drawn_w * _EMU_IN)),
                                Emu(int(drawn_h * _EMU_IN)))


_LAYOUTS = {
    'diagram': _diagram,
    'table': _table,
    'bullets': _bullets,
    'numbered_points': _numbered_points,
    'card_grid': _card_grid,
    'comparison_2': _comparison_2,
    'process_flow': _process_flow,
    'hero': _hero,
}


# ---------------------------------------------------------------------------
# Small helpers — tolerate both top-level keys and a nested `content` dict
# ---------------------------------------------------------------------------

def _get(entry: dict, key: str) -> list:
    return entry.get(key) or (entry.get('content') or {}).get(key) or []


def _f(obj, key: str) -> str:
    if isinstance(obj, dict):
        return str(obj.get(key) or '')
    return str(obj) if key in ('title', 'header') else ''


def _f2(obj) -> str:
    if isinstance(obj, dict):
        return str(obj.get('text') or obj.get('title') or obj.get('header') or '')
    return str(obj)


# Schedule/team layouts live in their own module (they draw in absolute
# template inches, not on this canvas); imported last so they can reuse the
# brand constants above.
from lib.gantt_layout import register as _register_gantt  # noqa: E402
from lib.org_chart_layout import register as _register_org_chart  # noqa: E402

_register_gantt(_LAYOUTS)
_register_org_chart(_LAYOUTS)
